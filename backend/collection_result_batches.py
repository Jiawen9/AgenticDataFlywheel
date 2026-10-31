"""Frozen per-dispatch collection inputs, independent of their reusable source."""
from __future__ import annotations

import copy
import hashlib
import io
from pathlib import Path
import uuid

from openpyxl import Workbook, load_workbook

from .collection_runs import CollectionRunError, _identifier
from .data_store import ArtifactStore, DATA_ROOT, RecordStore, RevisionConflict
from .data_store.locking import batch_lock
from .data_store.paths import contained_path
from .data_store.registry import utc_now
from .display_identifiers import DisplayIdentifierStore
from .file_io import io_path
from .identifiers import allocate_id
from .task_generation.collection_batches import COLLECTION_COLUMNS, payload_digest


class CollectionResultBatchStore:
    """One immutable result batch per source/request; no device work is performed."""

    def __init__(self, root: Path = DATA_ROOT, *, sources=None):
        self.root = contained_path(Path(root))
        self.records = RecordStore(self.root)
        self.artifacts = ArtifactStore(self.root)
        self.sources = sources

    @staticmethod
    def request_key(source_batch_id, dispatch_key):
        return payload_digest([source_batch_id, dispatch_key])

    def get(self, batch_id: str, *, require_workbook=False) -> dict | None:
        _identifier(batch_id, "采集结果批次编号")
        record = self.records.get("collection_result_batches", batch_id)
        if record is None:
            return None
        try:
            artifact = self.artifacts.get(batch_id, "00_collection")
            if artifact is None:
                raise ValueError("missing collection artifact")
            payload = self.artifacts.read_payload(artifact)
            if (payload.get("batch_id") != batch_id or payload.get("batch_role") != "collection_result"
                    or payload_digest(payload) != record["payload_sha256"]):
                raise ValueError("result batch checksum mismatch")
            path = contained_path(self.root, "batches", batch_id, "00_collection", payload["filename"])
            if require_workbook:
                path = self.artifacts.resolve_file(artifact, payload["filename"])
            return {**payload, "workbook_path": str(path)}
        except (OSError, ValueError, KeyError) as exc:
            raise CollectionRunError("采集结果批次快照缺失、损坏或校验失败") from exc

    def list(self, task_source_batch_id: str | None = None, *, include_published=True) -> list[dict]:
        from .batch_lifecycle import is_batch_active
        if task_source_batch_id is not None:
            _identifier(task_source_batch_id, "任务来源批次编号")
        values = [self.get(record["batch_id"]) for record in self.records.list("collection_result_batches")
                  if (task_source_batch_id is None or record["task_source_batch_id"] == task_source_batch_id)
                  and (include_published or is_batch_active(record["batch_id"], self.root))]
        return sorted(values, key=lambda value: (value["created_at"], value["batch_id"]))

    def _source(self, source_batch_id):
        if self.sources is None:
            from .collection_runs import CollectionRunStore
            self.sources = CollectionRunStore(self.root)
        try:
            source = self.sources.batch(source_batch_id, require_workbook=True)
        except CollectionRunError:
            raise
        except (OSError, ValueError) as exc:
            raise CollectionRunError("任务来源的冻结快照或 Excel 不可用") from exc
        if source.get("batch_role") == "collection_result":
            raise CollectionRunError("新采集请选择原任务来源，不能向已有结果批次追加采集")
        return source

    def _workbook(self, source):
        path = Path(source.get("workbook_path") or contained_path(self.root, "system", "task_generation",
            "collection_batches", source["batch_id"], source["filename"]))
        path = contained_path(self.root, str(path))
        try:
            content = io_path(path).read_bytes()
        except OSError as exc:
            raise CollectionRunError("任务来源的冻结 Excel 不可用") from exc
        if hashlib.sha256(content).hexdigest() != source.get("workbook_sha256"):
            raise CollectionRunError("任务来源 Excel 与冻结摘要不一致")
        return content

    @staticmethod
    def _selected(source, content, metadata):
        associations = metadata.get("phone_apps")
        if (not isinstance(associations, list) or not associations
                or any(not isinstance(row, dict) or not row.get("phone_id") or not row.get("app")
                       for row in associations)):
            raise CollectionRunError("采集结果批次需要冻结有效的手机/App 关联")
        apps = {str(row["app"]) for row in associations
                if (not metadata.get("phone_id") or row["phone_id"] == metadata["phone_id"])
                and (not metadata.get("app") or row["app"] == metadata["app"])}
        tasks = {str(task["collection_case_id"]): task for task in source["snapshot"]["tasks"]}
        selected, rows, seen = [], [], set()
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.active
                values = sheet.iter_rows(values_only=True)
                headers = list(next(values, ()))
                while headers and headers[-1] is None:
                    headers.pop()
                if headers != COLLECTION_COLUMNS:
                    raise CollectionRunError("任务来源必须保留完整17列表头及顺序")
                for row in values:
                    if not any(value is not None and value != "" for value in row):
                        continue
                    if any(value is not None and value != "" for value in row[17:]):
                        raise CollectionRunError("任务来源包含17列模板以外的数据")
                    if str(row[2]) not in apps:
                        continue
                    case_id = str(row[0]).strip()
                    if case_id not in tasks or case_id in seen:
                        raise CollectionRunError("任务来源 Excel 的用例编号与冻结快照不匹配")
                    if str(row[6]).strip() != str(tasks[case_id].get("task", "")).strip():
                        raise CollectionRunError("任务来源 Excel 的任务与冻结快照不匹配")
                    seen.add(case_id)
                    selected.append(copy.deepcopy(tasks[case_id]))
                    rows.append(list(row[:17]) + [None] * max(0, 17 - len(row)))
                title = sheet.title
            finally:
                book.close()
        except CollectionRunError:
            raise
        except Exception as exc:
            raise CollectionRunError("无法读取任务来源的17列工作簿") from exc
        if not selected:
            raise CollectionRunError("当前手机/App 条件没有匹配任务，未创建采集结果批次")
        return selected, rows, title

    @staticmethod
    def _write_workbook(path, rows, title):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        for row in [COLLECTION_COLUMNS, *rows]:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        try:
            book.save(io_path(path))
        finally:
            book.close()

    def prepare(self, source_batch_id: str, *, dispatch_key: str, metadata: dict,
                workbook_sha256: str | None = None) -> tuple[dict, bool]:
        _identifier(source_batch_id, "任务来源批次编号")
        _identifier(dispatch_key, "request_id")
        metadata = copy.deepcopy(metadata)
        request_key = self.request_key(source_batch_id, dispatch_key)
        # No active-source lock: a published frozen source remains reusable.
        with batch_lock(self.root, "collection_result_" + request_key):
            existing = self.records.get("collection_result_requests", request_key)
            if existing:
                request = {"task_source_batch_id": source_batch_id, "dispatch_key": dispatch_key,
                           "metadata": metadata, "workbook_sha256": workbook_sha256 or existing["request"]["workbook_sha256"]}
                if payload_digest(request) != existing["request_digest"]:
                    raise CollectionRunError("相同 request_id 对应的采集配置或任务来源不同")
                batch = self.get(existing["batch_id"], require_workbook=True)
                if batch is None:
                    raise CollectionRunError("已登记采集结果批次缺失，不能重新创建或重新采集")
                return batch, False
            source = self._source(source_batch_id)
            if workbook_sha256 is not None and workbook_sha256 != source["workbook_sha256"]:
                raise CollectionRunError("上传文件与冻结的任务来源 Excel 不一致")
            content = self._workbook(source)
            tasks, rows, sheet = self._selected(source, content, metadata)
            request = {"task_source_batch_id": source_batch_id, "dispatch_key": dispatch_key,
                       "metadata": metadata, "workbook_sha256": source["workbook_sha256"]}
            request_digest = payload_digest(request)
            try:
                batch_id = allocate_id("b", self.root, key=["collection_result", source_batch_id, dispatch_key],
                                       request_fingerprint=request_digest)
            except RevisionConflict as exc:
                raise CollectionRunError("相同 request_id 对应的采集配置或任务来源不同") from exc
            run_id = ("eval_" + payload_digest(dispatch_key)
                      if metadata.get("run_mode") == "modeliter" and metadata.get("collection_artifacts_version") == 1
                      else "cr_" + payload_digest([batch_id, dispatch_key]))
            filename = f"collection-batch-{batch_id}.xlsx"
            created = utc_now()
            snapshot = {**copy.deepcopy(source["snapshot"]), "batch_id": batch_id,
                        "batch_role": "collection_result", "task_source_batch_id": source_batch_id,
                        "tasks": tasks, "task_count": len(tasks)}
            temporary = contained_path(self.root, "tmp", "collection_result_batches", uuid.uuid4().hex + ".xlsx")
            io_path(temporary.parent).mkdir(parents=True, exist_ok=True)
            try:
                self._write_workbook(temporary, rows, sheet)
                checksum = hashlib.sha256(io_path(temporary).read_bytes()).hexdigest()
                payload = {"schema_version": 1, "batch_id": batch_id, "batch_role": "collection_result",
                    "task_source_batch_id": source_batch_id,
                    "task_source_snapshot_sha256": payload_digest(source["snapshot"]),
                    "task_source_workbook_sha256": source["workbook_sha256"],
                    "source_job_id": source.get("source_job_id"), "kind": source["kind"],
                    "job_status": source.get("job_status"), "knowledge_base_version": source.get("knowledge_base_version"),
                    "created_at": created, "task_count": len(tasks), "apps": list(dict.fromkeys(task["app"] for task in tasks)),
                    "filename": filename, "original_filename": source.get("original_filename", source["filename"]),
                    "description": source.get("description", ""),
                    "download_url": f"/api/phone-factory/batches/{batch_id}/workbook",
                    "sheet_name": sheet, "snapshot": snapshot, "warnings": source.get("warnings", []),
                    "workbook_sha256": checksum, "collection_run_id": run_id,
                    "agent_type": metadata.get("agent_type", "newagent"),
                    "execution_request_id": dispatch_key, "execution_metadata": metadata}
                entries = DisplayIdentifierStore(self.root).new_batch_entries(batch_id, tasks)
                self.artifacts.publish_many(batch_id, [{"stage": "00_collection", "payload": payload,
                    "workbooks": {filename: temporary}, "metadata": {"batch_role": "collection_result",
                        "task_source_batch_id": source_batch_id, "kind": source["kind"], "task_count": len(tasks)},
                    "source_refs": [{"kind": "collection_task_source", "batch_id": source_batch_id,
                        "snapshot_sha256": payload["task_source_snapshot_sha256"], "workbook_sha256": source["workbook_sha256"]}]}],
                    record_entries=[{"namespace": "collection_result_batches", "key": batch_id,
                        "payload": {"batch_id": batch_id, "task_source_batch_id": source_batch_id,
                            "collection_run_id": run_id, "payload_sha256": payload_digest(payload), "created_at": created},
                        "expected_revision": 0},
                        {"namespace": "collection_result_requests", "key": request_key,
                         "payload": {"batch_id": batch_id, "collection_run_id": run_id,
                             "request": request, "request_digest": request_digest}, "expected_revision": 0}, *entries])
            finally:
                io_path(temporary).unlink(missing_ok=True)
            return self.get(batch_id, require_workbook=True), True
