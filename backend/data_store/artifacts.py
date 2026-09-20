"""One recoverable current JSON/Excel artifact per batch and stage."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import uuid
from itertools import chain
from pathlib import Path
from typing import Any

from .paths import DATA_ROOT, contained_path
from .registry import RecordStore, utc_now, RevisionConflict
from .locking import batch_lock

_SAFE_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,127}\Z")
_RESERVED = re.compile(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?\Z", re.IGNORECASE)


def _identifier(value: str, label: str) -> str:
    if not isinstance(value, str) or not _SAFE_ID.fullmatch(value) or value.endswith(".") or _RESERVED.fullmatch(value):
        raise ValueError(f"Invalid {label}")
    return value


def _filename(value: str) -> str:
    if not isinstance(value, str) or not value or len(value) > 180 or value in {".", ".."} or re.search(r'[<>:"/\\|?*\x00-\x1f]', value) or value.endswith((".", " ")) or _RESERVED.fullmatch(value):
        raise ValueError("Invalid artifact filename")
    return value


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2) + "\n").encode("utf-8")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_bytes(path: Path, value: bytes) -> None:
    with path.open("xb") as output:
        output.write(value)
        output.flush()
        os.fsync(output.fileno())


def _write_tables(path: Path, tables: dict[str, list[dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    if not tables:
        raise ValueError("Excel tables must contain at least one worksheet")
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()
    try:
        for sheet_name, rows in tables.items():
            if not isinstance(sheet_name, str) or not sheet_name or len(sheet_name) > 31 or re.search(r"[\\/*?:\[\]]", sheet_name) or sheet_name.lower() in used_names:
                raise ValueError("Invalid or duplicate Excel worksheet name")
            used_names.add(sheet_name.lower())
            sheet = workbook.create_sheet(sheet_name)
            columns: list[str] = []
            for row in rows:
                if not isinstance(row, dict):
                    raise TypeError("Excel table rows must be dictionaries")
                for column in row:
                    if not isinstance(column, str):
                        raise TypeError("Excel column names must be strings")
                    if column not in columns:
                        columns.append(column)
            if len(rows) > 1048575 or len(columns) > 16384:
                raise ValueError("Excel worksheet exceeds its row or column limit")
            values_iter = chain([columns], ([row.get(column) for column in columns] for row in rows))
            for row_number, values in enumerate(values_iter, start=1):
                for column_number, value in enumerate(values, start=1):
                    if isinstance(value, (dict, list, tuple)):
                        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
                    if isinstance(value, str) and len(value) > 32767:
                        raise ValueError(f"Excel cell {sheet_name}!R{row_number}C{column_number} exceeds 32767 characters; original JSON text must be retained")
                    cell = sheet.cell(row_number, column_number, value)
                    # Treat arbitrary tasks, identifiers and model text as literal
                    # strings, including =, +, -, and @ prefixes.
                    if isinstance(value, str):
                        cell.data_type = "s"
                        cell.number_format = "@"
                    if row_number == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="0F766E")
            sheet.freeze_panes = "A2"
            if columns:
                sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)
    finally:
        workbook.close()
    with path.open("r+b") as output:
        os.fsync(output.fileno())


class ArtifactStore:
    def __init__(self, root: Path = DATA_ROOT):
        self.root = contained_path(Path(root))
        self.records = RecordStore(self.root)

    @staticmethod
    def _key(batch_id: str, stage: str, version: str | None = None) -> str:
        parts = [_identifier(batch_id, "batch ID"), _identifier(stage, "stage")]
        if version is not None:
            parts.append(_identifier(version, "version"))
        return "/".join(parts)

    def batch_lock(self, batch_id: str):
        return batch_lock(self.root, batch_id)

    def publish(self, batch_id: str, stage: str, payload: Any,
                workbooks: dict[str, Path] | None = None,
                tables: dict[str, list[dict]] | None = None,
                source_refs: list | None = None, metadata: dict | None = None,
                expected_version: str | None = None) -> dict:
        return self.publish_many(batch_id, [dict(stage=stage, payload=payload,
            workbooks=workbooks, tables=tables, source_refs=source_refs,
            metadata=metadata, expected_version=expected_version)])[0]

    def _remove(self, path: Path) -> None:
        target = contained_path(self.root, str(path))
        if target == self.root or not any(target.is_relative_to(self.root / part) for part in ("tmp", "batches")):
            raise ValueError("Refusing to remove outside artifact transaction directories")
        if target.is_dir():
            shutil.rmtree(target)

    def _recover(self, batch_id: str) -> None:
        """Called with batch lock held, before any reading or writing."""
        parent = contained_path(self.root, "tmp", "artifact_transactions", batch_id)
        if not parent.exists():
            return
        for transaction in sorted(parent.iterdir()):
            journal = transaction / "commit.json"
            if not journal.is_file():
                self._remove(transaction)
                continue
            state = json.loads(journal.read_text(encoding="utf-8"))
            if state.get("batch_id") != batch_id or state.get("op_id") != transaction.name:
                raise ValueError("Invalid artifact recovery journal")
            committed = self.records.get("artifact_commits", batch_id)
            saved = bool(committed and committed.get("op_id") == state["op_id"])
            for change in reversed(state["changes"]):
                stage = _identifier(change["stage"], "stage")
                final = contained_path(self.root, "batches", batch_id, stage)
                previous, prepared = transaction / "previous" / stage, transaction / "prepared" / stage
                if not saved:
                    if previous.exists():
                        self._remove(final)
                        previous.rename(final)
                    elif not change["had_previous"] and not prepared.exists():
                        self._remove(final)
            self._remove(transaction)

    def recover(self) -> None:
        parent = self.root / "tmp" / "artifact_transactions"
        if parent.is_dir():
            for directory in parent.iterdir():
                if directory.is_dir():
                    with self.batch_lock(directory.name):
                        self._recover(directory.name)

    @staticmethod
    def _source_refs(batch_id: str, stage: str, refs: list) -> list:
        result = []
        for ref in refs:
            # A replacement cannot reference its own superseded revision.
            if ref.get("batch_id") == batch_id and ref.get("stage") == stage:
                result.extend(ArtifactStore._source_refs(batch_id, stage, ref.get("source_refs", [])))
            else:
                result.append(ref)
        return result

    def publish_many(self, batch_id: str, entries: list[dict], *,
                     replace_legacy: bool = False, record_entries: list[dict] | None = None) -> list[dict]:
        """Commit whole stages together; version is a CAS token, never a folder.

        source_stages references earlier entries in this publication.
        replace_legacy is reserved for the backed-up offline migration.
        """
        _identifier(batch_id, "batch ID")
        if not entries:
            return []
        stages = [_identifier(item["stage"], "stage") for item in entries]
        if len(stages) != len(set(stages)):
            raise ValueError("A stage may only occur once in a publication")
        for entry in entries:
            _json_bytes(entry['payload'])
            names = {'result.json', 'manifest.json'} | ({'result.xlsx'} if entry.get('tables') is not None else set())
            for name in (entry.get('workbooks') or {}):
                _filename(name)
                if not name.lower().endswith('.xlsx') or name.casefold() in {value.casefold() for value in names}:
                    raise ValueError('Artifact workbooks require unique .xlsx filenames')
                names.add(name)
        with self.batch_lock(batch_id):
            self._recover(batch_id)
            from ..batch_lifecycle import ensure_batch_active
            ensure_batch_active(batch_id, self.root)
            op_id = uuid.uuid4().hex
            transaction = contained_path(self.root, "tmp", "artifact_transactions", batch_id, op_id)
            manifests, changes, writes = [], [], []
            try:
                for entry in entries:
                    stage = entry["stage"]
                    current = self.records.get("artifacts", self._key(batch_id, stage))
                    legacy = [item for item in self.records.list("artifacts")
                              if item.get("batch_id") == batch_id and item.get("stage") == stage
                              and item.get("schema_version", 1) < 2]
                    if legacy and not replace_legacy:
                        raise ValueError("旧阶段存在历史版本，请先完成单份过程件迁移")
                    expected = entry.get("expected_version")
                    if expected is not None and (current or {}).get("version") != expected:
                        raise RevisionConflict("当前阶段已更新，请刷新后重试")
                    payload_bytes = _json_bytes(entry["payload"])
                    metadata = json.loads(_json_bytes(entry.get("metadata") or {}))
                    refs = self._source_refs(batch_id, stage, json.loads(_json_bytes(entry.get("source_refs") or [])))
                    for source_stage in entry.get("source_stages", []):
                        source = next((item for item in manifests if item["stage"] == source_stage), None)
                        if source is None:
                            raise ValueError("source_stages must refer to a preceding publication")
                        refs.append(source)
                    workbooks, tables = entry.get("workbooks") or {}, entry.get("tables")
                    names = {"result.json", "manifest.json"}
                    if tables is not None:
                        names.add("result.xlsx")
                    workbook_hashes = {}
                    for name, source in workbooks.items():
                        _filename(name)
                        if not name.lower().endswith(".xlsx") or name.casefold() in {value.casefold() for value in names}:
                            raise ValueError("Artifact workbooks require unique .xlsx filenames")
                        names.add(name)
                        workbook_hashes[name] = _sha256(Path(source))
                    identity = {"payload": entry["payload"], "tables": tables, "workbooks": workbook_hashes,
                                "source_refs": refs, "metadata": {key: value for key, value in metadata.items()
                                if key not in {"job_id", "run_id", "preprocessing_job_id", "updated_at", "created_at"}}}
                    input_fingerprint = hashlib.sha256(_json_bytes(identity)).hexdigest()
                    if current and current.get("input_fingerprint") == input_fingerprint:
                        manifest = self._manifest(current)
                        for name in names - {"manifest.json"}:
                            self._resolve_file(manifest, name)
                        manifests.append(manifest)
                        continue
                    revision = int((current or {}).get("revision", 0)) + 1
                    final = contained_path(self.root, "batches", batch_id, stage)
                    prepared = transaction / "prepared" / stage
                    prepared.mkdir(parents=True, exist_ok=False)
                    _write_bytes(prepared / "result.json", payload_bytes)
                    if tables is not None:
                        _write_tables(prepared / "result.xlsx", tables)
                    for name, source in workbooks.items():
                        with Path(source).open("rb") as incoming, (prepared / name).open("xb") as output:
                            shutil.copyfileobj(incoming, output)
                            output.flush()
                            os.fsync(output.fileno())
                    files = [{"name": file.name, "kind": "json" if file.suffix == ".json" else "excel",
                              "path": (final / file.name).relative_to(self.root).as_posix(),
                              "sha256": _sha256(file), "size": file.stat().st_size}
                             for file in sorted(prepared.iterdir())]
                    manifest = {"schema_version": 2, "batch_id": batch_id, "stage": stage,
                                "version": f"r{revision:012d}", "revision": revision, "created_at": utc_now(),
                                "content_hash": hashlib.sha256(payload_bytes).hexdigest(),
                                "input_fingerprint": input_fingerprint, "source_refs": refs,
                                "metadata": metadata, "files": files}
                    if entry.get("legacy_aliases"):
                        manifest["legacy_aliases"] = list(entry["legacy_aliases"])
                    _write_bytes(prepared / "manifest.json", _json_bytes(manifest))
                    changes.append({"stage": stage, "had_previous": final.exists(),
                                    "legacy_keys": [self._key(batch_id, stage, item["version"]) for item in legacy]})
                    writes.append({"namespace": "artifacts", "key": self._key(batch_id, stage),
                                   "payload": manifest, "expected_revision": (current or {}).get("storage_revision", 0)})
                    manifests.append(manifest)
                if not changes:
                    if record_entries:
                        self.records.put_many(record_entries)
                    return manifests
                (transaction / "previous").mkdir(parents=True, exist_ok=True)
                _write_bytes(transaction / "commit.json", _json_bytes({"batch_id": batch_id, "op_id": op_id, "changes": changes}))
                for change in changes:
                    final = contained_path(self.root, "batches", batch_id, change["stage"])
                    final.parent.mkdir(parents=True, exist_ok=True)
                    if final.exists():
                        final.rename(transaction / "previous" / change["stage"])
                    (transaction / "prepared" / change["stage"]).rename(final)
                writes.extend(record_entries or [])
                writes.append({"namespace": "artifact_commits", "key": batch_id, "payload": {"op_id": op_id}})
                deletes = [{"namespace": "artifacts", "key": key} for change in changes for key in change["legacy_keys"]]
                self.records.put_many(writes, deletes=deletes)
                self._remove(transaction)
                return manifests
            except BaseException:
                if (transaction / "commit.json").exists():
                    self._recover(batch_id)
                else:
                    self._remove(transaction)
                raise

    def list(self, batch_id: str | None = None, stage: str | None = None) -> list[dict]:
        if batch_id is not None:
            _identifier(batch_id, "batch ID")
        if stage is not None:
            _identifier(stage, "stage")
        return [self._manifest(item) for item in self.records.list("artifacts")
                if (batch_id is None or item["batch_id"] == batch_id) and (stage is None or item["stage"] == stage)]

    def get(self, batch_id: str, stage: str, version: str | None = None) -> dict | None:
        record = self.records.get("artifacts", self._key(batch_id, stage))
        if record is not None:
            return self._manifest(record) if version is None or version == record["version"] or version in record.get("legacy_aliases", []) else None
        # Only read compatibility for stores awaiting offline migration.
        if version is not None:
            record = self.records.get("artifacts", self._key(batch_id, stage, version))
        else:
            candidates = self.list(batch_id, stage)
            record = max(candidates, key=lambda item: (item["created_at"], item["version"])) if candidates else None
        return self._manifest(record) if record else None

    @staticmethod
    def _manifest(record: dict) -> dict:
        return {key: value for key, value in record.items() if key != "storage_revision"}

    def resolve_file(self, manifest: dict, filename: str) -> Path:
        """Verify a path. Concurrent consumers should use read_file/read_payload."""
        with self.batch_lock(manifest["batch_id"]):
            self._recover(manifest["batch_id"])
            return self._resolve_file(manifest, filename)

    def read_file(self, manifest: dict, filename: str) -> bytes:
        with self.batch_lock(manifest["batch_id"]):
            self._recover(manifest["batch_id"])
            return self._resolve_file(manifest, filename).read_bytes()

    def read_payload(self, manifest: dict) -> Any:
        return json.loads(self.read_file(manifest, "result.json"))

    def _resolve_file(self, manifest: dict, filename: str) -> Path:
        _filename(filename)
        registered = self.get(manifest["batch_id"], manifest["stage"], manifest["version"])
        if registered is None:
            raise FileNotFoundError("Artifact snapshot is not registered; current stage changed")
        file = next((item for item in registered["files"] if item["name"] == filename), None)
        if file is None:
            raise FileNotFoundError("Artifact file is not in the snapshot manifest")
        path = contained_path(self.root, file["path"])
        parts = ["batches", registered["batch_id"], registered["stage"]]
        if registered.get("schema_version", 1) < 2:
            parts.append(registered["version"])
        if path.parent != contained_path(self.root, *parts):
            raise ValueError("Artifact file is outside its snapshot directory")
        if not path.is_file():
            raise FileNotFoundError("Artifact file is missing")
        if path.stat().st_size != file["size"] or _sha256(path) != file["sha256"]:
            raise ValueError("Artifact file checksum does not match the frozen snapshot")
        return path
