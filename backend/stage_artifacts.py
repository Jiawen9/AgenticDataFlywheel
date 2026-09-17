"""Shared adapters between trajectory workbooks and immutable stage versions.

New workbook sidecars are the authoritative structured values. Excel is an
export and is never implicitly imported after a structured version exists.
"""
from __future__ import annotations

import hashlib
import json
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .data_store import ArtifactStore, DATA_ROOT, RecordStore, rebase_data_path


def store_root(output: Path, data_root: Path | None = None) -> Path:
    if data_root is not None:
        return Path(data_root).expanduser().resolve()
    output = Path(output).resolve()
    if output.is_relative_to(DATA_ROOT.resolve()):
        return DATA_ROOT
    # Explicit test / CLI output directories stay isolated from production data.
    return (output.parent if output.suffix.lower() in {".xlsx", ".json"} else output) / "data"


def fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def workbook_payload(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets: dict[str, list[dict[str, Any]]] = {}
        columns: dict[str, list[str]] = {}
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "") for value in next(rows, ())]
            columns[sheet.title] = headers
            sheets[sheet.title] = [dict(zip(headers, row)) for row in rows]
        return {"schema_version": 1, "columns": columns, "sheets": sheets}
    finally:
        workbook.close()


def sidecar_path(path: Path) -> Path:
    return Path(path).with_suffix(".json")


def payload_workbook(payload: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in payload["sheets"].items():
        sheet = workbook.create_sheet(name)
        columns = payload.get("columns", {}).get(name) or list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(key) for key in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
    return workbook


def write_payload_workbook(path: Path, payload: dict[str, Any]) -> None:
    workbook = payload_workbook(payload)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    finally:
        workbook.close()


def write_sidecar(path: Path, payload: dict[str, Any], *, source_ref: dict[str, Any] | None = None) -> None:
    target = sidecar_path(path)
    value = {**payload, "workbook_sha256": fingerprint(path)}
    if source_ref:
        value["source_ref"] = source_ref
    temporary = target.with_name(f".{target.name}.{uuid.uuid4().hex}.tmp")
    try:
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)


def read_workbook_payload(path: Path, *, allow_excel_import: bool = False, data_root: Path | None = None) -> dict[str, Any]:
    path = rebase_data_path(path, data_root) if data_root is not None else Path(path)
    if path.suffix.lower() == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(value, dict) or not isinstance(value.get("sheets"), dict):
            raise ValueError(f"invalid trajectory snapshot: {path}")
        if value.get("source_ref", {}).get("stage") == "02_annotation":
            workbook_path = path.with_suffix(".xlsx")
            key = hashlib.sha256(str(workbook_path.resolve()).encode("utf-8")).hexdigest()
            stored = RecordStore(store_root(workbook_path, data_root)).get("trajectory_annotations", key)
            if stored is not None:
                return {**stored["payload"], "source_ref": stored["artifact"]}
        return value
    candidate = sidecar_path(path)
    if candidate.is_file():
        # Do not hash or parse the human-readable export on internal reads.
        # A deleted/edited Excel must not change its committed JSON version.
        return read_workbook_payload(candidate, data_root=data_root)
    if allow_excel_import:
        return workbook_payload(path)
    raise FileNotFoundError(f"Required JSON snapshot is missing: {candidate}")


def structured_input_exists(path: Path) -> bool:
    return sidecar_path(path).is_file()


def publish_workbook(path: Path, *, batch_id: str, stage: str, data_root: Path | None = None,
                     source_refs: list[dict[str, Any]] | None = None,
                     metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    payload = workbook_payload(path)
    root = store_root(path, data_root)
    artifact = ArtifactStore(root).publish(
        batch_id, stage, payload, workbooks={path.name: path},
        source_refs=source_refs or [], metadata=metadata or {},
    )
    if stage == "02_annotation":
        key = hashlib.sha256(str(Path(path).resolve()).encode("utf-8")).hexdigest()
        RecordStore(root).put("trajectory_annotations", key,
                             {"path": str(Path(path).resolve()), "payload": payload, "artifact": artifact})
    write_sidecar(path, payload, source_ref=artifact)
    return artifact


def observation_payload(grouped: dict[str, Any], confidence_threshold: float) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    trajectories, rows = [], []
    for task_id, items in grouped.items():
        for trajectory_id, steps in items:
            audit = [step.audit_dict(confidence_threshold) for step in steps]
            trajectories.append({"task_id": task_id, "trajectory_id": trajectory_id, "steps": audit})
            for step in audit:
                classification = step.get("classification") or {}
                rows.append({
                    "任务编号": task_id, "轨迹编号": trajectory_id, "步骤": step["step"],
                    "image": step["image"], "xml": step["xml"], "action": step["action_text"],
                    "summary": step["summary"], "actions_box": step["actions_box"],
                    "Observation": step["observation"], "中间态类别": classification.get("category"),
                    "是否中间态": classification.get("is_intermediate"), "置信度": classification.get("confidence"),
                    "判断原因": classification.get("reason"), "计入树": step["counted_in_tree"],
                    "决策来源": step["decision_source"], "中途终止": step["excluded_intermediate_terminate"],
                })
    return {"schema_version": 1, "trajectories": trajectories}, rows


def quality_tables(task_results: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    summaries, steps, rubrics = [], [], []
    for task in task_results:
        task_id = task["task_id"]
        for trajectory_id, evaluation in task.get("evaluations", {}).items():
            summaries.append({"任务编号": task_id, "轨迹编号": trajectory_id,
                              **{key: value for key, value in evaluation.items() if key != "step_evaluations"}})
            for step in evaluation.get("step_evaluations", []):
                steps.append({"任务编号": task_id, "轨迹编号": trajectory_id, **step})
        rubric = task.get("rubric", {})
        dimensions = rubric.get("dimensions", [])
        if dimensions:
            rubrics.extend({"任务编号": task_id, **item} for item in dimensions)
        else:
            rubrics.append({"任务编号": task_id, "评分标准": rubric})
    return {"轨迹汇总": summaries, "步骤评分": steps, "评分标准": rubrics}


def load_quality_objects(path: Path) -> tuple[dict[str, Any], list[Any]]:
    """Construct the existing AdaRubric models from the exact JSON workbook view."""
    from adarubric import TaskDescription, Trajectory, TrajectoryStep
    sheets = read_workbook_payload(path)["sheets"]
    if not {"Tasks", "Trajectories", "Steps"}.issubset(sheets):
        raise ValueError("quality snapshot requires Tasks, Trajectories and Steps")
    tasks = {}
    for row in sheets["Tasks"]:
        task = TaskDescription(task_id=str(row["task_id"]), instruction=str(row["instruction"]),
            domain=str(row["domain"]), complexity=str(row["complexity"]),
            context=json.loads(row["context_json"]), expected_tools=json.loads(row["expected_tools_json"]))
        tasks[task.task_id] = task
    by_trajectory: dict[str, list[Any]] = {}
    for row in sheets["Steps"]:
        by_trajectory.setdefault(str(row["trajectory_id"]), []).append(TrajectoryStep(
            step_id=int(row["step_id"]), action=str(row["action"]),
            action_input=json.loads(row["action_input_json"]), observation=str(row["observation"])))
    trajectories = []
    for row in sheets["Trajectories"]:
        if str(row["task_id"]) not in tasks:
            raise ValueError(f"unknown quality task: {row['task_id']}")
        trajectories.append(Trajectory(trajectory_id=str(row["trajectory_id"]), task_id=str(row["task_id"]),
            steps=sorted(by_trajectory.get(str(row["trajectory_id"]), []), key=lambda item: item.step_id),
            final_answer=row.get("final_answer"), metadata=json.loads(row["metadata_json"])))
    return tasks, trajectories
