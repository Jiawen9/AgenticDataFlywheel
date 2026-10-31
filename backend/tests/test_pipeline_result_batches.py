"""Production pipeline/result-batch wiring with fake collector transport only."""
from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest

from openpyxl import Workbook, load_workbook

from backend.batch_lifecycle import lifecycle
from backend.data_store import RecordStore
from backend.phone_factory import PhoneFactoryStore
from backend.pipeline_access import active_pipeline
from backend.pipelines import PipelineError, PipelineManager, PipelineRuntime
from backend.task_generation.collection_batches import COLLECTION_COLUMNS
from backend.tests.collector_protocol_support import current_capabilities


class PipelineResultBatchTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.records = RecordStore(self.root)
        self.calls, self.remote = [], {}
        self.factory = PhoneFactoryStore(self.root, run_client_fn=self.client)
        self.factory.initialize_settings({"phones": ["phone-a", "phone-b"], "apps": ["App A", "App B"],
            "phoneApps": [{"phone_id": "phone-a", "app": "App A"}, {"phone_id": "phone-b", "app": "App B"}],
            "vla": ["http://mock.invalid/vla"]})
        book = Workbook()
        book.active.append(COLLECTION_COLUMNS)
        for suffix in ("A", "B"):
            values = {"用例编号": "case-" + suffix, "涉及APP": "App " + suffix, "任务": "目标 " + suffix,
                      "页面关键元素": "元素 " + suffix, "任务结束标志": "结束 " + suffix,
                      "SOP": "步骤 " + suffix, "动态变量说明": "变量 " + suffix}
            book.active.append([values.get(column, "") for column in COLLECTION_COLUMNS])
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        added = self.factory.add_task({"filename": "source.xlsx", "description": "reusable source",
            "content_base64": base64.b64encode(stream.getvalue()).decode()})
        self.source = added["tasks"][0]["source_batch_id"]
        self.runtime = PipelineRuntime(self.root, preprocessing=None, tree=None, quality=None,
            cot=None, factory=self.factory, overview=None)
        self.runtime.configurations = lambda: {}
        self.manager = PipelineManager(self.runtime, self.root)

    def client(self, args):
        self.calls.append(args)
        response = {"ok": True}
        if args[0] == "capabilities":
            response.update(protocol_version=1, batch_results=True, agent_types={
                name: {"available": True, "collection": True, "capture_logprobs": True}
                for name in ("newagent", "pipelineagent")})
            current_capabilities(response)
        elif args[0] == "start-run":
            run_id = args[args.index("--collection-run-id") + 1]
            self.remote[run_id] = {"run_id": run_id, "collection_run_id": run_id,
                "run_mode": "generate", "status": "running", "errors": [],
                "agent_type": args[args.index("--agent-type") + 1] if "--agent-type" in args else "newagent"}
            response.update(self.remote[run_id])
        elif args[0] == "run":
            response.update(self.remote[args[1]])
        return {"ok": True, "output": json.dumps(response)}

    def request(self, request_id="first", **changes):
        return {"request_id": request_id, "batch_id": self.source, "name": "独立采集结果",
            "mode": "automatic", "start_mode": "collect", "threshold": 4,
            "collection_config": {"agent_type": "pipelineagent", "phone_id": "phone-a", "app": "App A",
                                  "vla": "http://mock.invalid/vla"}, **changes}

    def test_collect_reuses_published_source_and_freezes_selected_complete_rows(self):
        source = self.factory.collection_runs.batch(self.source, require_workbook=True)
        original = Path(source["workbook_path"]).read_bytes()
        self.records.put("batch_lifecycle", self.source, {"batch_id": self.source, "status": "published",
                         "release_id": "historical-release", "published_at": "historical"})
        pipeline = self.manager.create(self.request())
        self.assertNotEqual(pipeline["batch_id"], self.source)
        self.assertEqual(pipeline["task_source_batch_id"], self.source)
        self.assertEqual(pipeline["batch_role"], "collection_result")
        batch = self.factory.collection_runs.batch(pipeline["batch_id"], require_workbook=True)
        expected = [task for task in source["snapshot"]["tasks"] if task["app"] == "App A"]
        self.assertEqual(batch["snapshot"]["tasks"], expected)
        self.assertEqual(pipeline["task_ids"], [task["task_id"] for task in expected])
        book = load_workbook(batch["workbook_path"])
        try:
            rows = list(book.active.values)
        finally:
            book.close()
        self.assertEqual(list(rows[0]), COLLECTION_COLUMNS)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][12], "结束 A")
        self.assertEqual(rows[1][15:], ("步骤 A", "变量 A"))
        self.assertEqual(Path(source["workbook_path"]).read_bytes(), original)
        self.assertEqual(lifecycle(self.source, self.root)["status"], "published")
        self.assertIsNone(active_pipeline(self.source, self.root))
        self.assertEqual(active_pipeline(batch["batch_id"], self.root)["pipeline_id"], pipeline["pipeline_id"])
        self.assertFalse(self.factory.collection_runs.list_runs(batch["batch_id"]))

    def test_second_collect_uses_another_result_while_prior_pipeline_still_owns_its_result(self):
        first = self.manager.create(self.request())
        second = self.manager.create(self.request("second", collection_config={"agent_type": "newagent",
            "phone_id": "phone-b", "app": "App B", "vla": "http://mock.invalid/vla"}))
        self.assertNotEqual(first["batch_id"], second["batch_id"])
        self.assertNotEqual(first["task_ids"], second["task_ids"])
        self.assertEqual(first["task_source_batch_id"], second["task_source_batch_id"])
        self.assertEqual(len(self.records.list("collection_result_batches")), 2)
        for pipeline in (first, second):
            self.assertEqual(active_pipeline(pipeline["batch_id"], self.root)["pipeline_id"], pipeline["pipeline_id"])
        self.assertIsNone(active_pipeline(self.source, self.root))

    def test_active_source_owner_does_not_own_new_result_batch(self):
        source_flow = self.manager.create(self.request("historical-owner", start_mode="existing"))
        pipeline = self.manager.create(self.request())
        self.assertEqual(active_pipeline(self.source, self.root)["pipeline_id"], source_flow["pipeline_id"])
        self.assertEqual(active_pipeline(pipeline["batch_id"], self.root)["pipeline_id"], pipeline["pipeline_id"])
        self.assertNotEqual(pipeline["batch_id"], self.source)
        self.assertFalse(self.calls)

    def test_dispatch_and_restart_use_reserved_result_without_another_fork(self):
        pipeline = self.manager.create(self.request())
        self.manager.tick(pipeline["pipeline_id"])
        current = self.manager.get(pipeline["pipeline_id"])
        self.assertEqual(current["status"], "running", current.get("error"))
        self.assertEqual(len(current["collection_run_ids"]), 1)
        run = self.factory.collection_runs.get(current["collection_run_ids"][0])
        self.assertEqual(run["batch_id"], pipeline["batch_id"])
        self.assertEqual(run["task_source_batch_id"], self.source)
        self.assertEqual(run["request"]["metadata"]["agent_type"], "pipelineagent")
        start = [args for args in self.calls if args[0] == "start-run"]
        self.assertEqual(len(start), 1)
        self.assertEqual(start[0][start[0].index("--batch-id") + 1], pipeline["batch_id"])
        self.manager = PipelineManager(self.runtime, self.root)
        self.manager.tick(pipeline["pipeline_id"])
        replay = self.manager.create(self.request())
        self.assertEqual(replay["pipeline_id"], pipeline["pipeline_id"])
        self.assertEqual(len([args for args in self.calls if args[0] == "start-run"]), 1)
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)
        self.assertEqual(len(self.records.list("collection_runs")), 1)
        with self.assertRaises(PipelineError):
            self.manager.create(self.request(collection_config={"agent_type": "newagent"}))
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)

    def test_existing_result_continues_without_creating_another_result(self):
        pipeline = self.manager.create(self.request())
        self.manager.control(pipeline["pipeline_id"], "terminate", pipeline["storage_revision"])
        self.manager.tick(pipeline["pipeline_id"])
        existing = self.manager.create(self.request("continue", batch_id=pipeline["batch_id"], start_mode="existing"))
        self.assertEqual(existing["batch_id"], pipeline["batch_id"])
        self.assertEqual(existing["task_source_batch_id"], self.source)
        self.assertEqual(existing["batch_role"], "collection_result")
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)
        self.assertEqual(existing["collection_run_ids"], [])
        self.assertFalse(self.calls)

    def test_failed_pipeline_record_commit_reuses_prepared_result_on_retry(self):
        with patch.object(self.manager.records, "put_many", side_effect=OSError("interrupted record commit")):
            with self.assertRaises(OSError):
                self.manager.create(self.request())
        batch = self.records.list("collection_result_batches")[0]
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)
        self.assertEqual(self.records.list("pipelines"), [])
        source = self.factory.collection_runs.batch(self.source, require_workbook=True)
        source_path = Path(source["workbook_path"]).resolve()
        self.assertTrue(source_path.is_relative_to(self.root.resolve()))
        source_path.unlink()
        self.manager = PipelineManager(self.runtime, self.root)
        pipeline = self.manager.create(self.request())
        self.assertEqual(pipeline["batch_id"], batch["batch_id"])
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)
        with self.assertRaises(PipelineError):
            self.manager.create(self.request(name="changed request"))
        self.assertEqual(len(self.records.list("collection_result_batches")), 1)
        self.assertFalse(self.calls)

    def test_failed_result_guides_new_collection_but_keeps_historical_message(self):
        failed = {"collection_run_id": "failed-run", "status": "failed", "errors": []}
        pipeline = {"batch_id": "test-result", "start_mode": "collect",
                    "collection_run_ids": ["failed-run"], "task_ids": ["task"]}
        with patch.object(self.runtime, "collection_runs", return_value=[failed]):
            with self.assertRaisesRegex(PipelineError, "请从任务来源重新采集，并选择新结果批次继续"):
                self.runtime.input_ready({**pipeline, "batch_role": "collection_result"})
            with self.assertRaisesRegex(PipelineError, "请补采后接续已有结果"):
                self.runtime.input_ready(pipeline)


if __name__ == "__main__":
    unittest.main()
