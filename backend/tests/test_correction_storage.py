"""Temporary-data regression tests for SQLite drafts and immutable process artifacts."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

from backend.data_store import ArtifactStore, RecordStore, RevisionConflict
from backend.trajectory_correction import draft_store, service, cot_jobs
from backend.trajectory_correction.workbook import load_snapshot
from backend.data_publishing.service import DatasetReleaseRegistry
from backend.stage_artifacts import workbook_payload, write_sidecar


class DeferredExecutor:
    def submit(self, fn, *args):
        self.fn, self.args = fn, args
    def run(self):
        self.fn(*self.args)


class CorrectionStorageTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.source = self.root / "annotated.xlsx"
        book = Workbook()
        book.active.append(["文件夹名", "image", "action", "summary", "thought", "actions_box"])
        for step in (1, 2):
            book.active.append(["TASK-1", f"TASK/TASK-1/step{step:03d}_vla_input.jpg", '{"action":"wait"}', f"original {step}", "", "wait()"])
        book.save(self.source)
        book.close()
        self.snapshot = load_snapshot(self.source, allow_excel_import=True)
        self.sid = "a" * 16
        self.patches = [
            patch.object(draft_store, "CORRECTION_SESSIONS_DIR", self.root / "sessions"),
            patch.object(service, "_snapshot", side_effect=lambda _: deepcopy(self.snapshot)),
            patch.object(service, "_source_for_session", return_value=(self.source, self.root)),
            patch.object(service, "CORRECTION_EXPORTS_DIR", self.root / "exports"),
        ]
        for patcher in self.patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        draft_store.save_session({"session_id": self.sid, "tree_run_id": "batch-one", "row_edits": {},
                                  "group_exports": {}, "cot": {}, "exports": []})

    def test_concurrent_independent_fields_and_stale_save(self):
        stale = draft_store.load_session(self.sid)
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(lambda args: service.patch_row(self.sid, 2, args), [{"summary": "manual summary"}, {"thought": "manual thought"}]))
        current = draft_store.load_session(self.sid)
        self.assertEqual(current["row_edits"]["2"], {"summary": "manual summary", "thought": "manual thought"})
        stale["row_edits"] = {"2": {"summary": "stale"}}
        with self.assertRaises(RevisionConflict):
            draft_store.save_session(stale)
        # SQLite is authoritative even if a compatibility mirror is changed.
        (self.root / "sessions").mkdir(exist_ok=True)
        (self.root / "sessions" / f"{self.sid}.json").write_text(json.dumps(stale), encoding="utf-8")
        self.assertEqual(draft_store.load_session(self.sid)["row_edits"], current["row_edits"])
        self.assertTrue((self.root / "system" / "app.sqlite").is_file())

    def test_legacy_session_json_is_invisible_and_untouched(self):
        legacy_id = "b" * 16
        directory = self.root / "sessions"
        directory.mkdir(exist_ok=True)
        path = directory / f"{legacy_id}.json"
        path.write_text(json.dumps({"session_id": legacy_id, "row_edits": {}}), encoding="utf-8")
        content = path.read_bytes()
        self.assertIsNone(draft_store.load_session(legacy_id))
        self.assertEqual({s["session_id"] for s in draft_store.list_sessions()}, {self.sid})
        self.assertEqual(path.read_bytes(), content)
        self.assertIsNone(RecordStore(self.root).get("correction_sessions", legacy_id))

    def test_internal_reader_requires_json_and_legacy_assets_are_not_searched(self):
        from backend.trajectory_correction.assets import resolve_asset
        with self.assertRaisesRegex(FileNotFoundError, "JSON"):
            load_snapshot(self.source)
        self.patches[1].stop()
        with self.assertRaisesRegex(ValueError, "JSON"):
            service._snapshot(draft_store.load_session(self.sid))
        new_assets = self.root / "raw"
        unrelated = new_assets / "other" / "step.jpg"
        unrelated.parent.mkdir(parents=True)
        unrelated.write_bytes(b"wrong image")
        old_assets = self.root / "backend_workspace" / "rollout_trajectories" / "TASK"
        old_assets.mkdir(parents=True)
        old = old_assets / "step.jpg"
        old.write_bytes(b"legacy image")
        with self.assertRaises(FileNotFoundError):
            resolve_asset(new_assets, "TASK/step.jpg")
        self.assertEqual(old.read_bytes(), b"legacy image")

    def test_cot_job_json_is_not_loaded_or_rewritten_on_restart(self):
        jobs = self.root / "jobs"
        jobs.mkdir()
        job_id = "c" * 32
        path = jobs / f"{job_id}.json"
        path.write_text(json.dumps({"job_id": job_id, "status": "running"}), encoding="utf-8")
        before = path.read_bytes()
        manager = cot_jobs.CotJobManager(jobs, executor=DeferredExecutor())
        self.assertIsNone(manager.get(job_id))
        self.assertEqual(manager.list_jobs(), [])
        self.assertEqual(path.read_bytes(), before)

    def test_process_snapshot_keeps_one_current_result_and_all_audit_rows(self):
        service.patch_row(self.sid, 2, {"summary": "=literal", "thought": "final thought", "deleted": True})
        session = draft_store.load_session(self.sid)
        first = service.publish_stage_snapshot(session, self.snapshot, "06_correction")
        store = ArtifactStore(self.root)
        before = store.resolve_file(first, "result.json").read_bytes()
        book = load_workbook(store.resolve_file(first, "result.xlsx"))
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.cell(2, headers.index("summary") + 1).value, "=literal")
        self.assertEqual(sheet.cell(2, headers.index("summary") + 1).data_type, "s")
        self.assertTrue(sheet.cell(2, headers.index("deleted") + 1).value)
        book.close()
        service.patch_row(self.sid, 2, {"summary": "new"})
        second = service.publish_stage_snapshot(draft_store.load_session(self.sid), self.snapshot, "06_correction")
        self.assertNotEqual(first["version"], second["version"])
        with self.assertRaises(FileNotFoundError):
            store.resolve_file(first, "result.json")
        self.assertEqual(len(store.list(batch_id="batch-one", stage="06_correction")), 1)
        self.assertNotEqual(store.resolve_file(second, "result.json").read_bytes(), before)

    def run_cot(self, edit_during_generate):
        service.patch_row(self.sid, 2, {"actions": '{"action":"answer","text":"done"}', "summary": "before submit"})
        executor = DeferredExecutor()
        owner = self
        class Generator:
            model = "fake"
            def generate(self, **kwargs):
                edit_during_generate()
                return {"thought": "generated thought", "summary": "generated summary"}
        manager = cot_jobs.CotJobManager(self.root / "jobs", generator_factory=Generator, executor=executor)
        with patch.object(cot_jobs, "_snapshot", return_value=deepcopy(self.snapshot)), patch.object(cot_jobs, "session_asset", return_value=self.root / "unused.jpg"), patch.object(cot_jobs, "publish_cot_snapshot", return_value={"version": "test"}):
            job = manager.submit(self.sid)
            executor.run()
        return manager.get(job["job_id"])

    def test_cot_commits_against_latest_session_and_preserves_new_manual_values(self):
        def edit():
            service.patch_row(self.sid, 2, {"summary": "edited during generation"})
            service.patch_row(self.sid, 3, {"thought": "another row"})
        job = self.run_cot(edit)
        self.assertEqual(job["status"], "succeeded")
        current = draft_store.load_session(self.sid)
        self.assertEqual(current["row_edits"]["2"]["summary"], "edited during generation")
        self.assertEqual(current["row_edits"]["3"]["thought"], "another row")
        self.assertEqual(current["cot"]["2"]["thought"], "generated thought")

    def test_cot_rejects_changed_action_without_restoring_old_session(self):
        job = self.run_cot(lambda: service.patch_row(self.sid, 2, {"actions": '{"action":"wait"}'}))
        self.assertEqual(job["status"], "failed")
        current = draft_store.load_session(self.sid)
        self.assertNotIn("actions", current["row_edits"]["2"])
        self.assertNotIn("2", current["cot"])

    def test_new_session_reads_json_after_source_changes(self):
        # Use actual snapshot/source functions only in this test.
        self.patches[1].stop()
        self.patches[2].stop()
        write_sidecar(self.source, workbook_payload(self.source))
        self.source.unlink()
        selection = {"tree_run_id": "frozen-batch", "storage_batch_id": "upstream-batch", "tasks": [{"task_id": "TASK", "trajectory_id": "TASK-1"}]}
        assets = self.root / "raw"
        assets.mkdir()
        with patch.object(service, "CORRECTION_INPUTS_DIR", self.root / "inputs"), patch.object(service, "ensure_correction_dirs"), patch.object(service, "top1_selection_for_run", return_value=selection), patch("backend.trajectory_correction.quality_selection.source_for_tree_run", return_value=(self.source, assets)), patch("backend.trajectory_correction.assets.trajectory_source", return_value=assets):
            created = service.create_session("frozen-batch")
            saved = draft_store.load_session(created["session_id"])
            self.assertEqual(saved["storage_batch_id"], "upstream-batch")
            before = service._snapshot(saved)
            artifact = service.publish_stage_snapshot(saved, before, "06_correction")
            self.assertEqual(artifact["batch_id"], "upstream-batch")
            self.source.write_bytes(b"changed upstream")
            self.assertEqual(service._snapshot(saved), before)
            frozen, _ = service._source_for_session(saved)
            self.assertNotEqual(frozen.read_bytes(), b"changed upstream")
            frozen.unlink()
            self.assertEqual(service.get_session(saved["session_id"])["row_count"], 2)
            group = service.get_group(saved["session_id"], before["groups"][0]["group_id"])
            image = assets / group["rows"][0]["image"]
            image.parent.mkdir(parents=True, exist_ok=True)
            image.write_bytes(b"fake image")
            self.assertEqual(service.session_asset(saved["session_id"], group["rows"][0]["image"]), image)
            exported = service.export_dataset_session(saved["session_id"])
            output = self.root / "exports" / saved["session_id"] / exported["filename"]
            book = load_workbook(output)
            self.assertEqual(book.active.cell(2, 4).value, "original 1")
            book.close()
            # SQLite owns the current draft; export files may be regenerated.
            frozen.write_bytes(b"must not import this workbook")
            frozen.with_suffix(".json").unlink()
            service.export_dataset_session(saved["session_id"])
            (self.root / "inputs" / saved["session_id"] / saved["source_snapshot"]["json"]).unlink()
            self.assertEqual(len(service.get_group(saved["session_id"], before["groups"][0]["group_id"])["rows"]), 2)

    def test_upstream_json_ignores_excel_changes_and_json_alone_validates_quality(self):
        from backend.trajectory_correction.quality_selection import _verify_source_version
        payload = workbook_payload(self.source)
        write_sidecar(self.source, payload)
        source_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        tree_root = self.root / "trees"
        run = tree_root / "run"
        run.mkdir(parents=True)
        manifest = {"source_xlsx": {"sha256": source_hash}, "source_json": {"sha256": hashlib.sha256(self.source.with_suffix(".json").read_bytes()).hexdigest()}}
        (run / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        self.source.write_bytes(b"invalid Excel edit")
        self.assertEqual(load_snapshot(self.source)["groups"][0]["rows"][0]["summary"], "original 1")
        self.source.unlink()
        self.assertEqual(_verify_source_version("run", self.source, tree_root=tree_root), source_hash)

    def test_legacy_empty_row_keeps_original_row_numbers_after_json_freeze(self):
        book = load_workbook(self.source)
        book.active.insert_rows(3)
        book.save(self.source)
        book.close()
        self.snapshot = load_snapshot(self.source, allow_excel_import=True)
        self.assertEqual([row["excel_row"] for row in self.snapshot["groups"][0]["rows"]], [2, 4])
        write_sidecar(self.source, workbook_payload(self.source))
        self.source.unlink()
        service.patch_row(self.sid, 4, {"summary": "correct row four"})
        exported = service.export_dataset_session(self.sid)
        book = load_workbook(self.root / "exports" / self.sid / exported["filename"])
        self.assertEqual(book.active.cell(2, 4).value, "original 1")
        self.assertIsNone(book.active.cell(3, 4).value)
        self.assertEqual(book.active.cell(4, 4).value, "correct row four")
        book.close()

    def test_release_and_published_session_commit_together(self):
        from backend.batch_lifecycle import BatchPublishedError, lifecycle
        from backend.batch_results import annotation_task_fingerprints, tree_task_fingerprints, tree_result_hashes
        from backend.trajectory_correction.session_state import identified
        payload = workbook_payload(self.source)
        write_sidecar(self.source, payload)
        snapshot = identified(self.snapshot)
        tasks = {group["task_id"]: group for group in snapshot["groups"]}
        store = ArtifactStore(self.root)
        conversion = store.publish("batch-one", "01_conversion", payload)
        annotation = store.publish("batch-one", "02_annotation", payload, source_refs=[conversion])
        observation = store.publish("batch-one", "03_observation", {"trajectories": [
            {"task_id": task, "trajectory_id": group["meta_task"], "steps": group["rows"]}
            for task, group in tasks.items()]}, source_refs=[annotation])
        quality_input = {"sheets": {"Tasks": [{"task_id": task} for task in tasks],
            "Trajectories": [{"task_id": task, "trajectory_id": group["meta_task"]} for task, group in tasks.items()],
            "Steps": [{"trajectory_id": group["meta_task"], "step_id": row["step"]}
                      for group in tasks.values() for row in group["rows"]]}}
        trees = {task: {"task_id": task, "source_trajectories": [group["meta_task"]]} for task, group in tasks.items()}
        task_fingerprints = tree_task_fingerprints(payload, {"model": "offline"})
        tree_hashes = tree_result_hashes(trees, quality_input, task_fingerprints)
        tree = store.publish("batch-one", "04_tree", {
            "batch_id": "batch-one", "run_id": "batch-one", "raw_root": str(self.root),
            "trees": trees, "source_annotation": payload, "quality_input": quality_input,
            "tasks": [{"task_id": task, "trajectory_count": 1} for task in tasks],
            "source_task_fingerprints": annotation_task_fingerprints(payload),
            "task_fingerprints": task_fingerprints, "tree_hashes": tree_hashes}, source_refs=[observation])
        store.publish("batch-one", "05_quality", {"run_id": "batch-one", "source_tree_hashes": tree_hashes,
            "tasks": [{"task_id": task, "trajectory_count": 1, "evaluations": {
                group["meta_task"]: {"global_score": 5, "passed_threshold": True}}} for task, group in tasks.items()]},
            source_refs=[tree])
        draft_store.update_session(self.sid, lambda current: current.update(task_fingerprints=task_fingerprints))
        service.export_dataset_session(self.sid)
        registry = DatasetReleaseRegistry(releases_file=self.root / "release-index.json", project_root=self.root,
                                          data_root=self.root, trajectory_root=self.root,
                                          correction_exports_dir=self.root / "exports")
        with patch.object(registry._records, "put_many", side_effect=RevisionConflict("concurrent edit")):
            with self.assertRaises(RevisionConflict):
                registry.create("try", [self.sid])
        self.assertFalse(draft_store.load_session(self.sid).get("published"))
        self.assertEqual(lifecycle("batch-one", self.root)["status"], "active")
        self.assertEqual(registry.list_releases(), [])
        release = registry.create("success", [self.sid])
        self.assertEqual(draft_store.load_session(self.sid)["published_release_id"], release["release_id"])
        self.assertEqual(RecordStore(self.root).get("dataset_releases", release["release_id"])["source_refs"][0]["id"], self.sid)
        before = draft_store.load_session(self.sid)
        with self.assertRaises(BatchPublishedError):
            draft_store.update_session(self.sid, lambda current: current.update(row_edits={}))
        self.assertEqual(draft_store.load_session(self.sid), before)
        self.assertTrue(before["published"])
        with self.assertRaises(BatchPublishedError):
            registry.create("duplicate", [self.sid])
