from __future__ import annotations

from copy import deepcopy
from datetime import datetime
import hashlib
import json
import re
from pathlib import Path
import shutil
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

from backend.training_data_overview.external_workbook import parse_external_workbook
from backend.training_data_overview.converter import ConversionError, convert_release


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class _WorkbookFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.headers = ["trajectory_id", "image", "actions", "sop", "APP", "一级场景", "二级场景", "summary"]
        self.steps = [
            ["a", "missing/a/step1.jpg", '{"action":"CLICK"}', "打开", "表内 App", "生活", "查询", "第一步"],
            ["a", "missing/a/step2.jpg", '[{"action":"finish"}]', "结束", "", "", "", "第二步"],
            ["b", "missing/b/step1.jpg", "malformed", "另一个轨迹", "", "", "", ""],
        ]

    def workbook(self, *, rows=None, headers=None, name="input.xlsx", auxiliary=False, style=False, title="步骤"):
        path = self.root / name
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = title
        worksheet.append(self.headers if headers is None else headers)
        for row in self.steps if rows is None else rows:
            worksheet.append(row)
        if style:
            worksheet["A1"].font = Font(bold=True, color="FF0000")
            worksheet.column_dimensions["A"].width = 50
        if auxiliary:
            other = workbook.create_sheet("辅助")
            other.append(["trajectory_id", "action"])
            other.append(["other", '{"action":"wait"}'])
        workbook.save(path)
        workbook.close()
        return path

    def parse(self, path=None, **kwargs):
        return parse_external_workbook(path or self.workbook(), **{
            "data_source": "人工采集", "data_date": "2026-08-20", **kwargs})


class ExternalWorkbookTests(_WorkbookFixture):
    def test_grouped_table_classifications_override_fallback_and_manual_stays_unknown(self):
        result = self.parse(app="上传 App", level1="上传场景", level2="上传能力")
        self.assertTrue(result["valid"], result["errors"])
        self.assertEqual(result["sheet_name"], "步骤")
        self.assertEqual(result["sheets"], ["步骤"])
        self.assertEqual(result["summary"], {"trajectory_count": 2, "step_count": 3,
            "apps": [{"name": "表内 App", "trajectory_count": 1, "step_count": 2},
                     {"name": "上传 App", "trajectory_count": 1, "step_count": 1}],
            "scenes": [{"level1": "生活", "level2": "查询", "trajectory_count": 1, "step_count": 2},
                       {"level1": "上传场景", "level2": "上传能力", "trajectory_count": 1, "step_count": 1}]})
        first, second = result["rows"]
        self.assertEqual(first["action_box"], {"click": 1, "finish": 1})
        self.assertEqual(second["action_box"], {"unknown": 1})
        for row in result["rows"]:
            self.assertIsNone(row["人工精修步骤数量"])
            self.assertFalse(row["manual_known"])
            self.assertEqual(row["时间"], "2026-08-20")
            self.assertEqual(row["生产方式"], "人工采集")
            self.assertNotIn("release_id", row)
            self.assertNotIn("轨迹", row)
        self.assertTrue(any(warning["row"] == 4 and warning["field"] == "actions"
                            for warning in result["warnings"]))

    def test_missing_classifications_use_existing_unknown_labels(self):
        row = self.parse(data_source="")["rows"][1]
        self.assertEqual([row["APP"], row["一级场景"], row["二级场景"]], ["未记录 App", "未分类", "未分类"])
        self.assertFalse(row["app_known"])
        self.assertFalse(row["level1_known"])
        self.assertEqual(row["生产方式"], "人工采集")

    def test_active_sheet_default_and_explicit_selection(self):
        path = self.workbook(auxiliary=True)
        result = self.parse(path)
        self.assertTrue(result["valid"])
        self.assertEqual(result["summary"]["step_count"], 3)
        selected = self.parse(path, sheet_name="辅助")
        self.assertTrue(selected["valid"])
        self.assertEqual(selected["metadata"]["sheet_name"], "辅助")
        self.assertEqual(selected["summary"]["step_count"], 1)
        missing = self.parse(path, sheet_name="不存在")
        self.assertFalse(missing["valid"])
        self.assertEqual(missing["sheets"], ["步骤", "辅助"])

    def test_conflicts_are_reported_with_sheet_row_and_field(self):
        steps = deepcopy(self.steps)
        steps[1][4] = "冲突 App"
        result = self.parse(self.workbook(rows=steps))
        self.assertFalse(result["valid"])
        issue = next(issue for issue in result["errors"] if issue["field"] == "APP")
        self.assertEqual((issue["sheet"], issue["row"]), ("步骤", 3))
        self.assertIn("第 2 行", issue["message"])
        self.assertEqual(result["rows"], [])
        self.assertEqual(result["content_hash"], "")

    def test_metadata_date_errors_and_invalid_file_are_structured(self):
        path = self.workbook()
        for value in ("", "2026-02-30", "2026-8-20", "20260820", "2026-08-20T00:00:00"):
            with self.subTest(value=value):
                result = self.parse(path, data_date=value)
                self.assertFalse(result["valid"])
                self.assertTrue(any(issue["field"] == "data_date" for issue in result["errors"]))
        unsupported = self.root / "input.csv"
        unsupported.write_text("trajectory_id,action", encoding="utf-8")
        self.assertFalse(self.parse(unsupported)["valid"])
        invalid = self.root / "invalid.xlsx"
        invalid.write_text("not a workbook", encoding="utf-8")
        self.assertFalse(self.parse(invalid)["valid"])
        self.assertFalse(self.parse(self.root / "missing.xlsx")["valid"])

    def test_missing_headers_duplicates_empty_sheet_and_missing_identity_rejected(self):
        cases = [([], []), (["trajectory_id"], [["a"]]), (["action"], [['{"action":"click"}']]),
                 (["trajectory_id", "action", "action"], [["a", "x", "y"]]),
                 (["trajectory_id", "action"], [["", "x"]]),
                 (["trajectory_id", "action", ""], [["a", "x", "orphan"]])]
        for headers, rows in cases:
            with self.subTest(headers=headers, rows=rows):
                result = self.parse(self.workbook(headers=headers, rows=rows))
                self.assertFalse(result["valid"])
                self.assertTrue(result["errors"])
                self.assertTrue(all(set(issue) == {"sheet", "row", "field", "message"} for issue in result["errors"]))

    def test_identity_fallback_uses_image_parent_only_as_text(self):
        result = self.parse(self.workbook(headers=["文件夹名", "image", "action"], rows=[
            ["same", r"nonexistent\one\step1.jpg", '{"action":"wait"}'],
            ["same", "nonexistent/two/step1.jpg", '{"action":"wait"}'],
            ["same", "nonexistent/one/step2.jpg", '{"action":"finish"}']]))
        self.assertTrue(result["valid"])
        self.assertEqual([row["trajectory_identity"] for row in result["rows"]],
                         ["nonexistent/one/same", "nonexistent/two/same"])
        self.assertEqual([row["step数量"] for row in result["rows"]], [2, 1])

    def test_full_cell_content_and_order_change_fingerprint_even_with_same_counts(self):
        baseline = self.parse()
        variants = []
        changed_sop = deepcopy(self.steps)
        changed_sop[0][3] = "修改 SOP"
        variants.append(changed_sop)
        changed_summary = deepcopy(self.steps)
        changed_summary[0][-1] = "非统计列也修改"
        variants.append(changed_summary)
        changed_coordinate = deepcopy(self.steps)
        changed_coordinate[0][2] = '{"action":"CLICK","coordinate":[1,2]}'
        variants.append(changed_coordinate)
        variants.append([self.steps[1], self.steps[0], self.steps[2]])
        for rows in variants:
            with self.subTest(rows=rows):
                result = self.parse(self.workbook(rows=rows))
                self.assertEqual(result["summary"], baseline["summary"])
                self.assertNotEqual(result["content_hash"], baseline["content_hash"])
        for kwargs in ({"data_date": "2026-08-21"}, {"data_source": "人工复采"}, {"app": "补充 App"},
                       {"level1": "补充场景"}, {"level2": "补充能力"}):
            self.assertNotEqual(self.parse(**kwargs)["content_hash"], baseline["content_hash"])

    def test_filename_sheet_name_and_styles_do_not_change_fingerprint(self):
        original = self.parse()
        resaved = self.parse(self.workbook(name="另存.xlsx", title="重命名工作表", style=True))
        self.assertEqual(original["content_hash"], resaved["content_hash"])
        self.assertNotEqual(original["sheet_name"], resaved["sheet_name"])

    def test_typed_extra_cells_blank_rows_and_xlsm_are_supported(self):
        path = self.workbook(name="input.xlsm", headers=["trajectory_id", "action", "timestamp", "note"], rows=[
            [0, '{"action":"click"}', datetime(2026, 8, 20, 10, 30), 12.5], [],
            [0, '{"action":"finish"}', None, False]])
        result = self.parse(path)
        self.assertTrue(result["valid"], result["errors"])
        self.assertEqual(result["rows"][0]["trajectory_identity"], "0")
        self.assertEqual(result["rows"][0]["step数量"], 2)
        self.assertEqual(len(result["content_hash"]), 64)


    def test_zip_expansion_limit_is_checked_before_loading_workbook(self):
        path = self.workbook()
        with ZipFile(path) as archive:
            expanded_size = sum(entry.file_size for entry in archive.infolist())
        with patch("backend.training_data_overview.external_workbook.MAX_UNCOMPRESSED_BYTES", expanded_size - 1), \
             patch("backend.training_data_overview.external_workbook.load_workbook", side_effect=AssertionError("loaded too early")):
            result = self.parse(path)
        self.assertFalse(result["valid"])
        self.assertIn("512 MiB", result["errors"][0]["message"])
        with patch("backend.training_data_overview.external_workbook.MAX_UNCOMPRESSED_BYTES", expanded_size):
            self.assertTrue(self.parse(path)["valid"])

    def test_false_dimensions_neither_truncate_data_nor_iterate_empty_rectangle(self):
        original = self.workbook()
        baseline = self.parse(original)
        for dimension in ("A1:A1", "A1:XFD1048576"):
            with self.subTest(dimension=dimension):
                destination = self.root / ("small.xlsx" if dimension == "A1:A1" else "large.xlsx")
                with ZipFile(original) as source, ZipFile(destination, "w", ZIP_DEFLATED) as target:
                    for info in source.infolist():
                        data = source.read(info.filename)
                        if info.filename == "xl/worksheets/sheet1.xml":
                            data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="' + dimension.encode() + b'"', data)
                        target.writestr(info, data)
                result = self.parse(destination)
                self.assertTrue(result["valid"], result["errors"])
                self.assertEqual(result["summary"], baseline["summary"])
                self.assertEqual(result["content_hash"], baseline["content_hash"])

    def test_streamed_hash_preserves_v1_json_identity(self):
        result = self.parse()
        fingerprint = {"parser_version": 1,
                       "metadata": {key: value for key, value in result["metadata"].items() if key != "sheet_name"},
                       "headers": self.headers,
                       # Excel represents empty strings as empty cells on load.
                       "steps": [[None if value == "" else value for value in row] for row in self.steps]}
        expected = hashlib.sha256(json.dumps(fingerprint, ensure_ascii=False, sort_keys=True,
                                  separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()
        self.assertEqual(result["content_hash"], expected)

    def test_chart_sheet_selection_returns_validation_error_and_data_sheet_remains_selectable(self):
        path = self.root / "with-chart.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "步骤"
        sheet.append(["trajectory_id", "action", "number"])
        sheet.append(["a", '{"action":"click"}', 1])
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=2), titles_from_data=True)
        chartsheet = workbook.create_chartsheet("图表")
        chartsheet.add_chart(chart)
        workbook.active = chartsheet
        workbook.save(path)
        workbook.close()
        result = self.parse(path)
        self.assertFalse(result["valid"])
        self.assertEqual(result["errors"][0]["field"], "sheet_name")
        self.assertEqual(result["sheets"], ["步骤", "图表"])
        self.assertTrue(self.parse(path, sheet_name="步骤")["valid"])


class ExternalReleaseConverterTests(_WorkbookFixture):
    def setUp(self):
        super().setUp()
        self.registry = SimpleNamespace(data_root=self.root, resolve_project_path=lambda value: self.root / value)

    def reference(self, path):
        return {"path": path.relative_to(self.root).as_posix(), "sha256": sha(path)}

    def write_json(self, path, value):
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return self.reference(path)

    def release(self, release_id="rel_external"):
        path = self.workbook()
        parsed = self.parse(path)
        destination = self.root / "releases" / release_id
        (destination / "001").mkdir(parents=True)
        original = destination / "001" / "steps.xlsx"
        shutil.copyfile(path, original)
        original_ref = self.reference(original)
        canonical = {"schema_version": 1, "parser_version": 1,
                     **{key: parsed[key] for key in ("metadata", "rows", "warnings", "content_hash")}}
        canonical_ref = self.write_json(destination / "external-data.json", canonical)
        manifest = {"schema_version": 1, "parser_version": 1, "release_id": release_id,
                    "batch_id": "external_batch", "import_id": "import_one", "filename": "steps.xlsx",
                    "sheet_name": parsed["sheet_name"], "metadata": parsed["metadata"],
                    "source": original_ref, "canonical": canonical_ref, "content_hash": parsed["content_hash"],
                    "created_at": "2026-09-20T02:00:00Z"}
        manifest_ref = self.write_json(destination / "import-manifest.json", manifest)
        return {"release_id": release_id, "source_kind": "external_manual", "batch_ids": ["external_batch"],
                "created_at": "2026-09-20T02:00:00Z", "excel_paths": [original_ref], "trajectory_paths": [],
                "external_import": {**parsed["metadata"], "parser_version": 1, "import_id": "import_one",
                    "source_sha256": original_ref["sha256"], "content_hash": parsed["content_hash"],
                    "canonical": canonical_ref, "manifest": manifest_ref}}

    def update_json(self, release, field, transform):
        ref = release["external_import"][field]
        path = self.root / ref["path"]
        value = json.loads(path.read_text(encoding="utf-8"))
        transform(value)
        release["external_import"][field] = self.write_json(path, value)
        if field == "canonical":
            manifest_ref = release["external_import"]["manifest"]
            manifest_path = self.root / manifest_ref["path"]
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            manifest["canonical"] = release["external_import"]["canonical"]
            release["external_import"]["manifest"] = self.write_json(manifest_path, manifest)

    def test_external_conversion_uses_frozen_canonical_without_sessions_or_workbook_parse(self):
        release = self.release()
        def hashes():
            return {path.relative_to(self.root).as_posix(): sha(path) for path in self.root.rglob("*") if path.is_file()}
        before = hashes()
        with patch("backend.training_data_overview.converter._Sources", side_effect=AssertionError("session read")), \
             patch("backend.training_data_overview.converter.FrozenReleaseSources", side_effect=AssertionError("stage read")), \
             patch("backend.training_data_overview.converter.load_workbook", side_effect=AssertionError("workbook parse")):
            result = convert_release(self.registry, release)
        self.assertEqual(len(result["rows"]), 2)
        self.assertEqual(result["rows"][0]["轨迹"], "rel_external/001/a")
        self.assertEqual(result["rows"][0]["子任务轨迹"], "rel_external/001/a/subtask-1")
        self.assertEqual(result["rows"][0]["时间"], "2026-08-20")
        self.assertEqual(result["rows"][0]["source_file_index"], 1)
        self.assertTrue(all(isinstance(value, str) for value in result["warnings"]))
        self.assertEqual(hashes(), before)

    def test_new_release_uses_distinct_trajectory_identity(self):
        first = convert_release(self.registry, self.release("rel_first"))
        second = convert_release(self.registry, self.release("rel_second"))
        self.assertNotEqual(first["rows"][0]["轨迹"], second["rows"][0]["轨迹"])
        self.assertEqual(first["rows"][0]["trajectory_identity"], second["rows"][0]["trajectory_identity"])

    def test_all_frozen_files_require_matching_hashes(self):
        release = self.release()
        references = [release["excel_paths"][0], release["external_import"]["canonical"], release["external_import"]["manifest"]]
        for ref in references:
            with self.subTest(path=ref["path"]):
                path = self.root / ref["path"]
                before = path.read_bytes()
                try:
                    path.write_bytes(before + b" ")
                    with self.assertRaisesRegex(ConversionError, "SHA256"):
                        convert_release(self.registry, release)
                finally:
                    path.write_bytes(before)

    def test_all_frozen_paths_must_stay_in_the_release(self):
        release = self.release()
        for field in ("source", "canonical", "manifest"):
            with self.subTest(field=field):
                changed = deepcopy(release)
                ref = changed["excel_paths"][0] if field == "source" else changed["external_import"][field]
                path = self.root / ref["path"]
                outside = self.root / path.name
                shutil.copyfile(path, outside)
                ref["path"] = outside.name
                with self.assertRaisesRegex(ConversionError, "冻结文件"):
                    convert_release(self.registry, changed)

    def test_cross_release_and_registration_mismatch_rejected(self):
        for field, value in (("release_id", "other"), ("batch_id", "other"), ("import_id", "other"),
                             ("sheet_name", "other"), ("content_hash", "wrong"), ("parser_version", 2)):
            with self.subTest(field=field):
                release = self.release("rel_" + field)
                self.update_json(release, "manifest", lambda manifest: manifest.update({field: value}))
                with self.assertRaises(ConversionError):
                    convert_release(self.registry, release)
        release = self.release("rel_metadata")
        release["external_import"]["data_date"] = "2026-08-21"
        with self.assertRaisesRegex(ConversionError, "元数据"):
            convert_release(self.registry, release)

    def test_invalid_canonical_rows_rejected_even_with_registered_matching_hash(self):
        changes = [lambda row: row.update({"step数量": 200}), lambda row: row.update({"manual_known": True}),
                   lambda row: row.update({"人工精修步骤数量": 1}), lambda row: row.update({"时间": "2026-08-21"}),
                   lambda row: row.update({"APP": ""}), lambda row: row.update({"step数量": True})]
        for index, change in enumerate(changes):
            with self.subTest(index=index):
                release = self.release(f"rel_invalid_{index}")
                self.update_json(release, "canonical", lambda value: change(value["rows"][0]))
                with self.assertRaises(ConversionError):
                    convert_release(self.registry, release)
        release = self.release("rel_duplicate")
        self.update_json(release, "canonical", lambda value: value["rows"].append(value["rows"][0]))
        with self.assertRaisesRegex(ConversionError, "重复"):
            convert_release(self.registry, release)

    def test_missing_original_or_missing_import_registration_rejected(self):
        release = self.release()
        path = self.root / release["excel_paths"][0]["path"]
        path.unlink()
        with self.assertRaisesRegex(ConversionError, "不存在"):
            convert_release(self.registry, release)
        with self.assertRaises(ConversionError):
            convert_release(self.registry, {"release_id": "rel_x", "source_kind": "external_manual"})


if __name__ == "__main__":
    unittest.main()
