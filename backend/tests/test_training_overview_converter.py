from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from backend.data_store import ArtifactStore
from backend.data_store.release_provenance import freeze_existing_release, FrozenReleaseSources
from backend.training_data_overview.converter import ConversionError, convert_release


class ConverterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = ArtifactStore(self.root)
        self.registry = SimpleNamespace(data_root=self.root, resolve_project_path=lambda value: self.root / value)
        self.batch = "batch-a"
        self.run = "run-a"
        self.headers = ["文件夹名", "image", "action", "sop", "summary"]
        self.rows = [
            {"文件夹名": "same-name", "image": "runs/r1/case-a/same-name/step001.jpg",
             "action": '{"action":"click"}', "sop": "original", "summary": "summary"},
            {"文件夹名": "same-name", "image": "runs/r2/case-b/same-name/step001.jpg",
             "action": '{"action":"finish"}', "sop": "original", "summary": "summary"},
        ]

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows=None, auxiliary=False):
        path = self.root / "input.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "VLA trajectories"
        sheet.append(self.headers)
        for row in (self.rows if rows is None else rows):
            sheet.append([row.get(name) for name in self.headers])
        if auxiliary:
            other = workbook.create_sheet("辅助统计")
            other.append(self.headers)
            other.append(["extra", "extra/step.jpg", '{"action":"wait"}', "", ""])
        workbook.save(path)
        workbook.close()
        return path

    def release(self, *, rows=None, metadata=True, collection=True, edits=None,
                selected=None, auxiliary=False):
        values = self.rows if rows is None else rows
        workbook = self.workbook(values, auxiliary)
        release_id = "rel_one"
        destination = self.root / "releases" / release_id / "001" / "full.xlsx"
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(workbook, destination)
        release = {"release_id": release_id, "created_at": "2026-09-17T03:01:27Z",
                   "excel_paths": [{"path": destination.relative_to(self.root).as_posix(),
                     "sha256": hashlib.sha256(destination.read_bytes()).hexdigest()}],
                   "source_refs": []}
        if not metadata:
            return release
        identity_rows = []
        tasks = []
        for index, row in enumerate(values):
            identity_rows.append({**row, "trajectory_id": f"tr_{index}", "task_id": f"task_{index}",
                 "collection_case_id": f"case_{index}", "source_result_id": f"result_{index}"})
            tasks.append({"collection_case_id": f"case_{index}", "task_id": f"task_{index}",
                "source_result_id": f"result_{index}", "app": "爱奇艺",
                "scene": "视频娱乐", "capability": "播放", "sub_capability": "播放节目"})
        refs = []
        if collection:
            refs.append(self.store.publish(self.batch, "00_collection",
                {"batch_id": self.batch, "snapshot": {"tasks": tasks}}))
        conversion = self.store.publish(self.batch, "01_conversion",
            {"sheets": {"VLA trajectories": identity_rows}}, source_refs=refs)
        annotation = self.store.publish(self.batch, "02_annotation",
            {"sheets": {"VLA trajectories": identity_rows}}, source_refs=[conversion])
        # Current annotation keeps its real upstream, never a self-version link.
        edited_annotation = annotation
        observation = self.store.publish(self.batch, "03_observation", {},
            source_refs=[edited_annotation])
        self.tree = self.store.publish(self.batch, "04_tree", {"run_id": self.run},
            source_refs=[observation], metadata={"run_id": self.run})
        frozen = selected if selected is not None else [
            {"excel_row": index + 2, "image": row["image"], "values": deepcopy(row),
             "original_action": row["action"]} for index, row in enumerate(values)]
        payload = {"session_id": "session-a", "batch_id": self.batch,
                   "tree_run_id": self.run, "row_edits": edits or {}, "groups": [{"rows": frozen}]}
        self.cot = self.store.publish(self.batch, "07_cot", payload,
            workbooks={"full_dataset.xlsx": workbook},
            source_refs=[{"kind": "trajectory_tree", "id": self.run}])
        release["source_refs"] = [{"id": "session-a", "tree_run_id": self.run, "artifact": self.cot}]
        return freeze_existing_release(self.root, release)

    def bundle(self, release):
        return json.loads((self.root / release["provenance"]["path"]).read_text(encoding="utf-8"))

    def save_bundle(self, release, bundle):
        path = self.root / release["provenance"]["path"]
        data = json.dumps(bundle, ensure_ascii=False).encode()
        path.write_bytes(data)
        release["provenance"].update(sha256=hashlib.sha256(data).hexdigest(), size=len(data))

    def entry(self, bundle, stage):
        return next(item for item in bundle["artifacts"] if item["manifest"]["stage"] == stage)

    def test_full_business_sheet_and_precise_lineage(self):
        release = self.release(selected=[{"excel_row": 2, "image": self.rows[0]["image"],
                                          "values": self.rows[0]}], auxiliary=True)
        # A later unrelated tree and a newer changed classification must not leak in.
        self.store.publish(self.batch, "00_collection",
            {"batch_id": self.batch, "snapshot": {"tasks": [{"app": "wrong"}]}})
        self.store.publish(self.batch, "04_tree", {"run_id": "other"},
            metadata={"run_id": "other"})
        result = convert_release(self.registry, release)
        self.assertEqual(len(result["rows"]), 2)
        self.assertEqual(sum(row["step数量"] for row in result["rows"]), 2)
        self.assertEqual({row["APP"] for row in result["rows"]}, {"爱奇艺"})
        self.assertEqual({row["二级场景"] for row in result["rows"]}, {"播放"})
        self.assertTrue(all(row["manual_known"] and row["人工精修步骤数量"] == 0 for row in result["rows"]))
        self.assertEqual(result["rows"][0]["时间"], "2026-09-17 11:01:27")
        self.assertEqual(result, convert_release(self.registry, release))

    def test_legacy_without_sources_included_and_unknown_manual(self):
        result = convert_release(self.registry, self.release(metadata=False))
        self.assertEqual(len(result["rows"]), 2)
        self.assertEqual(len({row["轨迹"] for row in result["rows"]}), 2)
        for row in result["rows"]:
            self.assertIsNone(row["人工精修步骤数量"])
            self.assertFalse(row["manual_known"])
            self.assertEqual(row["APP"], "未记录 App")
            self.assertEqual(row["一级场景"], "未分类")
        self.assertTrue(result["warnings"])

    def test_metadata_without_collection_keeps_manual_known_and_missing_classification(self):
        result = convert_release(self.registry, self.release(collection=False))
        self.assertTrue(all(row["manual_known"] for row in result["rows"]))
        self.assertTrue(all(not row["app_known"] for row in result["rows"]))
        self.assertEqual(sum(row["step数量"] for row in result["rows"]), 2)

    def test_manual_only_effective_action_or_sop_once_per_row(self):
        originals = [deepcopy(self.rows[0]) for _ in range(5)]
        for index, row in enumerate(originals, 1):
            row["image"] = f"task/trajectory/step{index:03}.jpg"
        finals = deepcopy(originals)
        finals[0].update(action='{"action":"swipe"}', sop="changed")
        finals[1]["summary"] = "human summary"
        finals[2]["action"] = '{ "action": "click" }'
        finals[3]["sop"] = "original"  # Edit exists, but full export did not apply it.
        finals[4]["sop"] = "changed"
        edits = {"2": {"actions": finals[0]["action"], "original_actions": originals[0]["action"], "sop": "changed"},
                 "3": {"summary": "human summary", "actions_box": "manual bbox"},
                 "4": {"actions": finals[2]["action"], "original_actions": originals[2]["action"]},
                 "5": {"sop": "changed"},
                 "6": {"sop": "changed"}}
        selected = [{"excel_row": index + 2, "image": row["image"], "values": row,
                     "original_action": row["action"]} for index, row in enumerate(originals)]
        result = convert_release(self.registry, self.release(rows=finals, edits=edits, selected=selected))
        self.assertEqual(sum(row["人工精修步骤数量"] for row in result["rows"]), 2)

    def test_final_actions_not_bounding_boxes_and_no_terminal_step_subtraction(self):
        rows = [{"文件夹名": "t", "image": f"t/step{i}.jpg", "action": action}
                for i, action in enumerate(['{"action":"click"}', '{"action":"finish"}', "broken"])]
        result = convert_release(self.registry, self.release(rows=rows, metadata=False))
        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["step数量"], 3)
        self.assertEqual(result["rows"][0]["action_box"], {"click": 1, "finish": 1, "unknown": 1})

    def test_multiple_files_and_release_versions_never_merge(self):
        release = self.release(metadata=False)
        extra = self.root / "releases/rel_one/002/full.xlsx"
        extra.parent.mkdir()
        shutil.copyfile(self.root / release["excel_paths"][0]["path"], extra)
        release["excel_paths"].append({**release["excel_paths"][0], "path": extra.relative_to(self.root).as_posix()})
        result = convert_release(self.registry, release)
        self.assertEqual(len(result["rows"]), 4)
        self.assertEqual(len({row["轨迹"] for row in result["rows"]}), 4)

    def test_missing_tampered_or_wrong_release_file_rejected(self):
        release = self.release(metadata=False)
        path = self.root / release["excel_paths"][0]["path"]
        path.write_bytes(b"tampered")
        with self.assertRaisesRegex(ConversionError, "SHA256"):
            convert_release(self.registry, release)
        path.unlink()
        with self.assertRaisesRegex(ConversionError, "不存在"):
            convert_release(self.registry, release)
        release["excel_paths"][0]["path"] = "input.xlsx"
        with self.assertRaisesRegex(ConversionError, "冻结文件"):
            convert_release(self.registry, release)

    def test_frozen_registration_and_stage_sha_checked(self):
        release = self.release()
        changed = deepcopy(release)
        changed["source_refs"][0]["artifact"]["files"][0]["sha256"] = "0" * 64
        with self.assertRaisesRegex(ConversionError, "登记信息"):
            convert_release(self.registry, changed)
        store = FrozenReleaseSources(self.root, release)
        store.resolve_file(self.cot, "result.json").write_text("{}", encoding="utf-8")
        with self.assertRaisesRegex(ConversionError, "校验"):
            convert_release(self.registry, release)

    def test_explicit_tree_must_be_unique(self):
        release = self.release()
        bundle = self.bundle(release)
        duplicate = deepcopy(self.entry(bundle, "04_tree"))
        duplicate["manifest"]["version"] += "-copy"
        bundle["artifacts"].append(duplicate)
        self.save_bundle(release, bundle)
        with self.assertRaisesRegex(ConversionError, "唯一"):
            convert_release(self.registry, release)

    def test_explicit_tree_missing_rejected(self):
        release = self.release()
        bundle = self.bundle(release)
        bundle["artifacts"] = [item for item in bundle["artifacts"] if item["manifest"]["stage"] != "04_tree"]
        self.save_bundle(release, bundle)
        with self.assertRaisesRegex(ConversionError, "唯一"):
            convert_release(self.registry, release)

    def test_cross_batch_reference_rejected(self):
        release = self.release()
        bundle = self.bundle(release)
        foreign = deepcopy(self.entry(bundle, "03_observation"))
        foreign["manifest"]["batch_id"] = "other-batch"
        self.entry(bundle, "04_tree")["manifest"]["source_refs"] = [foreign["manifest"]]
        bundle["artifacts"].append(foreign)
        self.save_bundle(release, bundle)
        with self.assertRaisesRegex(ConversionError, "跨越批次"):
            convert_release(self.registry, release)

    def test_explicit_workbook_classification_retained(self):
        self.headers.extend(["APP", "一级场景", "二级场景"])
        rows = deepcopy(self.rows)
        for row in rows:
            row.update({"APP": "本表App", "一级场景": "Unclassified", "二级场景": ""})
        result = convert_release(self.registry, self.release(rows=rows, metadata=False))
        self.assertEqual(result["rows"][0]["APP"], "本表App")
        self.assertEqual(result["rows"][0]["一级场景"], "Unclassified")
        self.assertFalse(result["rows"][0]["level1_known"])
        self.assertFalse(result["rows"][0]["level2_known"])

    def test_singleton_action_list_and_case_are_normalized(self):
        rows = [{"文件夹名": "t", "image": f"t/step{i}.jpg", "action": action}
                for i, action in enumerate([
                    '[{"action":"CLICK","coordinate":[1,2]}]', '{"action":"Click"}',
                    '[{"action":"wait"},{"action":"finish"}]', '[]'])]
        result = convert_release(self.registry, self.release(rows=rows, metadata=False))
        self.assertEqual(result["rows"][0]["action_box"], {"click": 2, "unknown": 2})
        self.assertEqual(result["rows"][0]["step数量"], 4)

    def test_classification_placeholders_not_counted_as_real_coverage(self):
        self.headers.extend(["APP", "一级场景", "二级场景"])
        for placeholder in ("Unclassified", "unclassified", "未分类", "未知", "unknown", "未记录"):
            with self.subTest(placeholder=placeholder):
                rows = deepcopy(self.rows)
                for row in rows:
                    row.update({"APP": "未记录 App", "一级场景": placeholder, "二级场景": placeholder})
                result = convert_release(self.registry, self.release(rows=rows, metadata=False))
                converted = result["rows"][0]
                self.assertEqual(converted["一级场景"], placeholder)
                self.assertFalse(converted["app_known"])
                self.assertFalse(converted["level1_known"])
                self.assertFalse(converted["level2_known"])

    def test_explicit_classification_mismatch_rejected(self):
        release = self.release()
        bundle = self.bundle(release)
        entry = self.entry(bundle, "00_collection")
        file = entry["files"]["result.json"]
        path = self.root / "releases" / release["release_id"] / "provenance" / file["name"]
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["snapshot"]["tasks"] = [{"collection_case_id": "different", "task_id": "different"}]
        data = json.dumps(payload).encode()
        path.write_bytes(data)
        file.update(sha256=hashlib.sha256(data).hexdigest(), size=len(data))
        self.save_bundle(release, bundle)
        with self.assertRaisesRegex(ConversionError, "无法唯一关联"):
            convert_release(self.registry, release)

    def test_missing_explicit_upstream_snapshot_rejected(self):
        release = self.release()
        frozen = FrozenReleaseSources(self.root, release)
        conversion = frozen.list(self.batch, "01_conversion")[0]
        frozen.resolve_file(conversion, "result.json").unlink()
        with self.assertRaises(ConversionError):
            convert_release(self.registry, release)

    def test_stage_cycles_rejected(self):
        release = self.release()
        bundle = self.bundle(release)
        previous = self.entry(bundle, "02_annotation")["manifest"]
        previous["source_refs"] = [{key: previous[key] for key in ("batch_id", "stage", "version")}]
        self.save_bundle(release, bundle)
        with self.assertRaisesRegex(ConversionError, "循环"):
            convert_release(self.registry, release)

    def test_current_batch_files_can_be_removed_without_changing_release(self):
        release = self.release()
        expected = convert_release(self.registry, release)
        shutil.rmtree(self.root / "batches")
        self.assertEqual(convert_release(self.registry, release), expected)

    def test_precise_sources_require_independent_release_provenance(self):
        release = self.release()
        release.pop("provenance")
        with self.assertRaisesRegex(ConversionError, "迁移"):
            convert_release(self.registry, release)

    def test_input_files_unchanged(self):
        release = self.release()
        def hashes():
            return {path.relative_to(self.root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
                    for path in self.root.rglob("*") if path.is_file()}
        before = hashes()
        convert_release(self.registry, release)
        self.assertEqual(before, hashes())


if __name__ == "__main__":
    unittest.main()

