"""Offline platform HTTP -> actual Flask -> dispatcher -> report/archive integration.

Set PHONE_FACTORY_EVALUATION_ROOT to the reviewed second-stage collector working
copy. Phones and evaluator subprocesses are simulated; no network or ADB runs.
"""
from __future__ import annotations

import base64
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import urllib.error
from unittest.mock import patch
from urllib.parse import urlsplit
import zipfile

from openpyxl import Workbook, load_workbook
from backend import phonefactory_client
from backend.phone_factory import PhoneFactoryStore
from backend.task_generation.collection_batches import COLLECTION_COLUMNS
from backend.tests.test_trajectories_preprocessing import create_step

SNAPSHOT = os.environ.get('PHONE_FACTORY_EVALUATION_ROOT')


class FakeEvaluationProcess:
    serial = 800000

    def __init__(self, command, **kwargs):
        self.command, self.options, self.code = command, kwargs, None
        type(self).serial += 1
        self.pid = self.serial

    def poll(self):
        return self.code

    def wait(self, timeout=None):
        return self.code

    def argument(self, key):
        return self.command[self.command.index(key) + 1]


@unittest.skipUnless(SNAPSHOT, 'set PHONE_FACTORY_EVALUATION_ROOT for actual Flask integration')
class ModelIterationFlaskTests(unittest.TestCase):
    def setUp(self):
        self.stack = contextlib.ExitStack()
        self.addCleanup(self.stack.close)
        self.base = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.source = Path(SNAPSHOT)
        spec = importlib.util.spec_from_file_location('flask_modeliter_contract', self.source / 'phonefactory_manager.py')
        self.manager = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = self.manager
        spec.loader.exec_module(self.manager)
        self.collector = self.base / 'collector'
        self.inputs = self.base / 'offline_inputs'
        self.inputs.mkdir()
        template = self.collector / 'new-guigent'
        template.mkdir(parents=True)
        shutil.copy2(self.source / 'new-guigent/logprobs_export.py', template / 'logprobs_export.py')
        for key, value in {'BASE_DIR':self.collector, 'TEMPLATE_DIR':template,
                           'HTTP_RUNS_DIR':self.collector / '.http_runs',
                           'HTTP_ARCHIVES_DIR':self.collector / '.http_archives'}.items():
            self.stack.enter_context(patch.object(self.manager, key, value.as_posix()))
        self.stack.enter_context(patch.dict(os.environ, {
            'PHONE_FACTORY_EVALUATOR_DIR':str(self.source / 'evaluator_runtime'),
            'PHONE_FACTORY_EVALUATOR_PYTHON':sys.executable,
            'PHONE_FACTORY_EVALUATOR_PROFILES_DIR': str(self.source / 'evaluator_runtime/profiles'),
            'PHONE_FACTORY_EVALUATION_INPUT_ROOTS': json.dumps([str(self.inputs)]),
            'PHONE_FACTORY_EVALUATOR_TEXT_API_KEY':'offline-test-text',
            'PHONE_FACTORY_EVALUATOR_VISION_API_KEY':'offline-test-vision',
            'GUI_AGENT_VLA_CAPTURE_LOGPROBS':'true',
        }))
        self.launch = self.stack.enter_context(patch.object(self.manager, 'launch_onephonerun'))
        self.stack.enter_context(patch.object(self.manager, '_watch_http_run'))
        self.stack.enter_context(patch.object(self.manager, '_phone_process_stopped', return_value=False))
        self.dispatcher = self.manager._evaluation_dispatcher()
        self.stack.enter_context(patch.object(self.dispatcher, 'start'))
        self.jobs = []
        def spawn(command, **kwargs):
            job = FakeEvaluationProcess(command, **kwargs)
            self.jobs.append(job)
            return job
        self.dispatcher.process_factory = spawn
        self.stack.callback(self.dispatcher.close)
        self.remote = self.manager.app.test_client()
        self.factory = PhoneFactoryStore(self.base / 'platform', run_client_fn=self.client)
        self.factory.initialize_settings({'phones':['phone-a'], 'apps':['App'],
            'phoneApps':[{'phone_id':'phone-a', 'app':'App'}], 'vla':['target.invalid:8000']})
        book = Workbook()
        book.active.append(COLLECTION_COLUMNS)
        for case in ('A', 'A1'):
            book.active.append([case,'单APP','App','购物','搜索',None,'搜索商品' + case,
                                None,'必须出现商品' + case,2,'人工',None,None,None,None,None,None])
        output = io.BytesIO()
        book.save(output)
        book.close()
        self.factory.add_task({'filename':'evaluation.xlsx','description':'offline evaluation',
            'content_base64':base64.b64encode(output.getvalue()).decode()}, mode='modeliter')
        self.config = {role:{'base_url':f'https://judge-{role}.invalid/gateway/v1/',
            'model':f'judge-{role}', 'api_key_env':f'PHONE_FACTORY_EVALUATOR_{role.upper()}_API_KEY'}
            for role in ('text','vision')}
        self.factory.dispatch('evaluation-config','PUT',self.config,mode='modeliter')

    def urlopen(self, request, timeout=None):
        url = urlsplit(request.full_url)
        response = self.remote.open(url.path + ('?' + url.query if url.query else ''),
            method=request.get_method(), data=request.data, headers=dict(request.header_items()))
        body, headers, status = response.get_data(), response.headers, response.status_code
        response.close()
        if status >= 400:
            raise urllib.error.HTTPError(request.full_url,status,body.decode(),headers,io.BytesIO(body))
        stream = io.BytesIO(body)
        stream.headers = headers
        return stream

    def client(self, args):
        output = io.StringIO()
        with patch.object(sys,'argv',['phonefactory_client',*args]), \
                patch.object(phonefactory_client.urllib.request,'urlopen',self.urlopen), \
                contextlib.redirect_stdout(output):
            try:
                phonefactory_client.main()
            except SystemExit as exc:
                return {'ok':False,'output':str(exc)}
        return {'ok':True,'output':output.getvalue()}

    def start(self, request_id='one'):
        result = self.factory.remote_start({'filename':'evaluation.xlsx','request_id':request_id,
            'vla':'target.invalid:8000', 'config':{'capture_logprobs':True}},mode='modeliter')
        record = self.manager._load_http_run(result['run_id'])
        self.assertEqual(record['run_mode'],'modeliter')
        self.assertFalse(record['batch_id'])
        self.assertEqual(record['params']['evaluator_config'],self.config)
        self.assertEqual(record['params']['checklist_source'],'llm')
        self.assertTrue(record['params']['capture_logprobs'])
        self.assertEqual(self.factory.records.list('collection_runs'),[])
        return result,record

    def ready(self, record, index):
        case = record['evaluation_cases'][index]
        run_name = record['evaluation_source_run_id']
        directory = self.collector / ('new-guigent-' + case['phone_id']) / '.runs' / run_name
        folder = directory / case['expanded_case_id']
        folder.mkdir(parents=True,exist_ok=True)
        create_step(folder,1,{'action':'wait','time':1})
        session = case['expanded_case_id'] + '_' + run_name
        payload = {'task':case['source_row']['任务'],'session_id':session,
            'actions_flat':[{'global_step':1,'action':{'action':'wait','time':1}}]}
        trajectory = folder / '_trajectory_for_evaluate.json'
        trajectory.write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
        os.utime(trajectory,(record['created_epoch']+1,record['created_epoch']+1))
        trace = {'status':'ok','response_token_ids':[101+index],'behavior_logprobs':[-0.1],
            'response_loss_mask':[1],'sampling':{'temperature':0.7,'top_p':0.85}}
        response_path = folder / 'step001_vla_model_response.json'
        response = json.loads(response_path.read_text(encoding='utf-8'))
        response['behavior_trace'] = trace
        response_path.write_text(json.dumps(response),encoding='utf-8')
        with (directory / '_eval_queue.txt').open('a',encoding='utf-8') as queue:
            queue.write(session + '\n')
        return folder

    def finish_job(self, job, verdict='success', invalid=False):
        folder = Path(job.argument('--output-dir'))
        trajectory = Path(job.argument('--trajectory-path'))
        raw_config = json.loads(Path(job.argument('--config-json')).read_text(encoding='utf-8'))
        digest = hashlib.sha256(json.dumps(raw_config,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
        states = ['passed'] if verdict=='success' else ['failed'] if verdict=='fail' else ['passed','failed']
        report = {'session_id':job.argument('--session-id'),'task':'test','trajectory_sha256':hashlib.sha256(trajectory.read_bytes()).hexdigest(),
            'evaluator_config_digest':digest,'checklist_policy':'llm','checklist_source':'llm','evaluation_mode':'checklist',
            'evaluation_status':'complete','verdict':verdict,'failure_analysis':{'root_cause':'' if verdict=='success' else '验证目标未满足'},
            'key_elements_sha256':hashlib.sha256(job.argument('--key-elements').encode()).hexdigest(),
            'checklist_summary':{'total':len(states),'passed':states.count('passed'),'failed':states.count('failed'),'unchecked':0,
                'items':[{'id':f'c{i}','description':f'check{i}','status':status,'evidence':'offline'} for i,status in enumerate(states)]}}
        if invalid:
            report['trajectory_sha256'] = 'wrong-trajectory'
        if '--profile-snapshot' in job.command:
            profile = json.loads(Path(job.argument('--profile-snapshot')).read_text(encoding='utf-8'))
            report.update(checklist_policy='auto', checklist_source='llm',
                          **{key: profile[key] for key in ('profile_id', 'profile_digest', 'runtime_digest')})
        path = folder / '_report_offline.json'
        path.write_text(json.dumps(report,ensure_ascii=False),encoding='utf-8')
        (folder / '_evaluation_result.json').write_text(json.dumps({'status':'complete','report_path':str(path)}),encoding='utf-8')
        job.code = 0

    def results(self, run_id):
        return self.factory.dispatch(f'runs/{run_id}/evaluations','GET',{},mode='modeliter')

    def archive_collected(self, record, folders):
        for folder in folders:
            shutil.copytree(folder,Path(record['result_dir']) / folder.name,dirs_exist_ok=True)
        with patch.object(self.manager,'_phone_process_stopped',return_value=True), \
                patch.object(self.manager,'_phone_run_state',return_value={'status':'completed'}):
            return self.manager._finalize_http_run(record['collection_run_id'])

    def profile(self):
        value = self.factory.dispatch('evaluation-profiles', 'GET', {}, mode='modeliter')
        self.assertEqual(value['errors'], [], value)
        self.assertTrue(value['profiles'], value)
        return value['profiles'][0]

    def offline_input(self):
        for folder_name in ('first', 'second'):
            folder = self.inputs / folder_name
            folder.mkdir()
            create_step(folder, 1, {'action': 'wait', 'time': 1})
            payload = {'task_id': 'same-original-id', 'task': '搜索商品', 'session_id': folder_name,
                'actions_flat': [{'global_step': 1, 'action': {'action': 'wait', 'time': 1},
                                  'screenshot_before': r'Z:\old-server\step001_vla_input.jpg'}]}
            (folder / '_trajectory_for_evaluate.json').write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        return self.factory.dispatch('evaluation-preview', 'POST', {'source_path': str(self.inputs)}, mode='modeliter')

    def start_offline(self, profile, preview, request_id='offline'):
        return self.factory.dispatch('remote/start-evaluation', 'POST', {
            'request_id': request_id, 'evaluator_config': self.config,
            'evaluation_profile_id': profile['id'], 'evaluation_profile_digest': profile['digest'],
            'preview_id': preview['preview_id'], 'checklist_source': 'auto'}, mode='modeliter')

    def test_offline_http_profile_preview_progress_cancel_without_phone_launch(self):
        profile = self.profile()
        sources = self.factory.dispatch('evaluation-sources', 'GET', {}, mode='modeliter')
        self.assertEqual(sources['allowed_roots'], [str(self.inputs)])
        preview = self.offline_input()
        self.assertTrue(preview['valid'], preview)
        self.assertEqual(preview['trajectory_count'], 2)
        self.assertEqual(len({case['source_trajectory_id'] for case in preview['cases']}), 2)
        self.assertEqual({case['case_id'] for case in preview['cases']}, {'same-original-id'})
        # The independent entry must not depend on platform phone/VLA settings.
        self.factory._update(lambda state: state.update(phones=[], phoneApps=[], vla=[]))
        run = self.start_offline(profile, preview)
        self.assertEqual(run['execution_mode'], 'evaluate_only')
        self.assertEqual(self.launch.call_count, 0)
        record = self.manager._load_http_run(run['run_id'])
        self.assertEqual(record['phone_ids'], [])
        self.assertEqual(record['evaluation_profile']['profile_digest'], profile['digest'])
        for case in record['evaluation_cases']:
            self.assertTrue(Path(case['trajectory_path']).is_file())
            self.assertFalse(Path(case['trajectory_path']).is_relative_to(self.inputs))
        self.dispatcher.tick()
        self.assertEqual(len(self.jobs), 2)
        self.assertTrue(all(job.argument('--checklist-source') == 'auto' for job in self.jobs))
        self.assertTrue(all(job.options['env']['GUI_AGENT_VLA_CAPTURE_LOGPROBS'] == 'false' for job in self.jobs))
        state = self.results(run['run_id'])
        self.assertEqual(state['summary']['collection_completed'], 2)
        self.assertEqual(state['summary']['evaluation_running'], 2)
        self.assertEqual(state['source_digest'], preview['source_digest'])
        self.assertEqual({row['source_trajectory_id'] for row in state['cases']},
                         {row['source_trajectory_id'] for row in preview['cases']})
        with patch.object(self.dispatcher, '_stop_pid'), patch.object(self.manager, 'kill_processes_in') as phone_kill:
            cancelled = self.factory.dispatch(f"runs/{run['run_id']}/cancel", 'POST', {}, mode='modeliter')
        self.assertEqual(cancelled['status'], 'cancelled')
        phone_kill.assert_not_called()
        self.assertEqual(self.results(run['run_id'])['summary']['evaluation_failed'], 2)
        self.assertEqual(self.factory.records.list('collection_runs'), [])

    def test_offline_http_final_reports_archive_and_re_evaluation_are_separate(self):
        profile, preview = self.profile(), self.offline_input()
        run = self.start_offline(profile, preview)
        self.dispatcher.tick()
        self.finish_job(self.jobs[0], 'success')
        self.finish_job(self.jobs[1], 'partial_success')
        self.dispatcher.tick()
        # The scheduler finalizes independent runs only after their cases settle.
        record = self.manager._load_http_run(run['run_id'])
        if record['status'] not in self.manager.HTTP_TERMINAL:
            record = self.manager._finalize_http_run(run['run_id'])
        self.assertEqual(record['status'], 'completed', record.get('errors'))
        state = self.results(run['run_id'])
        self.assertEqual(state['summary']['success_rate'], 0.5)
        self.assertEqual(state['summary']['partial_success'], 1)
        self.assertEqual({row['checklist_source'] for row in state['cases']}, {'llm'})
        reports = self.factory.dispatch('reports', 'GET', {}, mode='modeliter')['folders']
        folder = next(item for item in reports if item.get('run_id') == run['run_id'])
        self.assertTrue(any(item['name'].endswith('.xlsx') for item in folder['files']))
        with zipfile.ZipFile(record['archive_path']) as archive:
            names = archive.namelist()
            self.assertTrue(any(name.startswith('evaluation_profile/') for name in names))
            self.assertEqual(sum(name.endswith('_trajectory_for_evaluate.json') for name in names), 2)
        sources = self.factory.dispatch('evaluation-sources', 'GET', {}, mode='modeliter')
        self.assertTrue(any(source['run_id'] == run['run_id'] for source in sources['sources']))
        archived_preview = self.factory.dispatch('evaluation-preview', 'POST', {'source_run_id': run['run_id']}, mode='modeliter')
        self.assertTrue(archived_preview['valid'], archived_preview)
        self.assertEqual(archived_preview['trajectory_count'], 2)
        # Selecting a completed run creates new output/checkpoint directories.
        second = self.start_offline(profile, archived_preview, request_id='offline-second')
        self.assertNotEqual(run['run_id'], second['run_id'])
        self.assertEqual(self.results(second['run_id'])['summary']['evaluation_completed'], 0)
        self.assertEqual(self.results(run['run_id']), state)
        self.assertEqual(self.launch.call_count, 0)

    def test_online_auto_profile_is_frozen_and_forwarded_to_executor(self):
        profile = self.profile()
        run = self.factory.remote_start({'filename': 'evaluation.xlsx', 'request_id': 'online-auto',
            'vla': 'target.invalid:8000', 'evaluator_config': self.config, 'checklist_source': 'auto',
            'evaluation_profile_id': profile['id'], 'evaluation_profile_digest': profile['digest']}, mode='modeliter')
        record = self.manager._load_http_run(run['run_id'])
        self.assertEqual(record['params']['checklist_source'], 'auto')
        self.assertEqual(record['params']['evaluation_profile_digest'], profile['digest'])
        self.assertEqual(record['evaluation_profile']['profile_id'], profile['id'])
        self.ready(record, 0)
        self.dispatcher.tick()
        self.assertEqual(len(self.jobs), 1)
        self.assertEqual(self.jobs[0].argument('--checklist-source'), 'auto')
        self.assertTrue(Path(self.jobs[0].argument('--profile-snapshot')).is_file())
        self.assertEqual(self.launch.call_count, 1)

    def test_online_auto_lost_ack_recovers_frozen_run_after_profile_catalog_changes(self):
        profile = self.profile()
        payload = {'filename': 'evaluation.xlsx', 'request_id': 'online-auto-lost-ack',
            'vla': 'target.invalid:8000', 'evaluator_config': self.config, 'checklist_source': 'auto',
            'evaluation_profile_id': profile['id'], 'evaluation_profile_digest': profile['digest']}
        def lose_start_ack(args):
            response = self.client(args)
            return {'ok': False, 'output': 'lost acknowledgement'} if args[0] == 'start-run' else response
        self.factory.run_client = lose_start_ack
        with self.assertRaisesRegex(Exception, 'lost acknowledgement'):
            self.factory.remote_start(payload, mode='modeliter')
        def changed_catalog(args):
            if args[0] == 'evaluation-profiles':
                return {'ok': True, 'output': json.dumps({'profiles': [{**profile, 'digest': '0' * 64}], 'errors': []})}
            return self.client(args)
        self.factory.run_client = changed_catalog
        run = self.factory.remote_start(payload, mode='modeliter')
        self.assertTrue(run['reused'])
        self.assertEqual(self.launch.call_count, 1)
        record = self.manager._load_http_run(run['run_id'])
        self.assertEqual(record['evaluation_profile']['profile_digest'], profile['digest'])

    def test_http_pipeline_overlaps_cases_and_freezes_complete_report_archive(self):
        run,record = self.start()
        first = self.ready(record,0)
        self.dispatcher.tick()
        self.assertEqual(len(self.jobs),1)
        self.assertEqual(self.jobs[0].options['env']['GUI_AGENT_VLA_CAPTURE_LOGPROBS'],'false')
        first_state = self.results(run['run_id'])
        self.assertEqual(first_state['summary']['collection_completed'],1)
        self.assertEqual(first_state['summary']['evaluation_running'],1)
        self.assertEqual(first_state['cases'][1]['collection_status'],'pending')
        with self.assertRaises(Exception):
            self.start('busy-phone')
        # Phone continues B while A's evaluation process is still running.
        second = self.ready(record,1)
        self.dispatcher.tick()
        self.assertEqual(len(self.jobs),2)
        self.assertIsNone(self.jobs[0].code)
        self.assertNotEqual(self.jobs[0].argument('--trajectory-path'),self.jobs[1].argument('--trajectory-path'))
        self.assertEqual([job.argument('--case-id') for job in self.jobs],['A','A1'])
        self.assertFalse(Path(record['archive_path']).exists())
        self.finish_job(self.jobs[0],'success')
        self.finish_job(self.jobs[1],'fail')
        self.dispatcher.tick()
        state = self.results(run['run_id'])
        self.assertEqual(state['summary']['evaluation_completed'],2)
        self.assertEqual(state['summary']['success_rate'],0.5)
        frozen = self.archive_collected(record,[first,second])
        self.assertEqual(frozen['status'],'completed',frozen.get('errors'))
        saved = self.factory.runtime.sync_evaluation(run['run_id'])
        self.assertEqual(saved['status'],'completed')
        self.assertTrue(self.factory.records.get('model_iter_results',run['run_id'])['final'])
        folders = self.factory.dispatch('reports','GET',{},mode='modeliter')['folders']
        report_folder = next(item for item in folders if item.get('run_id')==run['run_id'])
        self.assertEqual(sum(item['name'].endswith('.json') for item in report_folder['files']),2)
        excel = next(item for item in report_folder['files'] if item['name'].endswith('.xlsx'))
        response = self.remote.get('/report_download',query_string={'run_id':run['run_id'],'file_id':excel['file_id']})
        self.assertEqual(response.status_code,200)
        book = load_workbook(io.BytesIO(response.data),read_only=True)
        self.assertEqual(book.sheetnames,['summary','checklist_details'])
        self.assertEqual(book['summary'].max_row,3)
        book.close()
        response.close()
        archive = Path(frozen['archive_path']).read_bytes()
        self.assertEqual(frozen['manifest']['archive']['sha256'],hashlib.sha256(archive).hexdigest())
        with zipfile.ZipFile(io.BytesIO(archive)) as zipped:
            self.assertEqual(sum(name.endswith('_logprobs.json') for name in zipped.namelist()),2)
            self.assertGreaterEqual(sum(name.endswith('.json') and 'evaluation' in name for name in zipped.namelist()),2)
        self.assertEqual(self.factory.records.list('collection_runs'),[])
        # A new run uses a new result namespace even for the same case IDs.
        next_run,next_record = self.start('two')
        self.assertNotEqual(next_run['run_id'],run['run_id'])
        self.assertEqual(self.results(next_run['run_id'])['summary']['evaluation_completed'],0)

    def test_invalid_report_is_retryable_not_counted_as_model_failure(self):
        run,record = self.start()
        self.ready(record,0)
        self.dispatcher.tick()
        first = self.jobs[0]
        self.finish_job(first,invalid=True)
        self.dispatcher.tick()
        state = self.results(run['run_id'])
        self.assertEqual(state['cases'][0]['evaluation_status'],'retry_wait')
        self.assertEqual(state['summary']['evaluation_completed'],0)
        self.assertEqual(state['summary']['fail'],0)
        self.assertIsNone(state['summary']['success_rate'])
        self.dispatcher.clock = lambda: record['created_epoch']+100
        self.dispatcher.tick()
        self.assertEqual(len(self.jobs),2)
        self.assertEqual(first.argument('--output-dir'),self.jobs[1].argument('--output-dir'))
        self.assertEqual(first.argument('--trajectory-path'),self.jobs[1].argument('--trajectory-path'))
        self.assertIn('--resume-auto',self.jobs[1].command)
        self.assertEqual(self.launch.call_count,1)
        self.assertFalse(Path(record['archive_path']).exists())

    def test_cancel_freezes_collected_and_uncollected_cases_without_model_verdicts(self):
        run,record = self.start()
        self.ready(record,0)
        self.dispatcher.tick()
        with patch.object(self.dispatcher,'_stop_pid'), patch.object(self.manager,'kill_processes_in',return_value=0):
            response = self.remote.post('/runs/' + run['run_id'] + '/cancel')
        self.assertEqual(response.status_code,200,response.get_json())
        result = self.results(run['run_id'])
        self.assertFalse(result.get('stale'))
        self.assertEqual(result['summary']['collection_completed'],1)
        self.assertEqual(result['summary']['collection_failed'],1)
        self.assertEqual(result['summary']['evaluation_failed'],2)
        self.assertIsNone(result['summary']['success_rate'])
        self.assertEqual([case['evaluation_status'] for case in result['cases']],['needs_review','collection_failed'])
        saved = self.factory.records.get('model_iter_results',run['run_id'])
        self.assertTrue(saved['final'])
        self.assertEqual(self.factory.records.get('model_iter_runs',run['run_id'])['status'],'cancelled')
        self.assertFalse(self.dispatcher.active)


if __name__ == '__main__':
    unittest.main()
