"""Offline Pipeline selection, frozen publication and uncertain-commit recovery."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from backend import pipeline_release as publishing
from backend.batch_lifecycle import lifecycle, BatchNotReadyError
from backend.batch_results import annotation_task_fingerprints, digest
from backend.data_store import ArtifactStore, RecordStore, RevisionConflict
from backend.data_store.release_provenance import FrozenReleaseSources
from backend.data_publishing.service import DatasetReleaseRegistry
from backend.training_data_overview.converter import convert_release


class PipelineReleaseTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.raw = self.root / "raw" / "rollout_trajectories"
        self.raw.mkdir(parents=True)
        self.store, self.records = ArtifactStore(self.root), RecordStore(self.root)
        self.batch = "batch-pipeline"
        self.registry = DatasetReleaseRegistry(data_root=self.root, releases_file=self.root / "registry.json")
        self.seed()

    def seed(self, scores=None):
        scores = scores or {"A": {"a-second": 5, "a-first": 5}, "B": {"b-first": 2, "b-second": 1}}
        tasks, rows = [], []
        for task, choices in scores.items():
            tasks.append({"task_id": task, "trajectory_count": len(choices), "goal": "task " + task,
                "collection_case_id": "case-" + task, "app": "App " + task, "scene": "场景", "capability": "分类"})
            for trajectory in choices:
                for step in range(1, 4):
                    rows.append({"文件夹名": trajectory, "task_id": task, "trajectory_id": trajectory,
                        "collection_case_id": "case-" + task, "image": f"{task}/{trajectory}/step{step:03}.jpg",
                        "action": '{"action":"wait"}', "summary": "baseline", "sop": "original SOP", "actions_box": "wait()"})
        table = {"schema_version": 1, "columns": {"VLA trajectories": list(rows[0])}, "sheets": {"VLA trajectories": rows}}
        collection = self.store.publish(self.batch, "00_collection", {"batch_id": self.batch, "snapshot": {"tasks": tasks}})
        converted = self.store.publish(self.batch, "01_conversion", table, source_refs=[collection])
        annotated = self.store.publish(self.batch, "02_annotation", table, source_refs=[converted])
        fingerprints = annotation_task_fingerprints(table)
        trees = {item["task_id"]: {"steps": [], "task_id": item["task_id"]} for item in tasks}
        hashes = {task: digest(value) for task, value in trees.items()}
        refs = self.store.publish_many(self.batch, [
            {"stage": "03_observation", "payload": {}, "source_refs": [annotated]},
            {"stage": "04_tree", "payload": {"batch_id": self.batch, "run_id": self.batch,
                "trees": trees, "tasks": tasks, "task_fingerprints": fingerprints, "tree_hashes": hashes,
                "source_task_fingerprints": fingerprints, "source_annotation": table, "raw_root": str(self.raw)},
             "metadata": {"run_id": self.batch}, "source_stages": ["03_observation"]}])
        self.store.publish(self.batch, "05_quality", {"tasks": [
            {"task_id": task, "status": "succeeded", "evaluations": {
                trajectory: {"global_score": score} for trajectory, score in choices.items()}}
            for task, choices in scores.items()], "source_tree_hashes": hashes, "task_fingerprints": fingerprints},
            source_refs=[refs[-1]])

    def pipeline(self, mode="automatic", threshold=4):
        selection = publishing.build_selection(self.batch, ["A", "B"], mode, threshold, root=self.root)
        return {"pipeline_id": "pipe-test", "name": "Pipeline release", "batch_id": self.batch,
                "mode": mode, "selection": selection}

    def publish(self, pipeline):
        return publishing.publish_pipeline(pipeline, root=self.root, registry=self.registry)

    def test_threshold_tie_stability_and_automatic_no_session(self):
        pipeline = self.pipeline()
        self.assertEqual(pipeline["selection"]["selected_trajectories"], {"A": "a-first"})
        self.assertEqual(pipeline["selection"]["excluded"][0]["task_id"], "B")
        release = self.publish(pipeline)
        self.assertEqual((release["task_count"], release["trajectory_count"], release["step_count"]), (1, 1, 3))
        self.assertEqual(self.records.list("correction_sessions"), [])
        self.assertEqual(lifecycle(self.batch, self.root)["status"], "published")
        self.assertEqual(self.publish(pipeline)["release_id"], release["release_id"])
        self.assertEqual(len(self.registry.list_releases()), 1)
        workbook, _ = self.registry.excel_file(release["release_id"], 0)
        book = load_workbook(workbook)
        self.assertEqual(len(book.worksheets), 1)
        self.assertEqual(book.active.max_row, 4)
        columns = [cell.value for cell in book.active[1]]
        self.assertEqual(book.active.cell(2, columns.index("sop") + 1).value, "original SOP")
        book.close()
        converted = convert_release(self.registry, release)
        self.assertEqual(len(converted["rows"]), 1)
        self.assertEqual(converted["rows"][0]["step数量"], 3)
        self.assertEqual(converted["rows"][0]["人工精修步骤数量"], 0)
        self.assertEqual(converted["rows"][0]["APP"], "App A")
        frozen = FrozenReleaseSources(self.root, release)
        artifact = release["source_refs"][0]["artifact"]
        self.assertTrue(frozen.resolve_file(artifact, "result.json").is_file())

    def test_automatic_rebases_existing_edits_and_cot_without_using_old_export_choices(self):
        manual_selection = self.pipeline("manual")["selection"]
        created = publishing.prepare_manual(self.batch, manual_selection, root=self.root)
        session = self.records.get("correction_sessions", created["session_id"])
        current_groups = deepcopy(session["snapshot_payload"]["groups"])
        picked = next(group for group in current_groups if group["meta_task"] == "a-first")
        excluded = next(group for group in current_groups if group["meta_task"] == "a-second")
        # The saved workbook previously listed task B before A. Each task's
        # inputs and stable step identities are identical; Excel offsets differ.
        source_rows = session["workbook_payload"]["sheets"]["VLA trajectories"]
        session["workbook_payload"]["sheets"]["VLA trajectories"] = source_rows[6:] + source_rows[:6]
        mapping = {}
        for group in session["snapshot_payload"]["groups"]:
            for row in group["rows"]:
                previous = row["excel_row"] + (6 if group["task_id"] == "A" else -6)
                mapping[str(row["excel_row"])] = str(previous)
                row["excel_row"] = previous
        session["bbox_baselines"] = {mapping[number]: value for number, value in session["bbox_baselines"].items()}
        chosen_number = mapping[str(picked["rows"][0]["excel_row"])]
        excluded_number = mapping[str(excluded["rows"][0]["excel_row"])]
        final_action = '{"action":"wait","time":3}'
        session["row_edits"] = {
            chosen_number: {"actions": final_action, "original_actions": '{"action":"wait"}', "thought": "saved handwritten"},
            excluded_number: {"actions": '{"action":"click","coordinate":[1,2]}', "original_actions": '{"action":"wait"}'} }
        action_hash = publishing.fingerprint({"action": "wait", "time": 3})
        session["cot"] = {chosen_number: {"content_tag": "thought_summary", "action_hash": action_hash,
            "summary": "saved generated summary", "thought": "generated thought"}}
        session["group_exports"] = {group["group_id"]: group["meta_task"] in {"a-second", "b-first"} for group in current_groups}
        saved = self.records.put("correction_sessions", session["session_id"], session)
        release = self.publish(self.pipeline())
        self.assertEqual((release["task_count"], release["step_count"]), (1, 3))
        self.assertEqual(release["selection"]["selected_trajectories"], {"A": "a-first"})
        existing = self.records.list("correction_sessions")
        self.assertEqual(len(existing), 1)
        self.assertEqual(existing[0]["row_edits"], saved["row_edits"])
        self.assertEqual(existing[0]["group_exports"], saved["group_exports"])
        self.assertTrue(existing[0]["published"])
        book = load_workbook(self.registry.excel_file(release["release_id"], 0)[0])
        headers = [cell.value for cell in book.active[1]]
        row = dict(zip(headers, next(book.active.iter_rows(min_row=2, max_row=2, values_only=True))))
        book.close()
        self.assertEqual(row["trajectory_id"], "a-first")
        self.assertEqual(row["action"], final_action)
        self.assertEqual(row["thought"], "saved handwritten")
        self.assertEqual(row["summary"], "saved generated summary")
        payload = self.store.read_payload(self.store.get(self.batch, "07_cot"))
        self.assertEqual(set(payload["row_edits"]), {"2"})
        self.assertEqual(payload["cot"]["2"]["summary"], "saved generated summary")
        rows = convert_release(self.registry, release)["rows"]
        self.assertEqual(sum(item["人工精修步骤数量"] for item in rows), 1)
        self.assertEqual(sum(item["action_box"].get("click", 0) for item in rows), 0)

    def test_automatic_existing_pending_or_stale_edits_still_block_publication(self):
        created = publishing.prepare_manual(self.batch, self.pipeline("manual")["selection"], root=self.root)
        pipeline = self.pipeline()
        baseline = self.records.get("correction_sessions", created["session_id"])
        for changes in ({"pending_review": {"step": {"task_id": "A"}}}, {"stale_tasks": ["A"]}):
            with self.subTest(changes=changes):
                self.records.put("correction_sessions", baseline["session_id"], {**deepcopy(baseline), **changes})
                with self.assertRaises(BatchNotReadyError):
                    self.publish(pipeline)
                self.assertEqual(self.registry.list_releases(), [])
                self.assertEqual(lifecycle(self.batch, self.root)["status"], "active")
                self.assertEqual(len(self.records.list("correction_sessions")), 1)

    def test_manual_unmodified_top1_and_actual_edits_keep_whole_trajectory(self):
        pipeline = self.pipeline("manual")
        manual = publishing.prepare_manual(self.batch, pipeline["selection"], root=self.root)
        pipeline["session_id"] = manual["session_id"]
        session = self.records.get("correction_sessions", manual["session_id"])
        self.assertEqual(len(session["snapshot_payload"]["groups"]), 4)
        selected_group = next(group for group in session["snapshot_payload"]["groups"] if group["meta_task"] == "a-first")
        number = str(selected_group["rows"][1]["excel_row"])
        session["row_edits"][number] = {"actions": '{"action":"click","coordinate":[1,2]}', "original_actions": '{"action":"wait"}'}
        session = self.records.put("correction_sessions", session["session_id"], session, expected_revision=session["storage_revision"])
        pipeline["confirmation"] = publishing.confirm_manual(session["session_id"], expected_revision=session["storage_revision"], root=self.root)
        self.assertEqual(len(pipeline["confirmation"]["group_ids"]), 2)
        # A generated COT revision is legitimate after human approval.
        session["cot"][number] = {"content_tag": "thought_summary", "summary": "generated"}
        self.records.put("correction_sessions", session["session_id"], session, expected_revision=session["storage_revision"])
        release = self.publish(pipeline)
        self.assertEqual(release["step_count"], 6)
        rows = convert_release(self.registry, release)["rows"]
        self.assertEqual(sum(row["人工精修步骤数量"] for row in rows), 1)
        self.assertEqual(sum(row["step数量"] for row in rows), 6)
        self.assertEqual({row["APP"] for row in rows}, {"App A", "App B"})

    def test_manual_can_select_alternate_and_requires_exactly_one(self):
        pipeline = self.pipeline("manual")
        manual = publishing.prepare_manual(self.batch, pipeline["selection"], root=self.root)
        session = self.records.get("correction_sessions", manual["session_id"])
        for group in session["snapshot_payload"]["groups"]:
            if group["task_id"] == "A":
                session["group_exports"][group["group_id"]] = group["meta_task"] == "a-second"
        session = self.records.put("correction_sessions", session["session_id"], session)
        confirmation = publishing.confirm_manual(session["session_id"], root=self.root)
        self.assertEqual(confirmation["selection"]["selected_trajectories"]["A"], "a-second")
        pipeline.update(session_id=session["session_id"], confirmation=confirmation)
        session["row_edits"]["2"] = {"thought": "after confirmation"}
        self.records.put("correction_sessions", session["session_id"], session)
        with self.assertRaises(RevisionConflict):
            self.publish(pipeline)
        self.assertEqual(self.registry.list_releases(), [])
        for group in session["snapshot_payload"]["groups"]:
            session["group_exports"][group["group_id"]] = True
        self.records.put("correction_sessions", session["session_id"], session)
        with self.assertRaisesRegex(ValueError, "只能选择一条"):
            publishing.confirm_manual(session["session_id"], root=self.root)

    def test_deleted_step_remaps_manual_metadata_and_preserves_handwritten_cot(self):
        pipeline = self.pipeline("manual")
        manual = publishing.prepare_manual(self.batch, pipeline["selection"], root=self.root)
        session = self.records.get("correction_sessions", manual["session_id"])
        group = next(g for g in session["snapshot_payload"]["groups"] if g["meta_task"] == "a-first")
        first, second = [str(row["excel_row"]) for row in group["rows"][:2]]
        session["row_edits"][first] = {"deleted": True}
        session["row_edits"][second] = {"sop": "expert SOP", "thought": "handwritten"}
        session["cot"][second] = {"content_tag": "thought_summary", "action_hash": "outdated", "thought": "stale generated"}
        self.records.put("correction_sessions", session["session_id"], session)
        pipeline.update(session_id=session["session_id"], confirmation=publishing.confirm_manual(session["session_id"], root=self.root))
        release = self.publish(pipeline)
        self.assertEqual(release["step_count"], 5)
        converted = convert_release(self.registry, release)["rows"]
        self.assertEqual(sum(row["人工精修步骤数量"] for row in converted), 1)
        book = load_workbook(self.registry.excel_file(release["release_id"], 0)[0])
        headers = [cell.value for cell in book.active[1]]
        self.assertEqual(book.active.cell(2, headers.index("thought") + 1).value, "handwritten")
        book.close()

    def test_pipeline_session_read_preserves_all_candidates_and_revision(self):
        from backend.trajectory_correction import service, draft_store
        pipeline = self.pipeline("manual")
        manual = publishing.prepare_manual(self.batch, pipeline["selection"], root=self.root)
        self.records.put("pipelines", pipeline["pipeline_id"], {**pipeline, "status": "waiting_for_correction", "session_id": manual["session_id"]})
        self.records.put("pipeline_owners", self.batch, {"pipeline_id": pipeline["pipeline_id"]})
        with patch.object(draft_store, "CORRECTION_SESSIONS_DIR", self.root / "sessions"), patch.object(service, "CORRECTION_INPUTS_DIR", self.root / "system" / "trajectory_correction" / "inputs"):
            session = service.get_session(manual["session_id"])
        self.assertEqual(session["group_count"], 4)
        self.assertEqual(session["storage_revision"], manual["storage_revision"])
        again = publishing.prepare_manual(self.batch, pipeline["selection"], root=self.root)
        self.assertEqual(again["storage_revision"], manual["storage_revision"])
        # Ending orchestration must not turn a read into a Top1 reset.
        owner = self.records.get("pipelines", pipeline["pipeline_id"])
        self.records.put("pipelines", owner["pipeline_id"], {**owner, "status": "terminated"})
        saved = self.records.get("correction_sessions", manual["session_id"])
        alternate = next(g for g in saved["snapshot_payload"]["groups"] if g["meta_task"] == "a-second")
        saved["row_edits"][str(alternate["rows"][0]["excel_row"])] = {"thought": "saved alternate"}
        for group in saved["snapshot_payload"]["groups"]:
            if group["task_id"] == "A":
                saved["group_exports"][group["group_id"]] = group["group_id"] == alternate["group_id"]
        saved = self.records.put("correction_sessions", manual["session_id"], saved)
        with patch.object(draft_store, "CORRECTION_SESSIONS_DIR", self.root / "sessions"), patch.object(service, "CORRECTION_INPUTS_DIR", self.root / "system" / "trajectory_correction" / "inputs"):
            after = service.get_session(manual["session_id"])
        self.assertEqual(after["storage_revision"], saved["storage_revision"])
        self.assertEqual(self.records.get("correction_sessions", manual["session_id"])["row_edits"], saved["row_edits"])
        self.assertTrue(publishing.pipeline_session_source_current(saved, root=self.root))
        changed = self.store.read_payload(self.store.get(self.batch, "02_annotation"))
        changed["sheets"]["VLA trajectories"][0]["summary"] = "new source"
        self.store.publish(self.batch, "02_annotation", changed)
        self.assertFalse(publishing.pipeline_session_source_current(saved, root=self.root))

    def test_changed_quality_and_unselected_task_stale_cannot_publish(self):
        pipeline = self.pipeline()
        ref = self.store.get(self.batch, "05_quality")
        changed = self.store.read_payload(ref)
        changed["tasks"][1]["evaluations"]["b-first"]["global_score"] = 3
        self.store.publish(self.batch, "05_quality", changed)
        with self.assertRaises(RevisionConflict):
            self.publish(pipeline)
        pipeline = self.pipeline()
        changed["tasks"] = changed["tasks"][:1]
        self.store.publish(self.batch, "05_quality", changed)
        with self.assertRaises(BatchNotReadyError):
            self.publish(pipeline)
        self.assertEqual(lifecycle(self.batch, self.root)["status"], "active")

    def test_frozen_upstream_is_independent_of_current_artifact_changes(self):
        release = self.publish(self.pipeline())
        expected = convert_release(self.registry, release)
        with patch.object(ArtifactStore, "get", side_effect=AssertionError("must not consult current artifacts")):
            self.assertEqual(convert_release(self.registry, release), expected)

    def test_no_qualifiers_does_not_close_batch(self):
        self.seed({"A": {"a-second": 1}, "B": {"b-first": 2}})
        pipeline = self.pipeline(threshold=5)
        self.assertFalse(pipeline["selection"]["tasks"])
        with self.assertRaisesRegex(ValueError, "空发布"):
            self.publish(pipeline)
        self.assertEqual(lifecycle(self.batch, self.root)["status"], "active")

    def test_missing_quality_is_not_low_score_exclusion(self):
        ref = self.store.get(self.batch, "05_quality")
        quality = self.store.read_payload(ref)
        del quality["tasks"][1]["evaluations"]["b-second"]
        self.store.publish(self.batch, "05_quality", quality)
        with self.assertRaises(BatchNotReadyError):
            self.pipeline()

    def test_concurrent_publication_has_one_release(self):
        pipeline = self.pipeline()
        with ThreadPoolExecutor(max_workers=2) as pool:
            releases = list(pool.map(lambda _: self.publish(pipeline), range(2)))
        self.assertEqual(releases[0]["release_id"], releases[1]["release_id"])
        self.assertEqual(len(self.records.list("dataset_releases")), 1)

    def test_committed_transaction_with_lost_response_keeps_files(self):
        pipeline = self.pipeline()
        original = RecordStore.put_many
        def uncertain_commit(records, entries, **kwargs):
            result = original(records, entries, **kwargs)
            if any(item["namespace"] == "dataset_releases" for item in entries):
                raise OSError("lost committed response")
            return result
        with patch.object(RecordStore, "put_many", uncertain_commit):
            release = self.publish(pipeline)
        self.assertTrue(self.registry.excel_file(release["release_id"], 0)[0].is_file())
        self.assertEqual(len(convert_release(self.registry, release)["rows"]), 1)

    def test_failed_transaction_retry_keeps_batch_active(self):
        pipeline = self.pipeline()
        original = RecordStore.put_many
        def failed_commit(records, entries, **kwargs):
            if any(item["namespace"] == "dataset_releases" for item in entries):
                raise OSError("database unavailable before commit")
            return original(records, entries, **kwargs)
        with patch.object(RecordStore, "put_many", failed_commit):
            with self.assertRaises(OSError):
                self.publish(pipeline)
        self.assertEqual(lifecycle(self.batch, self.root)["status"], "active")
        self.assertEqual(self.registry.list_releases(), [])
        release = self.publish(pipeline)
        self.assertEqual(lifecycle(self.batch, self.root)["release_id"], release["release_id"])


class OrdinaryPublicationRecoveryTests(unittest.TestCase):
    from backend.tests.test_data_publishing import DatasetPublishingTests as Fixtures
    setUp = Fixtures.setUp
    tearDown = Fixtures.tearDown
    add_session = Fixtures.add_session
    registry = Fixtures.registry

    def test_committed_ordinary_release_files_and_session_survive_lost_response(self):
        session = self.add_session()
        registry = self.registry()
        original = registry._records.put_many
        def committed(entries, **kwargs):
            original(entries, **kwargs)
            raise OSError("commit response lost")
        with patch.object(registry._records, "put_many", committed):
            release = registry.create("ordinary publication", [session["session_id"]])
        self.assertTrue(registry.excel_file(release["release_id"], 0)[0].exists())
        self.assertTrue(session["published"])
        self.assertEqual(lifecycle(session["batch_id"], registry.data_root)["status"], "published")

    def test_ordinary_publication_cannot_bypass_pipeline_ownership(self):
        from backend.pipeline_access import PipelineManagedError
        session = self.add_session()
        registry = self.registry()
        records = registry._records
        pipeline = {"pipeline_id": "pipe-managed", "batch_id": session["batch_id"], "status": "running", "mode": "automatic"}
        records.put("pipelines", pipeline["pipeline_id"], pipeline)
        records.put("pipeline_owners", pipeline["batch_id"], {"pipeline_id": pipeline["pipeline_id"]})
        with self.assertRaises(PipelineManagedError):
            registry.create("bypass", [session["session_id"]])
        self.assertEqual(registry.list_releases(), [])


if __name__ == "__main__":
    unittest.main()
