"""Real platform persistence with an in-process collector transport double."""
from __future__ import annotations

import base64
import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from backend import phonefactory_client
from backend.phone_factory import DEFAULT_CONFIG, PhoneFactoryError
from backend.pipelines import PipelineRuntime
from backend.task_generation.collection_batches import payload_digest
from backend.tests import test_phone_factory_runtime as fixtures


def streaming_capabilities(value):
    value.update(collection_artifacts_version=1)
    value.setdefault("agent_types", {})
    for engine in ("newagent", "pipelineagent"):
        value["agent_types"].setdefault(engine, {})
        value["agent_types"][engine].update(available=True, collection=True,
            model_iteration_evaluation=True, capture_logprobs=True,
            collection_artifacts_version=1, repeat_times=True, max_repeat_times=100)
    return value


class CollectionProtocolRuntimeTests(unittest.TestCase):
    setUp = fixtures.FactoryRuntimeTests.setUp
    request = fixtures.FactoryRuntimeTests.request

    def client(self, args):
        result = fixtures.FactoryRuntimeTests.client(self, args)
        value = json.loads(result["output"])
        if args[0] == "capabilities":
            streaming_capabilities(value)
        elif args[0] == "start-run":
            value["agent_type"] = args[args.index("--agent-type") + 1]
            value["batch_id"] = args[args.index("--batch-id") + 1] if "--batch-id" in args else ""
            self.remote[value["run_id"]].update(value, status="running")
        elif args[0] == "cancel-run":
            self.remote[args[1]].update(status="cancelled", collection_status="cancelled")
            value = {"ok": True, **self.remote[args[1]]}
        return {**result, "output": json.dumps(value)}

    def test_repeat_defaults_are_generate_only_and_old_saves_preserve_them(self):
        self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 7})
        self.store.dispatch("config", "POST", DEFAULT_CONFIG)
        self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 1}, mode="modeliter")
        self.assertEqual(self.store.config()["repeat_times"], 7)
        self.assertEqual(self.store.config("modeliter")["repeat_times"], 1)
        for invalid in (True, 0, 101, 1.5, "2", None):
            with self.subTest(invalid=invalid), self.assertRaises(PhoneFactoryError):
                self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": invalid})
        with self.assertRaises(PhoneFactoryError):
            self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 2}, mode="modeliter")

    def test_both_engines_repeat_frozen_dispatch_and_changed_defaults_do_not_change_replay(self):
        self.store.add_task(self.upload)
        for engine in ("newagent", "pipelineagent"):
            request = self.request(request_id=engine, agent_type=engine, repeat_times=3)
            first = self.store.remote_start(request)
            metadata = first["request"]["metadata"]
            self.assertEqual((metadata["repeat_times"], metadata["collection_artifacts_version"]), (3, 1))
            self.assertNotIn("repeat_times", metadata["config"])
            args = [call for call in self.calls if call[0] == "start-run"][-1]
            self.assertEqual(args[args.index("--repeat-times") + 1], "3")
            self.assertEqual(args[args.index("--collection-artifacts-version") + 1], "1")
            self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 9})
            replay = self.store.remote_start(request)
            self.assertTrue(replay["reused"])
            self.assertEqual(replay["request"], first["request"])
            with self.assertRaises(ValueError):
                self.store.remote_start({**request, "repeat_times": 4})
        self.assertEqual(sum(call[0] == "start-run" for call in self.calls), 2)

    def test_modeliter_shares_run_with_real_batch_and_preserves_upload(self):
        book = load_workbook(io.BytesIO(fixtures.workbook_bytes()))
        book.active.insert_cols(1)
        book.active.cell(1, 1, "priority")
        book.active.cell(2, 1, "must retain")
        output = io.BytesIO(); book.save(output); book.close()
        content = output.getvalue()
        self.store.add_task({**self.upload, "content_base64": base64.b64encode(content).decode()}, mode="modeliter")
        run = self.store.remote_start(self.request(repeat_times=1), mode="modeliter")
        collection = self.store.collection_runs.get(run["run_id"])
        self.assertEqual(run["batch_id"], collection["batch_id"])
        batch = self.store.collection_runs.batch(run["batch_id"], require_workbook=True)
        self.assertEqual(batch["collection_run_id"], run["run_id"])
        original = collection["request"]["metadata"]["original_workbook"]
        self.assertEqual((self.root / original["relative_path"]).read_bytes(), content)
        args = [call for call in self.calls if call[0] == "start-run"][-1]
        self.assertEqual(Path(args[1]).read_bytes(), content)
        self.assertEqual(args[args.index("--batch-id") + 1], run["batch_id"])
        self.assertEqual(args[args.index("--repeat-times") + 1], "1")
        self.assertEqual(self.store.runtime.run_views("generate"), [])
        self.assertEqual(self.store.runtime.run_views("modeliter")[0]["batch_id"], run["batch_id"])
        self.assertEqual(self.store.remote_start(self.request(repeat_times=1), mode="modeliter")["run_id"], run["run_id"])
        self.assertEqual(sum(call[0] == "start-run" for call in self.calls), 1)
        with self.assertRaisesRegex(PhoneFactoryError, "一条轨迹"):
            self.store.remote_start(self.request(request_id="wrong", repeat_times=2), mode="modeliter")
        with self.assertRaises(PhoneFactoryError):
            self.store.runtime.remote_start_evaluation({"repeat_times": 1})

    def test_old_collector_rejected_before_phone_dispatch_or_new_result_batch(self):
        self.store.add_task(self.upload)
        self.store.run_client = lambda args: {"ok": True, "output": json.dumps(
            {"protocol_version": 1, "batch_results": True, "agent_types": {"newagent": {"available": True, "collection": True}}})}
        with self.assertRaisesRegex(PhoneFactoryError, "逐条轨迹回传"):
            self.store.remote_start(self.request())
        self.assertEqual(self.store.collection_runs.list_runs(), [])
        self.assertEqual(self.store.collection_runs.result_batches.list(), [])

    def test_legacy_generate_replays_without_upgrade_or_new_defaults(self):
        self.store.add_task(self.upload)
        first = self.store.runtime.remote_start(self.request(), "generate", legacy_batch=True)
        def make_legacy(value):
            value["request"]["metadata"].pop("repeat_times", None)
            value["request_digest"] = payload_digest(value["request"])
        frozen = self.store.records.update("collection_runs", first["run_id"], make_legacy)
        before = len(self.calls)
        self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 8})
        replay = self.store.remote_start(self.request(repeat_times=1))
        self.assertEqual(replay["request"], frozen["request"])
        self.assertEqual(len(self.calls), before)
        with self.assertRaises(ValueError):
            self.store.remote_start(self.request(repeat_times=2))

    def test_historical_modeliter_replay_does_not_create_collection_data(self):
        self.store.add_task(self.upload, mode="modeliter")
        request = self.request()
        options = self.store.runtime._options(request, "modeliter")
        options.pop("repeat_times")
        import hashlib
        frozen = {**options, "filename": "tasks.xlsx", "workbook_sha256": hashlib.sha256(
            self.store.task_path("tasks.xlsx").read_bytes()).hexdigest()}
        run_id = "eval_" + payload_digest(request["request_id"])
        self.store.records.put("model_iter_runs", run_id, {"run_id": run_id, "collection_run_id": run_id,
            "status": "running", "created_at": "2026-10-30T00:00:00Z", "request": frozen,
            "request_digest": payload_digest(frozen), "dispatch_error": None})
        with patch.object(self.store, "run_client", side_effect=AssertionError("historical replay contacted collector")):
            replay = self.store.remote_start({**request, "repeat_times": 1}, mode="modeliter")
        self.assertTrue(replay["reused"])
        self.assertEqual(replay["request"], frozen)
        self.assertEqual(self.store.collection_runs.list_runs(), [])
        self.assertEqual(self.store.collection_runs.result_batches.list(), [])

    def test_existing_result_batch_missing_repeat_accepts_one_without_rewriting(self):
        task = self.store.add_task(self.upload)["imported_task"]
        options = self.store.runtime._options(self.request(), "generate")
        options.pop("repeat_times")
        old, _ = self.store.collection_runs.create_result_run(task["source_batch_id"],
            dispatch_key="request-1", metadata={**options, "filename": task["filename"]},
            workbook_sha256=task["content_sha256"])
        self.store.collection_runs.dispatched(old["run_id"], {"ok": True})
        snapshot = self.store.collection_runs.get(old["run_id"])
        replay = self.store.remote_start(self.request(repeat_times=1))
        self.assertTrue(replay["reused"])
        self.assertEqual(self.store.collection_runs.get(old["run_id"]), snapshot)

    def test_stop_keeps_collection_transfer_alive_and_modeliter_status_separate(self):
        self.store.add_task(self.upload)
        run = self.store.remote_start(self.request())
        stopped = self.store.runtime.cancel_collection(run["run_id"])
        self.assertEqual((stopped["status"], stopped["collection_status"]), ("running", "cancelled"))
        self.assertNotEqual(self.store.runtime.sync_policy.view("collection", run["run_id"])["state"], "completed")
        self.assertEqual(self.store.runtime.cancel_collection(run["run_id"])["stop_requested_at"], stopped["stop_requested_at"])

    def test_modeliter_retry_transfer_and_listing_do_not_overwrite_evaluation(self):
        self.store.add_task(self.upload, mode="modeliter")
        run = self.store.remote_start(self.request(), mode="modeliter")
        self.store.records.update("model_iter_runs", run["run_id"], lambda item: item.update(status="completed"))
        self.store.records.update("collection_runs", run["run_id"], lambda item: item.update(
            collection_status="running", transfer_status="failed", transferred_trajectory_count=1))
        with patch.object(self.store.runtime.transfer, "sync", return_value={}) as sync:
            view = self.store.dispatch("runs/" + run["run_id"] + "/retry-transfer", "POST", {}, mode="modeliter")
        sync.assert_called_once_with(run["run_id"], force=True)
        self.assertEqual(view["status"], "completed")
        self.assertEqual(view["collection_status"], "running")
        self.assertEqual(view["collection_summary"]["transferred_trajectory_count"], 1)

    def test_modeliter_cancel_does_not_cancel_transfer_and_list_retains_published_permissions(self):
        self.store.add_task(self.upload, mode="modeliter")
        run = self.store.remote_start(self.request(), mode="modeliter")
        self.store.dispatch("runs/" + run["run_id"] + "/cancel", "POST", {}, mode="modeliter")
        collection = self.store.collection_runs.get(run["run_id"])
        self.assertEqual((collection["status"], collection["collection_status"]), ("running", "cancelled"))
        self.store.records.put("batch_lifecycle", run["batch_id"], {
            "batch_id": run["batch_id"], "status": "published", "release_id": "release-one", "published_at": "yesterday"})
        view = self.store.runtime.run_views("modeliter")[0]
        self.assertEqual((view["batch_status"], view["release_id"]), ("published", "release-one"))

    def test_pipeline_defaults_freeze_once(self):
        runtime = PipelineRuntime(self.root, preprocessing=None, tree=None, quality=None,
                                  cot=None, factory=self.store, overview=None)
        self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 4})
        options = runtime.collection_options({"vla": "http://mock.invalid/vla"})
        self.assertEqual((options["repeat_times"], options["collection_artifacts_version"]), (4, 1))
        self.store.dispatch("config", "POST", {**DEFAULT_CONFIG, "repeat_times": 8})
        self.assertEqual(runtime.collection_options(options, new_run=False), options)
        legacy = {key: value for key, value in options.items() if key not in {"repeat_times", "collection_artifacts_version"}}
        self.assertEqual(runtime.collection_options(legacy, new_run=False), legacy)

    def test_cli_artifact_commands_quote_ids_and_forward_options(self):
        for command, expected in (
            (["collection-artifacts", "run-a", "--cursor", "4"], "/runs/run-a/collection-artifacts?cursor=4&limit=100"),
            (["collection-artifact", "run-a", "artifact-a"], "/runs/run-a/collection-artifacts/artifact-a"),
            (["collection-manifest", "run-a"], "/runs/run-a/collection-manifest"),
        ):
            with self.subTest(command=command), patch("sys.argv", ["client", *command]), \
                    patch.object(phonefactory_client, "get_json", return_value={}) as get, patch("builtins.print"):
                phonefactory_client.main()
            get.assert_called_once_with(expected)
        for command, expected in (
            (["collection-artifact-archive", "run-a", "artifact-a", "--out", "file.zip"],
             "/runs/run-a/collection-artifacts/artifact-a/archive"),
            (["collection-archive", "run-a", "--out", "file.zip"], "/runs/run-a/collection-archive"),
        ):
            with self.subTest(command=command), patch("sys.argv", ["client", *command]), \
                    patch.object(phonefactory_client, "download_file", return_value={}) as download, patch("builtins.print"):
                phonefactory_client.main()
            download.assert_called_once_with(expected, "file.zip")


if __name__ == "__main__":
    unittest.main()
