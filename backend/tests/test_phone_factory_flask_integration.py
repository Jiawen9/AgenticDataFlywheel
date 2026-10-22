"""Offline cross-repository check of the deployed Flask collector contract.

Set PHONE_FACTORY_SNAPSHOT_ROOT to a reviewed collector snapshot containing
phonefactory_manager.py and new-guigent/logprobs_export.py. Without that variable
this optional suite is explicitly skipped. Flask and pandas must be installed in
the test runtime when the snapshot is selected. No server or device is started.
"""
from __future__ import annotations

import ast
import base64
import contextlib
import hashlib
import importlib.util
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
import urllib.error
from functools import partial
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook

from backend import phonefactory_client, pipeline_release
from backend.batch_results import current_tree_payload, merge_quality_results, quality_task_fingerprints
from backend.data_publishing.service import DatasetReleaseRegistry
from backend.data_store import ArtifactStore
from backend.data_store.registry import utc_now
from backend.phone_factory import DEFAULT_CONFIG, PhoneFactoryStore
from backend.preprocessing_jobs import PreprocessingJobManager
from backend.quality_input_builder import build_quality_workbook
from backend.task_generation.collection_batches import COLLECTION_COLUMNS
from backend.tests.test_preprocessing_jobs import CONFIG, ManualQueue
from backend.tests.test_trajectories_preprocessing import AcceptingReviewer, create_step
from backend.tests.test_tree_build_service import FakeAlignmentReviewer, FakeClassifier, FakeSummarizer
from backend.tree_build_service import build_tree_run


SNAPSHOT_ROOT = os.environ.get('PHONE_FACTORY_SNAPSHOT_ROOT')
SHEET = 'VLA trajectories'
BASE_COLUMNS = ['文件夹名', 'image', 'xml', 'action', 'summary', 'logprobs_path']


@unittest.skipUnless(SNAPSHOT_ROOT, 'set PHONE_FACTORY_SNAPSHOT_ROOT for actual Flask collector integration')
class FlaskCollectorIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.stack = contextlib.ExitStack()
        self.addCleanup(self.stack.close)
        self.base = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        snapshot = Path(SNAPSHOT_ROOT)
        manager_path = snapshot / 'phonefactory_manager.py'
        helper_path = snapshot / 'new-guigent' / 'logprobs_export.py'
        self.assertTrue(manager_path.is_file(), str(manager_path))
        self.assertTrue(helper_path.is_file(), str(helper_path))
        spec = importlib.util.spec_from_file_location('phone_factory_flask_snapshot', manager_path)
        self.manager = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.manager)
        self.collector_root = self.base / 'collector'
        template = self.collector_root / 'new-guigent'
        template.mkdir(parents=True)
        shutil.copy2(helper_path, template / helper_path.name)
        for key, value in {'BASE_DIR': self.collector_root, 'TEMPLATE_DIR': template,
                           'HTTP_RUNS_DIR': self.collector_root / '.http_runs',
                           'HTTP_ARCHIVES_DIR': self.collector_root / '.http_archives'}.items():
            self.stack.enter_context(patch.object(self.manager, key, value.as_posix()))
        self.launch = self.stack.enter_context(patch.object(self.manager, 'launch_onephonerun'))
        self.stack.enter_context(patch.object(self.manager, '_watch_http_run'))
        self.remote = self.manager.app.test_client()
        self.calls = []
        self.factory = PhoneFactoryStore(self.base / 'platform', run_client_fn=self.client)
        self.factory.initialize_settings({'phones': ['phone-a', 'phone-b'], 'apps': ['AppA', 'AppB'],
            'phoneApps': [{'phone_id': 'phone-a', 'app': 'AppA'}, {'phone_id': 'phone-b', 'app': 'AppB'}],
            'vla': ['model.invalid:8000']})
        book = Workbook()
        book.active.append(COLLECTION_COLUMNS)
        for case, app in [('case-a', 'AppA'), ('case-b', 'AppB')]:
            book.active.append([case, '单APP', app, '购物', '搜索', None, '搜索商品',
                                None, None, 2, '人工', None, None, None, None, None, None])
        data = io.BytesIO()
        book.save(data)
        book.close()
        task = self.factory.add_task({'filename': 'flask.xlsx', 'description': 'Flask integration',
            'content_base64': base64.b64encode(data.getvalue()).decode()})['imported_task']
        self.filename, self.batch_id = task['filename'], task['source_batch_id']
        self.saved = {**DEFAULT_CONFIG, 'capture_logprobs': True,
                      'sampling_enabled': True, 'temperature': 0.35, 'top_p': 0.75}
        self.factory.dispatch('config', 'POST', self.saved)
        self.stack.enter_context(patch('backend.trajectories_preprocessing.read_env_file', return_value={}))

    def urlopen(self, request, timeout=None):
        url = urlsplit(request.full_url)
        response = self.remote.open(url.path + ('?' + url.query if url.query else ''),
            method=request.get_method(), data=request.data, headers=dict(request.header_items()))
        content, headers, status = response.get_data(), response.headers, response.status_code
        response.close()
        if status >= 400:
            raise urllib.error.HTTPError(request.full_url, status, content.decode(), headers, io.BytesIO(content))
        stream = io.BytesIO(content)
        stream.headers = headers
        return stream

    def client(self, args):
        self.calls.append(list(args))
        output = io.StringIO()
        with patch.object(sys, 'argv', ['phonefactory_client', *args]), \
                patch.object(phonefactory_client.urllib.request, 'urlopen', self.urlopen), \
                contextlib.redirect_stdout(output):
            try:
                phonefactory_client.main()
            except SystemExit as exc:
                return {'ok': False, 'output': str(exc)}
        return {'ok': True, 'output': output.getvalue()}

    def start(self, request_id, *, vla='model.invalid:8000', capture=None):
        params = {'filename': self.filename, 'request_id': request_id, 'vla': vla}
        if capture is not None:
            params['config'] = {'capture_logprobs': capture}
        run = self.factory.remote_start(params)
        record = self.manager._load_http_run(run['collection_run_id'])
        self.assertEqual(record['params']['temperature'], 0.35)
        self.assertEqual(record['params']['top_p'], 0.75)
        for phone in record['phone_ids']:
            script = (self.collector_root / ('new-guigent-' + phone) / 'onephonerun.py').read_text(encoding='utf-8')
            env = {}
            for node in ast.walk(ast.parse(script)):
                if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Subscript):
                    target = node.targets[0]
                    if ast.unparse(target.value) == 'os.environ' and isinstance(target.slice, ast.Constant) and isinstance(node.value, ast.Constant):
                        env[target.slice.value] = node.value.value
            self.assertEqual(env['GUI_AGENT_VLA_API_BASE_URL'], record['params']['vla'])
            self.assertEqual(env['GUI_AGENT_CORE_VLA_TEMPERATURE'], '0.35')
            self.assertEqual(env['GUI_AGENT_CORE_VLA_TOP_P'], '0.75')
        return run, record

    def populate(self, record, traces):
        responses = {}
        for phone, cases in record['phone_cases'].items():
            for case in cases:
                folder = Path(record['result_dir']) / case
                folder.mkdir(parents=True, exist_ok=True)
                action = {'action': 'click', 'coordinate': [50, 50]}
                create_step(folder, 1, action)
                response = folder / 'step001_vla_model_response.json'
                payload = json.loads(response.read_text(encoding='utf-8'))
                if traces.get(phone) is not None:
                    payload['behavior_trace'] = traces[phone]
                    response.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
                (folder / '_trajectory_for_evaluate.json').write_text(
                    json.dumps({'actions_flat': [{'global_step': 1, 'action': action}]}), encoding='utf-8')
                responses[response] = response.read_bytes()
        return responses

    def finish(self, run, expected_status, responses):
        record = self.manager._finalize_http_run(run['collection_run_id'])
        self.assertEqual(record['status'], expected_status, record.get('errors'))
        archive = Path(record['archive_path']).read_bytes()
        self.assertEqual(record['manifest']['archive'], {'size': len(archive), 'sha256': hashlib.sha256(archive).hexdigest()})
        for path, original in responses.items():
            self.assertEqual(path.read_bytes(), original)
        completed = self.factory.runtime.transfer.sync(run['collection_run_id'])
        self.assertEqual(completed['status'], 'completed')
        self.assertEqual(completed['remote_status'], expected_status)
        self.assertEqual(completed['transfer_status'], 'completed')
        self.assertEqual(len(completed['trajectories']), 2)
        self.assertEqual(len({item['relative_dir'] for item in completed['trajectories']}), 2)
        return completed

    def preprocess(self):
        queue = ManualQueue()
        manager = PreprocessingJobManager(self.factory.root, source_store=self.factory.collection_runs,
            executor=queue, config_loader=lambda: dict(CONFIG), reviewer_factory=lambda **_: AcceptingReviewer())
        job = manager.submit(self.batch_id)
        queue.drain()
        job = manager.get(job['job_id'])
        self.assertEqual(job['status'], 'succeeded', job.get('error'))
        store = ArtifactStore(self.factory.root)
        payloads = [store.read_payload(ref) for ref in job['artifacts']]
        for index, payload in enumerate(payloads):
            self.assertEqual(payload['columns'][SHEET], BASE_COLUMNS + (['actions_box'] if index else []))
            self.assertEqual(len(payload['sheets'][SHEET]), 2)
        self.assertEqual([row['logprobs_path'] for row in payloads[0]['sheets'][SHEET]],
                         [row['logprobs_path'] for row in payloads[1]['sheets'][SHEET]])
        return job, payloads[-1]

    def test_enabled_flask_archive_preprocessing_and_publication(self):
        run, record = self.start('enabled')
        self.assertTrue(record['params']['capture_logprobs'])
        self.assertEqual(record['params']['vla'], 'http://model.invalid:8000/v1/')
        traces = {phone: {'status': 'ok', 'response_token_ids': [101 + index],
            'behavior_logprobs': [-0.25], 'response_loss_mask': [1],
            'sampling': {'temperature': 0.35, 'top_p': 0.75}}
            for index, phone in enumerate(('phone-a', 'phone-b'))}
        completed = self.finish(run, 'completed', self.populate(record, traces))
        self.assertFalse(completed['errors'])
        self.assertEqual(self.launch.call_count, 2)
        self.assertTrue(self.factory.remote_start({'filename': self.filename,
            'request_id': 'enabled', 'vla': 'model.invalid:8000'})['reused'])
        self.assertEqual(self.launch.call_count, 2)
        job, table = self.preprocess()
        rows = table['sheets'][SHEET]
        paths = {row['logprobs_path'] for row in rows}
        self.assertEqual(len(paths), 2)
        for row in rows:
            reference = Path(row['logprobs_path'])
            self.assertEqual(reference.name, 'step001_logprobs.json')
            self.assertEqual(reference.parent, Path(row['image']).parent)
            source = Path(table['raw_root']) / reference
            trace = json.loads(source.read_text(encoding='utf-8'))
            response = json.loads(source.with_name('step001_vla_model_response.json').read_text(encoding='utf-8'))
            self.assertEqual(trace, response['behavior_trace'])
            self.assertIn(trace, traces.values())

        root = self.factory.root
        task_ids = sorted({row['task_id'] for row in rows})
        configuration = {'model': 'offline-tree', 'version': 1}
        with patch('backend.tree_build_service.tree_build_config', return_value=configuration), \
             patch('backend.tree_build_service.configure_reviewer_environment', return_value='fake-model'), \
             patch('backend.tree_build_service.QwenIntermediateStateClassifier', FakeClassifier), \
             patch('backend.tree_build_service.QwenStateAlignmentReviewer', FakeAlignmentReviewer):
            build_tree_run(task_ids, job_id='flask-tree', progress=lambda _: None, batch_id=self.batch_id,
                annotation_version=job['annotation_version'], data_root=root, env_path=self.base / 'no-network.env',
                quality_builder=partial(build_quality_workbook, summarizer=FakeSummarizer()))
        tree = current_tree_payload(self.batch_id, root)
        self.assertEqual({row['logprobs_path'] for row in tree['source_annotation']['sheets'][SHEET]}, paths)
        evaluations = [{'task_id': task, 'status': 'succeeded', 'evaluations': {
            row['trajectory_id']: {'global_score': 5, 'passed_threshold': True}
            for row in rows if row['task_id'] == task}} for task in task_ids]
        merge_quality_results(self.batch_id, evaluations, tree['tree_hashes'], quality_task_fingerprints(tree),
            job_id='flask-quality', completed_at=utc_now(), root=root)
        selection = pipeline_release.build_selection(self.batch_id, task_ids, 'automatic', 4, root=root)
        registry = DatasetReleaseRegistry(data_root=root, releases_file=root / 'registry.json')
        release = pipeline_release.publish_pipeline({'pipeline_id': 'flask-offline', 'name': 'Flask integration',
            'batch_id': self.batch_id, 'mode': 'automatic', 'selection': selection}, root=root, registry=registry)
        workbook = load_workbook(registry.excel_file(release['release_id'], 0)[0], read_only=True)
        try:
            exported = list(workbook.active.values)
            header = list(exported[0])
            self.assertEqual(header.count('logprobs_path'), 1)
            self.assertEqual({row[header.index('logprobs_path')] for row in exported[1:]}, paths)
        finally:
            workbook.close()

    def test_disabled_override_and_legacy_default_keep_ordinary_trajectories(self):
        run, record = self.start('disabled', vla='https://model.invalid/custom/v1', capture=False)
        self.assertFalse(record['params']['capture_logprobs'])
        self.assertEqual(record['params']['vla'], 'https://model.invalid/custom/v1/')
        self.assertTrue(self.factory.config()['capture_logprobs'])
        self.finish(run, 'completed', self.populate(record, {}))
        _, table = self.preprocess()
        self.assertTrue(all(row['logprobs_path'] == '' for row in table['sheets'][SHEET]))
        def legacy_settings(state):
            state['config'].pop('capture_logprobs')
        self.factory._update(legacy_settings)
        _, legacy = self.start('legacy-default')
        self.assertFalse(legacy['params']['capture_logprobs'])
        self.assertNotIn('capture_logprobs', self.factory._current()['config'])

    def test_missing_and_invalid_logprobs_keep_both_trajectories_and_report_steps(self):
        self.factory.dispatch('config', 'POST', {**self.saved, 'capture_logprobs': False})
        run, record = self.start('partial', capture=True)
        invalid = {'status': 'invalid', 'error': 'test invalid trace', 'response_token_ids': [202],
                   'behavior_logprobs': [], 'response_loss_mask': [1]}
        completed = self.finish(run, 'partial', self.populate(record, {'phone-b': invalid}))
        self.assertFalse(self.factory.config()['capture_logprobs'])
        self.assertEqual(self.factory.state()['tasks'][0]['status'], '部分完成')
        self.assertEqual(len(completed['errors']), 2)
        self.assertTrue(all('step001' in item['error'] for item in completed['errors']))
        _, table = self.preprocess()
        references = [row['logprobs_path'] for row in table['sheets'][SHEET]]
        self.assertEqual(sum(bool(path) for path in references), 1)
        source = Path(table['raw_root']) / next(path for path in references if path)
        self.assertEqual(json.loads(source.read_text(encoding='utf-8')), invalid)


if __name__ == '__main__':
    unittest.main()
