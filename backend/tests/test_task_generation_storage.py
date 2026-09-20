from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from backend.task_generation.collection_input import CollectionInputError
from backend.task_generation.jobs import TaskGenerationJobManager


class NoExecution:
    def submit(self, *args, **kwargs):
        raise AssertionError("存储操作不得重新调用模型")


class TaskGenerationStorageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.manager = self.make_manager()
        self.job = {"job_id": "generation-1", "kind": "task_generation", "status": "succeeded",
                    "stage": "succeeded", "errors": [], "warnings": [], "knowledge_base_version": "kb-1"}
        self.rows = [
            {"result_id": f"result-{i}", "task_uuid": f"task-{i}", "task": f"任务 {i}",
             "app": "App", "scene": "L1", "capability": "L2", "sub_capability": "L3",
             "pre_dependency": "zero", "deleted": False} for i in range(12)]
        self.manager._write_job(self.job)
        self.manager._write_json(self.manager._results_path("generation-1"), self.rows)

    def make_manager(self):
        manager = TaskGenerationJobManager(
            self.base / "jobs", self.base / "runs", self.base / "exports",
            self.base / "kb", self.base / "logs", executor=NoExecution(), data_root=self.base / "data")
        self.addCleanup(manager.shutdown)
        return manager

    def test_sqlite_is_authoritative_after_missing_or_modified_mirrors_and_restart(self):
        self.assertFalse(self.manager._results_path("generation-1").exists())
        self.assertFalse(self.manager._job_path("generation-1").exists())
        self.manager._results_path("generation-1").parent.mkdir(parents=True, exist_ok=True)
        self.manager._results_path("generation-1").write_text("broken JSON", encoding="utf-8")
        restarted = self.make_manager()
        self.assertEqual(restarted.results("generation-1"), self.rows)
        self.assertEqual(restarted.collection_input("generation-1")["task_count"], 12)
        self.assertEqual(self.manager._results_path("generation-1").read_text(encoding="utf-8"), "broken JSON")

    def test_transactional_edits_from_independent_managers_keep_all_rows(self):
        managers = [self.manager, self.make_manager(), self.make_manager()]
        with ThreadPoolExecutor(max_workers=6) as executor:
            futures = [executor.submit(managers[i % 3].patch_result, "generation-1",
                                       f"result-{i}", {"task": f"人工任务 {i}"}) for i in range(12)]
            for future in futures:
                future.result()
        self.assertEqual([row["task"] for row in self.manager.results("generation-1")],
                         [f"人工任务 {i}" for i in range(12)])

    def test_current_snapshot_replaces_fixed_files_and_preserves_excel_literal_text(self):
        first = self.manager.snapshot("generation-1")
        first_json = self.manager.artifacts.resolve_file(first, "result.json")
        original_bytes = first_json.read_bytes()
        self.manager.patch_result("generation-1", "result-0", {"task": "=原文不能变公式"})
        second = self.manager.snapshot("generation-1")
        self.assertNotEqual(first["version"], second["version"])
        self.assertEqual(self.manager.artifacts.resolve_file(second, "result.json"), first_json)
        self.assertNotEqual(first_json.read_bytes(), original_bytes)
        self.assertEqual(len(self.manager.artifacts.list(first["batch_id"], first["stage"])), 1)
        self.assertIsNone(self.manager.artifacts.get(first["batch_id"], first["stage"], first["version"]))
        with self.assertRaises(FileNotFoundError):
            self.manager.artifacts.resolve_file(first, "result.json")
        payload = json.loads(self.manager.artifacts.resolve_file(second, "result.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["results"][0]["task"], "=原文不能变公式")
        workbook = load_workbook(self.manager.artifacts.resolve_file(second, "result.xlsx"), data_only=False)
        try:
            sheet = workbook["任务"]
            headers = [cell.value for cell in sheet[1]]
            cell = sheet.cell(2, headers.index("task") + 1)
            self.assertEqual(cell.value, payload["results"][0]["task"])
            self.assertEqual(cell.data_type, "s")
        finally:
            workbook.close()

    def test_failed_snapshot_keeps_saved_results_and_can_retry_without_model(self):
        with patch.object(self.manager.artifacts, "publish", side_effect=OSError("磁盘不可用")):
            self.manager._save_snapshot_or_warning("generation-1")
        self.assertIn("磁盘不可用", self.manager.get("generation-1")["artifact_error"])
        self.assertEqual(self.manager.results("generation-1"), self.rows)
        self.manager.snapshot("generation-1")
        self.assertIsNone(self.manager.get("generation-1")["artifact_error"])

    def test_collection_artifact_failure_leaves_no_selectable_batch(self):
        with patch.object(self.manager.artifacts, "publish", side_effect=sqlite3.OperationalError("database locked")):
            with self.assertRaises(CollectionInputError):
                self.manager.submit_collection_batch("generation-1")
        self.assertEqual(self.manager.collection_batches(), [])
        self.assertFalse(self.manager._collection_batch_dir("generation-1").exists())
        batch, created = self.manager.submit_collection_batch("generation-1")
        self.assertTrue(created)
        self.assertEqual(batch["task_count"], 12)
        manifests = self.manager.artifacts.list(batch_id="generation-1", stage="00_collection")
        self.assertEqual(len(manifests), 1)

    def test_same_backend_managers_submit_one_collection_batch(self):
        other = self.make_manager()
        with ThreadPoolExecutor(max_workers=2) as executor:
            futures = [executor.submit(manager.submit_collection_batch, "generation-1")
                       for manager in (self.manager, other)]
            results = [future.result() for future in futures]
        self.assertEqual(sorted(created for _, created in results), [False, True])
        self.assertEqual(results[0][0], results[1][0])
        self.assertEqual(len(self.manager.artifacts.list(batch_id="generation-1", stage="00_collection")), 1)

    def test_malformed_legacy_job_does_not_prevent_startup(self):
        (self.base / "jobs" / "broken.json").write_text("{}", encoding="utf-8")
        restarted = self.make_manager()
        self.assertEqual([item["job_id"] for item in restarted.list_jobs()], ["generation-1"])

    def test_augmentation_preview_reads_saved_seeds_without_json_mirror(self):
        self.manager._write_job({**self.job, "kind": "augmentation", "augmentation_preview_version": 1})
        seeds = [{"seed_id": "seed-1", "task": "失败任务", "app": "App",
                  "classification_status": "classified", "generation_status": "waiting"}]
        self.manager._write_json(self.manager._seeds_path("generation-1"), seeds)
        self.assertFalse(self.manager._seeds_path("generation-1").exists())
        preview = self.manager.augmentation_preview("generation-1", include_tree=False)
        self.assertTrue(preview["available"])
        self.assertEqual(preview["seeds"], seeds)

    def test_old_frozen_library_is_not_automatically_imported(self):
        legacy = self.base / "legacy" / "task_generation"
        library = legacy / "runs" / "old-job" / "KnowledgeBase"
        library.mkdir(parents=True)
        (library / "scene.json").write_text('{"version":"frozen"}', encoding="utf-8")
        with self.assertRaises(FileNotFoundError):
            self.manager._execution_library("old-job")
        self.assertFalse((self.base / "runs" / "old-job").exists())
        self.assertFalse((library.parent / "model_calls").exists())

    def test_old_jobs_and_json_mirrors_are_invisible_and_not_modified(self):
        legacy = self.base / "legacy" / "task_generation"
        (legacy / "jobs").mkdir(parents=True)
        (legacy / "runs" / "old-job").mkdir(parents=True)
        old_job = legacy / "jobs" / "old-job.json"
        old_rows = legacy / "runs" / "old-job" / "results.json"
        old_job.write_text(json.dumps({**self.job, "job_id": "old-job"}), encoding="utf-8")
        old_rows.write_text(json.dumps(self.rows), encoding="utf-8")
        before = (old_job.read_bytes(), old_rows.read_bytes())
        (self.base / "jobs" / "old-job.json").write_bytes(old_job.read_bytes())
        self.assertEqual({job["job_id"] for job in self.manager.list_jobs()}, {"generation-1"})
        self.assertIsNone(self.manager.get("old-job"))
        with self.assertRaises(FileNotFoundError):
            self.manager.collection_input("old-job")
        self.assertIsNone(self.manager.store.get("task_generation.results", "old-job"))
        with self.assertRaises(FileNotFoundError):
            self.manager.patch_result("old-job", "result-0", {"task": "不应被写入"})
        self.assertEqual((old_job.read_bytes(), old_rows.read_bytes()), before)


if __name__ == "__main__":
    unittest.main()
