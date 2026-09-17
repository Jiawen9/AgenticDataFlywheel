from __future__ import annotations

from concurrent.futures import Future, ThreadPoolExecutor
from copy import deepcopy
from functools import partial
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image

from backend.data_store import ArtifactStore
from backend.stage_artifacts import write_payload_workbook
from backend.trajectory_context import AnnotationVersionConflict, resolve_batch_context
from backend.trajectory_data import load_annotated_trajectories, resolve_image_asset, task_summaries, update_action_bbox
from backend.tree_build_jobs import TreeBuildJobManager
from backend.tree_build_service import build_tree_run
from backend.quality_input_builder import build_quality_workbook
from backend.tests.test_tree_build_service import FakeClassifier, FakeAlignmentReviewer, FakeSummarizer


COLUMNS = ["文件夹名", "image", "xml", "action", "summary", "actions_box"]


class DeferredExecutor:
    def __init__(self):
        self.work = []

    def submit(self, fn, *args):
        future = Future()
        self.work.append((fn, args, future))
        return future

    def finish(self):
        for fn, args, future in self.work:
            try:
                future.set_result(fn(*args))
            except Exception as exc:
                future.set_exception(exc)


class TrajectoryBatchContextTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "data"
        self.store = ArtifactStore(self.root)

    def make_batch(self, batch="batch-a", *, historical=False):
        raw = self.root / "raw" / ("rollout_trajectories" if historical else f"collection_batches/{batch}")
        rows = []
        for number in (1, 2):
            run = f"run-{number}"
            original = "TASK-1" if historical else "repeated-original"
            directory = Path("TASK") / f"TASK-{number}" if historical else Path("runs") / run / "case-number" / original
            target = raw / directory
            target.mkdir(parents=True)
            Image.new("RGB", (100, 200), "white").save(target / "step001_vla_input.jpg")
            (target / "step001_vla_input_ui.xml").write_text("<hierarchy/>", encoding="utf-8")
            row = {"文件夹名": f"TASK-{number}" if historical else original,
                   "image": (directory / "step001_vla_input.jpg").as_posix(),
                   "xml": (directory / "step001_vla_input_ui.xml").as_posix(),
                   "action": '{"action":"click","coordinate":[10,20]}', "summary": batch,
                   "actions_box": "click(bbox=<bbox>[1,2,30,40]</bbox>)"}
            if not historical:
                identity = hashlib.sha256(json.dumps([batch, run, "TASK", f"case-number/{original}"],
                                                     ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
                row.update(task_id="TASK", trajectory_id="tr_" + identity, source_trajectory_id=original,
                           collection_run_id=run, collected_at=f"2026-09-15T00:00:0{number}Z")
            rows.append(row)
        payload = {"schema_version": 1, "columns": {"VLA trajectories": COLUMNS.copy()},
                   "sheets": {"VLA trajectories": rows}}
        conversion = self.store.publish(batch, "01_conversion", payload,
                                        source_refs=[{"kind": "raw_trajectories", "path": str(raw)}])
        if not historical:
            collection = self.store.publish(batch, "00_collection", {"schema_version": 1,
                "snapshot": {"tasks": [{"task_id": "TASK", "collection_case_id": "case-number", "task": "Frozen goal " + batch}]}})
            conversion = self.store.publish(batch, "01_conversion", payload,
                source_refs=[collection, {"kind": "raw_trajectories", "path": str(raw)}])
        workbook = self.root / f"{batch}.xlsx"
        write_payload_workbook(workbook, payload)
        annotation = self.store.publish(batch, "02_annotation", payload, workbooks={workbook.name: workbook},
                                        source_refs=[conversion], metadata={"raw_root": str(raw)})
        return payload, annotation, raw

    def edit(self, batch, annotation, payload, **kwargs):
        row = payload["sheets"]["VLA trajectories"][0]
        return update_action_bbox("TASK", row["trajectory_id"], 1, 2, (5, 6, 40, 50),
            batch_id=batch, expected_annotation_version=annotation["version"], data_root=self.root, **kwargs)

    def test_two_batches_and_repeated_source_ids_remain_separate(self):
        _, a, root_a = self.make_batch()
        self.make_batch("batch-b")
        values = load_annotated_trajectories(batch_id="batch-a", data_root=self.root)
        self.assertEqual(len(values["TASK"]), 2)
        self.assertEqual(len({item["trajectory_id"] for item in values["TASK"]}), 2)
        self.assertEqual({item["source_trajectory_id"] for item in values["TASK"]}, {"repeated-original"})
        for trajectory in values["TASK"]:
            step = trajectory["steps"][0]
            self.assertEqual(step["action_summary"], "batch-a")
            self.assertIn("batch_id=batch-a", step["image_url"])
            self.assertIn(a["version"], step["image_url"])
            path = resolve_image_asset(step["image"], batch_id="batch-a", data_root=self.root)
            self.assertTrue(path.is_relative_to(root_a))
        self.assertEqual(task_summaries(batch_id="batch-b", data_root=self.root)[0]["goal"], "Frozen goal batch-b")

    def test_edit_creates_new_version_retaining_hidden_json_identity_and_excel_columns(self):
        payload, annotation, _ = self.make_batch()
        before = self.store.resolve_file(annotation, "result.json").read_bytes()
        result = self.edit("batch-a", annotation, payload)
        current = resolve_batch_context("batch-a", root=self.root)
        self.assertEqual(current.annotation_version, result["annotation_version"])
        self.assertEqual(self.store.resolve_file(annotation, "result.json").read_bytes(), before)
        old_row, new_row = payload["sheets"]["VLA trajectories"][0], current.payload["sheets"]["VLA trajectories"][0]
        self.assertEqual({k: v for k, v in new_row.items() if k != "actions_box"},
                         {k: v for k, v in old_row.items() if k != "actions_box"})
        self.assertEqual(current.annotation_ref["source_refs"][0]["version"], annotation["version"])
        excel = self.store.resolve_file(current.annotation_ref, "annotated_trajectories.xlsx")
        book = load_workbook(excel, read_only=True)
        try:
            self.assertEqual([cell.value for cell in book.active[1]], COLUMNS)
        finally:
            book.close()
        with self.assertRaises(AnnotationVersionConflict):
            self.edit("batch-a", annotation, payload)

    def test_concurrent_edits_based_on_same_version_only_publish_once(self):
        payload, annotation, _ = self.make_batch()
        def attempt():
            try:
                return self.edit("batch-a", annotation, payload)
            except AnnotationVersionConflict:
                return "conflict"
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: attempt(), range(2)))
        self.assertEqual(sum(isinstance(item, dict) for item in results), 1)
        self.assertEqual(results.count("conflict"), 1)
        self.assertEqual(len(self.store.list("batch-a", "02_annotation")), 2)

    def test_historical_registered_batch_keeps_original_ids_without_zero_stage(self):
        _, annotation, raw = self.make_batch("validation-batch", historical=True)
        context = resolve_batch_context("validation-batch", annotation["version"], self.root)
        self.assertEqual(context.raw_root, raw)
        self.assertFalse(self.store.list("validation-batch", "00_collection"))
        values = load_annotated_trajectories(batch_id="validation-batch", data_root=self.root)
        self.assertEqual([item["trajectory_id"] for item in values["TASK"]], ["TASK-1", "TASK-2"])

    def test_json_required_and_verified_even_when_excel_exists(self):
        _, annotation, _ = self.make_batch()
        path = self.store.resolve_file(annotation, "result.json")
        path.write_text("{}", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "checksum"):
            resolve_batch_context("batch-a", root=self.root)
        path.unlink()
        with self.assertRaises(FileNotFoundError):
            resolve_batch_context("batch-a", root=self.root)

    def test_missing_excel_does_not_change_json_and_cross_batch_source_rejected(self):
        _, annotation, raw = self.make_batch()
        excel = next(item for item in annotation["files"] if item["kind"] == "excel")
        (self.root / excel["path"]).unlink()
        self.assertEqual(resolve_batch_context("batch-a", root=self.root).raw_root, raw)
        self.store.publish("other", "02_annotation", {"schema_version": 1, "sheets": {}},
                           source_refs=[annotation])
        with self.assertRaisesRegex(ValueError, "跨越批次"):
            resolve_batch_context("other", root=self.root)

    def test_job_freezes_annotation_when_queued_before_a_later_edit(self):
        payload, annotation, _ = self.make_batch()
        executor = DeferredExecutor()
        observed = {}
        def runner(_task_ids, **kwargs):
            context = resolve_batch_context(kwargs["batch_id"], kwargs["annotation_version"], kwargs["data_root"])
            observed.update(version=context.annotation_version, row=context.payload["sheets"]["VLA trajectories"][0])
            return "mock-run", {}
        manager = TreeBuildJobManager(self.root / "system/jobs", runner, executor, data_root=self.root)
        job = manager.submit(["TASK"], batch_id="batch-a")
        self.edit("batch-a", annotation, payload)
        executor.finish()
        self.assertEqual(job["annotation_version"], annotation["version"])
        self.assertEqual(observed["version"], annotation["version"])
        self.assertEqual(observed["row"]["actions_box"], payload["sheets"]["VLA trajectories"][0]["actions_box"])
        self.assertEqual(manager.get(job["job_id"])["status"], "succeeded")

    def test_job_listing_filters_batches_and_orders_newest_first(self):
        self.make_batch()
        self.make_batch("batch-b")
        manager = TreeBuildJobManager(self.root / "system/jobs", executor=DeferredExecutor(), data_root=self.root)
        first = manager.submit(["TASK"], batch_id="batch-a")
        second = manager.submit(["TASK"], batch_id="batch-b")
        third = manager.submit(["TASK"], batch_id="batch-a")
        self.assertEqual([item["job_id"] for item in manager.list_jobs()],
                         [third["job_id"], second["job_id"], first["job_id"]])
        self.assertEqual([item["job_id"] for item in manager.list_jobs("batch-a")],
                         [third["job_id"], first["job_id"]])
        self.assertEqual(manager.list_jobs("missing"), [])

    def test_tree_observation_and_quality_json_preserve_all_identity_fields(self):
        payload, annotation, raw = self.make_batch()
        runs = self.root / "system/trajectory_tree_runs"
        with patch("backend.tree_build_service.configure_reviewer_environment", return_value="fake-model"), \
             patch("backend.tree_build_service.QwenIntermediateStateClassifier", FakeClassifier), \
             patch("backend.tree_build_service.QwenStateAlignmentReviewer", FakeAlignmentReviewer):
            run_id, manifest = build_tree_run(["TASK"], job_id="test-job", progress=lambda _: None,
                batch_id="batch-a", annotation_version=annotation["version"], data_root=self.root, runs_dir=runs,
                env_path=self.root / "test.env", classification_cache=self.root / "cache/classification.json",
                alignment_cache=self.root / "cache/alignment.json",
                quality_builder=partial(build_quality_workbook, summarizer=FakeSummarizer()))
        self.assertEqual(manifest["batch_id"], "batch-a")
        self.assertEqual(manifest["raw_root"], str(raw))
        self.assertEqual(manifest["source_refs"][0]["version"], annotation["version"])
        frozen = json.loads((runs / run_id / "source_annotated.json").read_text(encoding="utf-8"))
        self.assertEqual(frozen["sheets"]["VLA trajectories"], payload["sheets"]["VLA trajectories"])
        observation = self.store.list("batch-a", "03_observation")[0]
        obs = json.loads(self.store.resolve_file(observation, "result.json").read_text(encoding="utf-8"))
        self.assertEqual(len(obs["trajectories"]), 2)
        self.assertEqual({item["steps"][0]["collection_run_id"] for item in obs["trajectories"]}, {"run-1", "run-2"})
        tree = json.loads((runs / run_id / "TASK.json").read_text(encoding="utf-8"))
        self.assertEqual(len(tree["source_trajectories"]), 2)
        quality = json.loads((runs / run_id / "rubric_trajectories.json").read_text(encoding="utf-8"))
        book = load_workbook(runs / run_id / "rubric_trajectories.xlsx", read_only=True)
        try:
            values = list(book["Tasks"].iter_rows(values_only=True))
            self.assertEqual([dict(zip(values[0], row)) for row in values[1:]], quality["sheets"]["Tasks"])
        finally:
            book.close()
        for row in quality["sheets"]["Steps"]:
            source = json.loads(row["action_input_json"])["source_identity"]
            self.assertEqual(source["trajectory_id"], row["trajectory_id"])
            self.assertEqual(source["source_trajectory_id"], "repeated-original")
            self.assertTrue(source["collected_at"])
        self.assertEqual(len({row["trajectory_id"] for row in quality["sheets"]["Trajectories"]}), 2)


if __name__ == "__main__":
    unittest.main()

