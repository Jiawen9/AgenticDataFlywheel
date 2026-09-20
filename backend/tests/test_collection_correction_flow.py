"""New collection JSON identities and real image files survive correction/COT."""
from __future__ import annotations

import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from backend import stage_artifacts
from backend.data_store import ArtifactStore, RecordStore
from backend.stage_artifacts import fingerprint, write_payload_workbook, write_sidecar
from backend.tests.test_collection_runs import raw_trajectory
from backend.tests.test_correction_storage import DeferredExecutor
from backend.trajectory_correction import cot_jobs, draft_store, quality_selection, service
from backend.trajectory_correction.assets import registered_asset_root
from backend.trajectory_correction.router import router
from backend.trajectory_correction.workbook import load_snapshot, open_source_workbook


class CollectionCorrectionFlowTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="collection-correction-")
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name) / "external-data"
        self.raw = self.root / "raw/collection_batches/batch-one"
        self.tree_root = self.root / "system/trajectory_tree_runs"
        self.quality_root = self.root / "system/trajectory_quality_results"
        self.tree_run = self.tree_root / "tree-one"
        self.tree_run.mkdir(parents=True)
        self.source = self.tree_run / "source_annotated.xlsx"
        self.inputs = self.root / "system/trajectory_correction/inputs"
        self.exports = self.root / "system/trajectory_correction/exports"
        self.patches = [
            patch.object(draft_store, "CORRECTION_SESSIONS_DIR", self.root / "sessions"),
            patch.object(quality_selection, "TREE_RUNS_DIR", self.tree_root),
            patch.object(quality_selection, "QUALITY_RESULTS_DIR", self.quality_root),
            # These fixtures explicitly exercise the legacy source adapter.
            patch.object(quality_selection, "_current_quality_runs", side_effect=lambda: quality_selection._completed_quality_runs(self.quality_root, self.tree_root)),
            patch.object(stage_artifacts, "DATA_ROOT", self.root),
            patch.object(service, "CORRECTION_INPUTS_DIR", self.inputs),
            patch.object(service, "CORRECTION_EXPORTS_DIR", self.exports),
            patch.object(service, "ensure_correction_dirs", side_effect=lambda: self.inputs.mkdir(parents=True, exist_ok=True)),
        ]
        for patcher in self.patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.columns = ["文件夹名", "image", "xml", "action", "summary", "actions_box"]
        self.identities = ["tr_" + "a" * 64, "tr_" + "b" * 64]
        self.rows = []
        for index, run_id in enumerate(("run-one", "run-two")):
            raw_trajectory({"output_dir": str(self.raw / "runs" / run_id)})
            prefix = f"runs/{run_id}/task-one/original-run/"
            self.rows.append({
                "文件夹名": "original-run", "image": prefix + "step001_vla_input.jpg",
                "xml": prefix + "step001_vla_input_ui.xml",
                "action": '{"action":"click","coordinate":[10,10]}', "summary": f"original {index}",
                "actions_box": "click(bbox=<bbox>[0,0,20,30]</bbox>)",
                "task_id": "task-one", "trajectory_id": self.identities[index],
                "source_trajectory_id": "original-run", "collection_run_id": run_id,
                "collection_case_id": "task-one", "source_result_id": "result-one",
                "collected_at": "2026-09-15T13:00:00+08:00",
            })
        self.payload = {"schema_version": 1, "columns": {"VLA trajectories": self.columns},
                        "sheets": {"VLA trajectories": self.rows}}
        write_payload_workbook(self.source, self.payload)
        write_sidecar(self.source, self.payload)
        self.manifest = {"run_id": "tree-one", "batch_id": "batch-one", "raw_root": str(self.raw),
            "annotation_version": "annotation-one", "completed_at": "2026-09-15T14:00:00+08:00",
            "source_annotated_file": self.source.name,
            "source_annotated_json": self.source.with_suffix(".json").name,
            "source_xlsx": {"sha256": fingerprint(self.source)},
            "source_json": {"sha256": fingerprint(self.source.with_suffix(".json"))},
            "tasks": [{"task_id": "task-one", "goal": "Open the settings page", "tree_file": "task-one.json", "trajectory_count": 2}]}
        (self.tree_run / "manifest.json").write_text(json.dumps(self.manifest), encoding="utf-8")
        (self.tree_run / "task-one.json").write_text('{"task_id":"task-one"}', encoding="utf-8")
        records = RecordStore(self.root)
        records.put("quality_manifests", "tree-one", {"run_id": "tree-one", "updated_at": "2026-09-15T15:00:00+08:00", "tasks": [{"task_id": "task-one", "trajectory_count": 2}]})
        records.put("quality_results", "tree-one:task-one", {"run_id": "tree-one", "task_id": "task-one", "trajectory_count": 2, "evaluations": {
            self.identities[0]: {"trajectory_id": self.identities[0], "global_score": 2.0, "passed_threshold": False},
            self.identities[1]: {"trajectory_id": self.identities[1], "global_score": 5.0, "passed_threshold": True},
        }})

    def test_hidden_json_identity_preserves_two_runs_without_adding_excel_columns(self):
        before = self.source.with_suffix(".json").read_bytes()
        self.source.unlink()
        snapshot = load_snapshot(self.source, asset_root=self.raw)
        self.assertEqual([group["meta_task"] for group in snapshot["groups"]], self.identities)
        self.assertEqual([group["task_id"] for group in snapshot["groups"]], ["task-one", "task-one"])
        self.assertEqual(snapshot["headers"], self.columns)
        self.assertEqual([group["rows"][0]["source_trajectory_id"] for group in snapshot["groups"]], ["original-run", "original-run"])
        selection = quality_selection.top1_selection_for_run("tree-one")
        self.assertEqual(selection["selected_trajectories"], {"task-one": self.identities[1]})
        self.assertEqual(selection["tasks"][0]["trajectory_count"], 2)
        workbook = open_source_workbook(self.source)
        self.assertEqual([cell.value for cell in workbook.active[1]], self.columns)
        workbook.close()
        self.assertEqual(self.source.with_suffix(".json").read_bytes(), before)

    def test_session_image_endpoint_and_mock_cot_read_selected_run_after_reload(self):
        created = service.create_session("tree-one")
        sid = created["session_id"]
        saved = draft_store.load_session(sid)
        self.assertEqual(saved["source_snapshot"]["raw_root"], str(self.raw))
        self.assertEqual(created["row_count"], 1)
        group = service.get_group(sid, "group_0")
        row = group["rows"][0]
        self.assertEqual(row["trajectory_id"], self.identities[1])
        expected_image = self.raw / self.rows[1]["image"]
        self.assertEqual(service.session_asset(sid, row["image"]), expected_image)
        # A future tree source change must not retarget this frozen session.
        modified = {**self.manifest, "raw_root": str(self.root / "raw/other-batch")}
        (self.tree_run / "manifest.json").write_text(json.dumps(modified))
        self.source.unlink()
        self.source.with_suffix(".json").unlink()
        self.assertEqual(service.get_session(sid)["row_count"], 1)
        app = FastAPI()
        app.include_router(router)
        with TestClient(app) as client:
            response = client.get(row["image_url"])
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
            self.assertEqual(response.content, expected_image.read_bytes())
            self.assertEqual(response.headers["content-type"], "image/jpeg")
        service.patch_row(sid, row["excel_row"], {"actions": '{"action":"wait"}'})
        cot_row = service.get_cot(sid)["groups"][0]["rows"][0]
        self.assertEqual(cot_row["task_id"], "task-one")
        self.assertEqual(cot_row["trajectory_id"], self.identities[1])
        owner, calls = self, []
        class Generator:
            model = "offline-mock"
            def generate(self, **kwargs):
                calls.append(kwargs)
                owner.assertEqual(kwargs["image"], expected_image)
                owner.assertEqual(kwargs["trajectory_id"], owner.identities[1])
                owner.assertEqual(kwargs["task"], "Open the settings page")
                with Image.open(kwargs["image"]) as image:
                    owner.assertEqual(image.size, (20, 30))
                return {"thought": "generated thought", "summary": "generated summary"}
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "cot_jobs", generator_factory=Generator, executor=queue)
        job = manager.submit(sid)
        queue.run()
        done = manager.get(job["job_id"])
        self.assertEqual(done["status"], "succeeded", done.get("error"))
        self.assertEqual(len(calls), 1)
        self.assertEqual(service.get_group(sid, "group_0")["rows"][0]["thought"], "generated thought")
        artifact = ArtifactStore(self.root).list("batch-one", "07_cot")[0]
        workbook = load_workbook(ArtifactStore(self.root).resolve_file(artifact, "full_dataset.xlsx"), read_only=True)
        try:
            headers = list(next(workbook.active.values))
            self.assertEqual(headers[:len(self.columns)], self.columns)
            self.assertNotIn("task_id", headers)
            self.assertNotIn("trajectory_id", headers)
            self.assertEqual(workbook.active.max_row, 3)
            self.assertEqual(workbook.active.cell(2, headers.index("summary") + 1).value, "original 0")
            self.assertEqual(workbook.active.cell(3, headers.index("summary") + 1).value, "generated summary")
        finally:
            workbook.close()

    def test_new_collection_identity_missing_is_error_and_old_rows_are_not_augmented(self):
        changed = copy.deepcopy(self.payload)
        changed["sheets"]["VLA trajectories"][0].pop("task_id")
        write_sidecar(self.source, changed)
        with self.assertRaisesRegex(ValueError, "task_id"):
            load_snapshot(self.source, asset_root=self.raw)
        legacy = {"schema_version": 1, "columns": {"VLA trajectories": self.columns},
                  "sheets": {"VLA trajectories": [{column: self.rows[0][column] for column in self.columns}]}}
        legacy["sheets"]["VLA trajectories"][0]["image"] = "OLD/OLD-1/step001_vla_input.jpg"
        legacy["sheets"]["VLA trajectories"][0]["文件夹名"] = "OLD-1"
        write_sidecar(self.source, legacy)
        snapshot = load_snapshot(self.source, asset_root=self.raw)
        self.assertEqual(snapshot["groups"][0]["meta_task"], "OLD-1")
        self.assertNotIn("task_id", snapshot["groups"][0])
        self.assertNotIn("collection_run_id", snapshot["groups"][0]["rows"][0])

    def test_manifest_resource_root_outside_data_is_rejected(self):
        outside = self.root.parent / "outside"
        outside.mkdir()
        self.manifest["raw_root"] = str(outside)
        (self.tree_run / "manifest.json").write_text(json.dumps(self.manifest))
        with self.assertRaisesRegex(ValueError, "raw"):
            quality_selection.source_for_tree_run("tree-one")
        with self.assertRaises(ValueError):
            registered_asset_root(str(outside), self.root)

    def test_frozen_session_survives_root_relocation_without_changing_inputs(self):
        created = service.create_session("tree-one")
        sid = created["session_id"]
        old_root = self.root
        relocated = old_root.parent / "renamed-workspace"
        before = {path.relative_to(old_root): path.read_bytes()
                  for path in (self.inputs / sid).iterdir() if path.is_file()}
        old_root.rename(relocated)
        RecordStore(relocated).put("data_root_relocations", "move", {
            "old_root": str(old_root), "new_root": str(relocated), "status": "applied"})
        with patch.object(draft_store, "CORRECTION_SESSIONS_DIR", relocated / "sessions"), \
             patch.object(service, "CORRECTION_INPUTS_DIR", relocated / self.inputs.relative_to(old_root)), \
             patch.object(service, "CORRECTION_EXPORTS_DIR", relocated / self.exports.relative_to(old_root)):
            self.assertEqual(service.get_session(sid)["row_count"], 1)
            row = service.get_group(sid, "group_0")["rows"][0]
            image = service.session_asset(sid, row["image"])
            self.assertEqual(image, relocated / self.raw.relative_to(old_root) / self.rows[1]["image"])
            exported = service.export_dataset_session(sid)
            self.assertTrue(service.download_export(sid, exported["filename"]).is_file())
            self.assertEqual(draft_store.load_session(sid)["source_snapshot"]["raw_root"], str(self.raw))
        self.assertEqual({relative: (relocated / relative).read_bytes() for relative in before}, before)


if __name__ == "__main__":
    unittest.main()
