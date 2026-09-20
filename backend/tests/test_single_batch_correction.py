"""A batch keeps one draft while task inputs and spreadsheet positions change."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from backend.data_store import ArtifactStore, RecordStore
from backend.stage_artifacts import write_payload_workbook, write_sidecar
from backend.trajectory_correction import service, draft_store, cot_jobs, quality_selection
from backend.trajectory_correction.session_state import migrate_session_identity
from backend.tests.test_correction_storage import DeferredExecutor


class SingleBatchCorrectionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source.xlsx"
        self.raw = self.root / "raw"
        self.raw.mkdir()
        self.selection = {}
        for patcher in (
            patch.object(draft_store, "CORRECTION_SESSIONS_DIR", self.root / "sessions"),
            patch.object(service, "CORRECTION_INPUTS_DIR", self.root / "inputs"),
            patch.object(service, "CORRECTION_EXPORTS_DIR", self.root / "exports"),
            patch.object(service, "ensure_correction_dirs"),
            patch.object(service, "top1_selection_for_run", side_effect=lambda _: deepcopy(self.selection)),
            patch("backend.trajectory_correction.quality_selection.source_for_tree_run", return_value=(self.source, self.raw)),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    def inputs(self, tasks, *, changed=None, alternate=None):
        changed, alternate = changed or set(), alternate or set()
        rows = []
        for task in tasks:
            trajectory = task + ("-2" if task in alternate else "-1")
            rows.append({"文件夹名": trajectory, "task_id": task, "trajectory_id": trajectory,
                "image": f"{task}/{trajectory}/step001_vla_input.jpg", "action": '{"action":"wait"}',
                "summary": "changed baseline" if task in changed else "baseline", "sop": "original",
                "actions_box": "wait()"})
        self.payload = {"schema_version": 1, "columns": {"VLA trajectories": list(rows[0])},
                        "sheets": {"VLA trajectories": rows}}
        write_payload_workbook(self.source, self.payload)
        write_sidecar(self.source, self.payload)
        self.selection = {"tree_run_id": "tree-current", "storage_batch_id": "batch-one",
            "tasks": [{"task_id": row["task_id"], "trajectory_id": row["trajectory_id"], "goal": row["task_id"], "trajectory_count": 1} for row in rows],
            "selected_trajectories": {row["task_id"]: row["trajectory_id"] for row in rows},
            "task_fingerprints": {task: task + ("-changed" if task in changed or task in alternate else "-original") for task in tasks}}

    def create(self):
        return service.create_session("tree-current")

    def test_http_mutations_require_revision_and_reject_stale_pages_without_saving(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from backend.trajectory_correction import router as routes
        self.inputs(["A"])
        created = self.create()
        sid, revision = created["session_id"], created["storage_revision"]
        group_id = created["groups"][0]["group_id"]
        app = FastAPI()
        app.include_router(routes.router)
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", executor=queue)
        self.addCleanup(manager.shutdown)
        with patch.object(routes, "_cot_job_manager", manager), TestClient(app) as client:
            prefix = f"/api/correction/sessions/{sid}"
            missing = client.patch(prefix + "/rows/2", json={"summary": "missing token"})
            self.assertEqual(missing.status_code, 422)
            updated = client.patch(prefix + "/rows/2", json={"summary": "saved expert edit", "expected_revision": revision})
            self.assertEqual(updated.status_code, 200, updated.text)
            before = draft_store.load_session(sid)
            requests = [
                ("patch", prefix + "/rows/2", {"summary": "stale edit"}),
                ("patch", prefix + f"/tasks/{group_id}/export", {"export": False}),
                ("patch", prefix + f"/tasks/{group_id}/review", {"decision": "discard"}),
                ("post", prefix + "/export", {}),
                ("post", prefix + "/dataset-export", {}),
                ("post", "/api/correction/cot-jobs", {"session_id": sid}),
            ]
            for method, url, data in requests:
                response = getattr(client, method)(url, json={**data, "expected_revision": revision})
                self.assertEqual(response.status_code, 409, (url, response.text))
                self.assertEqual(draft_store.load_session(sid), before)
            self.assertEqual(manager.list_jobs(), [])

    def test_appending_task_and_reordering_rows_preserves_existing_manual_work(self):
        self.inputs(["A"])
        first = self.create()
        sid = first["session_id"]
        service.patch_row(sid, 2, {"summary": "expert A"})
        self.inputs(["B", "A"])
        second = self.create()
        self.assertEqual(first["session_id"], second["session_id"])
        self.assertEqual(len(draft_store.list_sessions()), 1)
        saved = draft_store.load_session(sid)
        self.assertEqual(saved["row_edits"], {"3": {"summary": "expert A"}})
        self.assertFalse(saved["pending_review"])
        self.assertEqual(self.create()["storage_revision"], second["storage_revision"])

    def test_changed_task_waits_for_review_while_other_task_exports(self):
        self.inputs(["A", "B"])
        first = self.create()
        sid = first["session_id"]
        service.patch_row(sid, 2, {"summary": "expert A"})
        service.patch_row(sid, 3, {"summary": "expert B"})
        service.on_batch_tasks_invalidated("batch-one", ["A"], root=self.root)
        self.inputs(["A", "B"], changed={"A"})
        current = self.create()
        group_a = next(group for group in current["groups"] if group["task_id"] == "A")
        self.assertTrue(group_a["pending_review"])
        self.assertTrue(group_a["can_adopt_review"])
        exported = service.export_dataset_session(sid)
        path = self.root / "exports" / sid / exported["filename"]
        book = load_workbook(path)
        self.assertEqual(book.active.max_row, 2)
        headers = [cell.value for cell in book.active[1]]
        self.assertEqual(book.active.cell(2, headers.index("task_id") + 1).value, "B")
        self.assertEqual(book.active.cell(2, headers.index("summary") + 1).value, "expert B")
        book.close()
        artifact = ArtifactStore(self.root).read_payload(exported["artifact"])
        self.assertEqual([group["task_id"] for group in artifact["groups"]], ["B"])
        self.assertEqual(artifact["row_edits"]["2"]["summary"], "expert B")
        result = service.review_group(sid, group_a["group_id"], "adopt")
        self.assertFalse(result["group"]["pending_review"])
        self.assertEqual(draft_store.load_session(sid)["row_edits"]["2"]["summary"], "expert A")

    def test_changed_top1_never_applies_edits_to_different_trajectory(self):
        self.inputs(["A"])
        sid = self.create()["session_id"]
        service.patch_row(sid, 2, {"thought": "old trajectory edit"})
        self.inputs(["A"], alternate={"A"})
        current = self.create()
        group = current["groups"][0]
        self.assertFalse(group["can_adopt_review"])
        with self.assertRaisesRegex(ValueError, "原步骤"):
            service.review_group(sid, group["group_id"], "adopt")
        service.review_group(sid, group["group_id"], "discard")
        self.assertEqual(draft_store.load_session(sid)["row_edits"], {})

    def test_cot_result_from_replaced_input_is_rejected(self):
        self.inputs(["A"])
        sid = self.create()["session_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        owner = self
        class Generator:
            model = "offline"
            def generate(self, **kwargs):
                service.on_batch_tasks_invalidated("batch-one", ["A"], root=owner.root)
                return {"summary": "obsolete", "thought": "obsolete"}
        executor = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", generator_factory=Generator, executor=executor)
        job = manager.submit(sid)
        with patch.object(cot_jobs, "session_asset", return_value=self.root / "mock.jpg"):
            executor.run()
        self.assertEqual(manager.get(job["job_id"])["status"], "failed")
        self.assertFalse(draft_store.load_session(sid)["cot"])

    def test_current_exports_reuse_identical_input_and_replace_only_their_kind(self):
        self.inputs(["A"])
        sid = self.create()["session_id"]
        first = service.export_dataset_session(sid)
        self.assertEqual(service.export_dataset_session(sid)["export_id"], first["export_id"])
        service.patch_row(sid, 2, {"summary": "later"})
        second = service.export_dataset_session(sid)
        self.assertNotEqual(second["export_id"], first["export_id"])
        self.assertFalse((self.root / "exports" / sid / first["filename"]).exists())
        self.assertEqual(len(draft_store.load_session(sid)["exports"]), 1)
        with self.assertRaises(FileNotFoundError):
            service.download_export(sid, first["filename"])

    def test_pure_migration_keeps_manual_and_cot_values(self):
        self.inputs(["A"])
        sid = self.create()["session_id"]
        service.patch_row(sid, 2, {"summary": "expert"})
        existing = draft_store.load_session(sid)
        existing.update(published=True, published_release_id="rel_old", cot={"2": {"thought": "old generated"}})
        updated = migrate_session_identity(existing, service._snapshot(existing), batch_id="batch-one",
            table_payload=self.payload, task_fingerprints=self.selection["task_fingerprints"])
        self.assertEqual(updated["row_edits"], existing["row_edits"])
        self.assertEqual(updated["cot"], existing["cot"])
        self.assertTrue(updated["published"])
        self.assertTrue(updated["published_content_fingerprint"])

    def publishable_session(self):
        from backend.batch_results import annotation_task_fingerprints
        self.inputs(["A", "B"])
        store = ArtifactStore(self.root)
        conversion = store.publish("batch-one", "01_conversion", self.payload)
        annotation = store.publish("batch-one", "02_annotation", self.payload, source_refs=[conversion])
        observation = store.publish("batch-one", "03_observation", {}, source_refs=[annotation])
        task_fingerprints = {"A": "tree-A", "B": "tree-B"}
        tree = store.publish("batch-one", "04_tree", {
            "run_id": "batch-one", "batch_id": "batch-one", "raw_root": str(self.raw),
            "trees": {task: {"task_id": task} for task in ("A", "B")},
            "tasks": [{"task_id": task, "goal": task, "trajectory_count": 1} for task in ("A", "B")],
            "task_fingerprints": task_fingerprints,
            "source_task_fingerprints": annotation_task_fingerprints(self.payload),
            "tree_hashes": {"A": "hash-A", "B": "hash-B"},
            "source_annotation": self.payload,
            "quality_input": {"sheets": {"Tasks": [{"task_id": "A"}, {"task_id": "B"}]}}},
            source_refs=[observation], metadata={"run_id": "batch-one"})
        store.publish("batch-one", "05_quality", {
            "run_id": "batch-one", "source_tree_hashes": {"A": "hash-A", "B": "hash-B"},
            "tasks": [{"task_id": task, "trajectory_count": 1, "evaluations": {
                task + "-1": {"global_score": 5, "passed_threshold": True}}} for task in ("A", "B")]},
            source_refs=[tree])
        with patch.object(service, "top1_selection_for_run", quality_selection.top1_selection_for_run):
            public = service.create_session(batch_id="batch-one")
            self.assertEqual(public["batch_id"], "batch-one")
            self.assertEqual(len(public["groups"]), 2)
            self.assertEqual(quality_selection.correction_batches()["default_batch_id"], "batch-one")
        return public, store

    def release_registry(self):
        from backend.data_publishing.service import DatasetReleaseRegistry
        return DatasetReleaseRegistry(releases_file=self.root / "registry.json", project_root=self.root,
            data_root=self.root, trajectory_root=self.raw, correction_exports_dir=self.root / "exports")

    def test_published_batch_rejects_later_changes_and_keeps_frozen_statistics(self):
        from backend.batch_lifecycle import BatchPublishedError
        from backend.training_data_overview.converter import convert_release
        public, store = self.publishable_session()
        sid = public["session_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        exported = service.export_dataset_session(sid)
        registry = self.release_registry()
        release = registry.create("frozen release", [sid])
        expected = convert_release(registry, release)
        self.assertEqual(sum(row["人工精修步骤数量"] for row in expected["rows"]), 1)
        session_before = draft_store.load_session(sid)
        artifacts_before = store.list("batch-one")
        with self.assertRaises(BatchPublishedError):
            store.publish("batch-one", "02_annotation", self.payload)
        with self.assertRaises(BatchPublishedError):
            service.on_batch_tasks_invalidated("batch-one", ["A"], root=self.root)
        with self.assertRaises(BatchPublishedError):
            draft_store.save_session({**session_before, "published": False})
        self.assertEqual(store.list("batch-one"), artifacts_before)
        self.assertEqual(draft_store.load_session(sid), session_before)
        self.assertEqual(convert_release(registry, release), expected)
        self.assertTrue(session_before["published"])
        self.assertTrue(service.download_export(sid, exported["filename"]).is_file())

    def test_closed_workbench_http_rejects_reads_and_mutations_without_reconciling(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from backend.batch_lifecycle import install_lifecycle_handlers
        from backend.trajectory_correction import router as routes
        public, store = self.publishable_session()
        sid = public["session_id"]
        group_id = public["groups"][0]["group_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        exported = service.export_dataset_session(sid)
        release = self.release_registry().create("complete", [sid])
        RecordStore(self.root).put("batch_run_aliases", "legacy-run", {"batch_id": "batch-one"})
        before = draft_store.load_session(sid)
        artifacts = store.list("batch-one")
        app = FastAPI()
        install_lifecycle_handlers(app)
        app.include_router(routes.router)
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", executor=queue)
        self.addCleanup(manager.shutdown)
        with patch.object(routes, "_cot_job_manager", manager), TestClient(app) as client:
            prefix = f"/api/correction/sessions/{sid}"
            for url in (prefix, prefix + "/tasks", prefix + f"/tasks/{group_id}", prefix + "/cot",
                        "/api/correction/recommendation?batch_id=batch-one",
                        "/api/correction/recommendation?tree_run_id=legacy-run"):
                response = client.get(url)
                self.assertEqual(response.status_code, 409, response.text)
                self.assertEqual(response.json()["detail"]["code"], "batch_published")
                self.assertEqual(response.json()["detail"]["release_id"], release["release_id"])
            writes = [
                ("post", "/api/correction/sessions", {"batch_id": "batch-one"}),
                ("post", "/api/correction/sessions", {"tree_run_id": "batch-one"}),
                ("post", "/api/correction/sessions", {"tree_run_id": "legacy-run"}),
                ("patch", prefix + "/rows/2", {"summary": "forbidden"}),
                ("patch", prefix + f"/tasks/{group_id}/export", {"export": False}),
                ("patch", prefix + f"/tasks/{group_id}/review", {"decision": "discard"}),
                ("post", prefix + "/export", {}),
                ("post", prefix + "/dataset-export", {}),
                ("post", "/api/correction/cot-jobs", {"session_id": sid}),
            ]
            for method, url, payload in writes:
                response = getattr(client, method)(url, json={**payload, "expected_revision": before["storage_revision"]})
                self.assertEqual(response.status_code, 409, (url, response.text))
                self.assertEqual(response.json()["detail"]["code"], "batch_published")
            self.assertEqual(client.get("/api/correction/sessions").json()["sessions"], [])
            self.assertEqual(client.get("/api/correction/batches").json()["batches"], [])
            self.assertEqual(client.get(exported["download_url"]).status_code, 200)
        self.assertEqual(draft_store.load_session(sid), before)
        self.assertEqual(store.list("batch-one"), artifacts)
        self.assertEqual(manager.list_jobs(), [])

    def test_closed_cot_request_cannot_reuse_a_successful_job(self):
        from backend.batch_lifecycle import BatchPublishedError
        public, store = self.publishable_session()
        sid = public["session_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        class Generator:
            model = "offline"
            def generate(self, **kwargs):
                return {"summary": "generated", "thought": "generated"}
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", generator_factory=Generator, executor=queue)
        job = manager.submit(sid)
        with patch.object(cot_jobs, "session_asset", return_value=self.root / "mock.jpg"):
            queue.run()
        self.assertEqual(manager.get(job["job_id"])["status"], "succeeded")
        service.export_dataset_session(sid)
        self.release_registry().create("complete", [sid])
        before = draft_store.load_session(sid)
        with self.assertRaises(BatchPublishedError):
            manager.submit(sid)
        self.assertEqual(draft_store.load_session(sid), before)
        self.assertEqual(len(manager.list_jobs()), 1)
        self.assertEqual(manager.list_jobs(active_only=True), [])

    def test_queued_cot_start_after_publication_skips_model_and_draft_write(self):
        public, store = self.publishable_session()
        sid = public["session_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        service.export_dataset_session(sid)
        class Generator:
            def __init__(self):
                raise AssertionError("closed batch must not construct model")
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", generator_factory=Generator, executor=queue)
        job = manager.submit(sid)
        # An interrupted job may still have a queued executor callback in another process.
        manager._progress(job["job_id"], {"status": "interrupted"})
        self.release_registry().create("complete", [sid])
        before = draft_store.load_session(sid)
        artifacts = store.list("batch-one")
        queue.run()
        self.assertEqual(manager.get(job["job_id"])["stage"], "batch_published")
        self.assertEqual(draft_store.load_session(sid), before)
        self.assertEqual(store.list("batch-one"), artifacts)

    def test_cot_response_after_publication_cannot_overwrite_draft_or_stage(self):
        public, store = self.publishable_session()
        sid = public["session_id"]
        service.patch_row(sid, 2, {"actions": '{"action":"finish"}'})
        service.export_dataset_session(sid)
        registry = self.release_registry()
        before = {}
        class Generator:
            model = "offline"
            def generate(inner, **kwargs):
                manager._progress(job["job_id"], {"status": "interrupted"})
                registry.create("complete", [sid])
                before["session"] = draft_store.load_session(sid)
                before["artifacts"] = store.list("batch-one")
                return {"summary": "late forbidden", "thought": "late forbidden"}
        queue = DeferredExecutor()
        manager = cot_jobs.CotJobManager(self.root / "jobs", generator_factory=Generator, executor=queue)
        job = manager.submit(sid)
        with patch.object(cot_jobs, "session_asset", return_value=self.root / "mock.jpg"):
            queue.run()
        self.assertEqual(manager.get(job["job_id"])["stage"], "batch_published")
        self.assertEqual(draft_store.load_session(sid), before["session"])
        self.assertEqual(store.list("batch-one"), before["artifacts"])
