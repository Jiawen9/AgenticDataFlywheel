from __future__ import annotations

import io
import json
import shutil
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.task_generation.jobs import TaskGenerationJobManager
from backend.task_generation.router import router


EXPECTED_COLUMNS = [
    "用例编号", "单APP跨APP", "涉及APP", "一级场景", "二级场景", "三级场景", "任务",
    "难易程度", "页面关键元素", "预期步数", "用例来源", "构建人", "任务结束标志",
    "任务复杂度", "关键动作", "SOP", "动态变量说明",
]
SUMMARY_FIELDS = {
    "schema_version", "batch_id", "source_job_id", "kind", "job_status", "knowledge_base_version",
    "created_at", "task_count", "apps", "filename", "download_url",
}


class NoExecution:
    def submit(self, *args, **kwargs):
        raise AssertionError("冻结采集批次不能启动模型或采集任务")


class CollectionBatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.manager = self.make_manager()
        self.sequence = 0

    def make_manager(self, **kwargs):
        manager = TaskGenerationJobManager(
            self.base / "jobs", self.base / "runs", self.base / "exports", self.base / "kb", self.base / "logs",
            executor=NoExecution(), **kwargs,
        )
        self.addCleanup(manager.shutdown)
        return manager

    @staticmethod
    def row(result_id="result-1", **changes):
        return {"result_id": result_id, "task_uuid": f"uuid-{result_id}", "task": "在 AppA 中查找视频",
                "app": "AppA", "scene": "场景A", "capability": "能力A", "sub_capability": "子能力A",
                "pre_dependency": "zero", "pre_task_uuid": None, "deleted": False, **changes}

    def job(self, rows=None, *, kind="task_generation", status="succeeded", **changes):
        self.sequence += 1
        job_id = f"source-job-{self.sequence}"
        self.manager._write_job({"job_id": job_id, "kind": kind, "status": status, "stage": status,
                                 "knowledge_base_version": "kb-before-submit", "errors": [], "warnings": [], **changes})
        self.manager._write_json(self.manager._results_path(job_id), [self.row()] if rows is None else rows)
        return job_id

    def all_source_files(self):
        batch_root = self.manager.collection_batches_dir
        return {str(path.relative_to(self.base)): (path.read_bytes(), path.stat().st_mtime_ns)
                for path in self.base.rglob("*") if path.is_file() and not path.is_relative_to(batch_root)}

    def workbook_rows(self, batch_id):
        workbook = load_workbook(self.manager.collection_batch_workbook(batch_id), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Sheet1"])
        sheet = workbook["Sheet1"]
        self.assertEqual(sheet.max_column, 17)
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), EXPECTED_COLUMNS)
        return sheet, rows[1:]

    def test_submit_creates_exact_summary_and_frozen_input_without_changing_sources(self):
        job_id = self.job([self.row("one"), self.row("deleted", deleted=True), self.row("two", app="AppB")])
        # An existing legacy export is unrelated and must not be rewritten.
        (self.manager.exports_dir / "old.xlsx").write_bytes(b"legacy-export")
        before = self.all_source_files()
        current = self.manager.collection_input(job_id)
        detail, created = self.manager.submit_collection_batch(job_id)
        self.assertTrue(created)
        self.assertEqual(set(detail), SUMMARY_FIELDS | {"snapshot"})
        self.assertEqual(detail["batch_id"], job_id)
        self.assertEqual(detail["source_job_id"], job_id)
        self.assertEqual(detail["schema_version"], 1)
        self.assertEqual(detail["job_status"], "succeeded")
        self.assertEqual(detail["knowledge_base_version"], "kb-before-submit")
        self.assertEqual(detail["task_count"], 2)
        self.assertEqual(detail["apps"], ["AppA", "AppB"])
        self.assertEqual(detail["filename"], f"collection-batch-{job_id}.xlsx")
        self.assertEqual(detail["download_url"], f"/api/task-generation/collection-batches/{job_id}/workbook")
        self.assertTrue(detail["created_at"])
        self.assertEqual(set(detail["snapshot"]), set(current))
        for source, frozen in zip(current["tasks"], detail["snapshot"]["tasks"]):
            self.assertEqual({key: value for key, value in frozen.items() if key != "collection_case_id"}, source)
            self.assertEqual(frozen["collection_case_id"], source["task_id"])
        self.assertEqual(self.all_source_files(), before)
        directory = self.manager.collection_batches_dir / job_id
        self.assertEqual({path.name for path in directory.iterdir()}, {"batch.json", detail["filename"]})
        stored = json.loads((directory / "batch.json").read_text(encoding="utf-8"))
        integrity = stored.pop("_integrity")
        self.assertEqual(stored, detail)
        self.assertEqual(set(integrity), {"payload_sha256", "workbook_sha256"})
        self.assertRegex(integrity["payload_sha256"], r"^[0-9a-f]{64}$")
        self.assertRegex(integrity["workbook_sha256"], r"^[0-9a-f]{64}$")

    def test_repeated_submission_is_immutable_after_edit_delete_and_source_state_change(self):
        job_id = self.job()
        original, _ = self.manager.submit_collection_batch(job_id)
        workbook = self.manager.collection_batch_workbook(job_id)
        original_bytes = workbook.read_bytes()
        original_mtime = workbook.stat().st_mtime_ns
        self.manager.patch_result(job_id, "result-1", {"task": "后续人工修改", "deleted": True})
        self.manager._progress(job_id, {"status": "failed", "knowledge_base_version": "kb-after-submit"})
        with patch("backend.task_generation.collection_batches.write_collection_workbook", side_effect=AssertionError("重复提交不得重写表格")), \
             patch.object(self.manager, "collection_input", side_effect=AssertionError("重复提交不得重新冻结源数据")):
            repeated, created = self.manager.submit_collection_batch(job_id)
        self.assertFalse(created)
        self.assertEqual(repeated, original)
        self.assertEqual(workbook.read_bytes(), original_bytes)
        self.assertEqual(workbook.stat().st_mtime_ns, original_mtime)
        self.assertEqual(len(self.manager.collection_batches()), 1)

    def test_concurrent_submissions_publish_only_one_batch(self):
        job_id = self.job()
        from backend.task_generation.collection_batches import write_collection_workbook
        with patch("backend.task_generation.collection_batches.write_collection_workbook", wraps=write_collection_workbook) as writer:
            with ThreadPoolExecutor(max_workers=8) as executor:
                outcomes = list(executor.map(lambda _index: self.manager.submit_collection_batch(job_id), range(16)))
        self.assertEqual(sum(created for _detail, created in outcomes), 1)
        self.assertEqual(writer.call_count, 1)
        self.assertTrue(all(detail == outcomes[0][0] for detail, _created in outcomes))
        self.assertEqual(len(self.manager.collection_batches()), 1)

    def test_restart_reads_original_batch_and_workbook_without_source_reconstruction(self):
        job_id = self.job()
        detail, _ = self.manager.submit_collection_batch(job_id)
        expected_workbook = self.manager.collection_batch_workbook(job_id).read_bytes()
        self.manager._results_path(job_id).write_text("corrupt after submission", encoding="utf-8")
        restored = self.make_manager()
        self.assertEqual(restored.collection_batch(job_id), detail)
        self.assertEqual(restored.collection_batch_workbook(job_id).read_bytes(), expected_workbook)
        self.assertEqual(restored.submit_collection_batch(job_id), (detail, False))

    def test_initial_workbook_maps_seven_values_and_keeps_ten_columns_blank(self):
        rows = [self.row("one", **{"用例编号": "普通生成不得选此编号"}), self.row("two", scene=None, capability=None, sub_capability=None)]
        job_id = self.job(rows)
        self.manager.submit_collection_batch(job_id)
        _sheet, values = self.workbook_rows(job_id)
        self.assertEqual(list(values[0][:7]), ["uuid-one", "单APP", "AppA", "场景A", "能力A", "子能力A", rows[0]["task"]])
        self.assertEqual(list(values[1][:7]), ["uuid-two", "单APP", "AppA", None, None, None, rows[1]["task"]])
        self.assertTrue(all(all(value is None for value in row[7:]) for row in values))

    def test_augmentation_case_id_preferred_with_legacy_fallback_and_unknown_dependency(self):
        rows = [
            {"result_id": "variant-a", "task": "变体任务甲", "app": "AppA", "用例编号": "CASE-A", "scene": "Unclassified"},
            {"result_id": "variant-b", "task": "变体任务乙", "app": "AppB", "用例编号": ""},
            {"result_id": "variant-c", "task": "变体任务丙", "app": "AppB", "用例编号": "  "},
        ]
        job_id = self.job(rows, kind="augmentation", knowledge_base_version=None)
        detail, _ = self.manager.submit_collection_batch(job_id)
        frozen = detail["snapshot"]["tasks"]
        self.assertEqual([task["collection_case_id"] for task in frozen], ["CASE-A", "variant-b", "variant-c"])
        self.assertTrue(all(task["pre_dependency"] == "unknown" for task in frozen))
        self.assertTrue(all(task["source_seed_id"] is None for task in frozen))
        self.assertEqual(frozen[0]["scene"], "Unclassified")
        self.assertIsNone(detail["knowledge_base_version"])
        _sheet, values = self.workbook_rows(job_id)
        self.assertEqual([row[0] for row in values], ["CASE-A", "variant-b", "variant-c"])
        self.assertEqual(values[0][3], "Unclassified")
        self.assertIsNone(values[1][3])

    def test_snapshot_retains_weak_strong_and_partial_metadata_but_excel_does_not_invent_extra_values(self):
        errors = [{"stage": "generating", "error": "有一条任务生成失败"}]
        rows = [
            self.row("pre", pre_dependency="pre_node"),
            self.row("main", pre_dependency="weak", pre_task_uuid="uuid-pre"),
            self.row("strong", pre_dependency="strong", status="-2", dependency_error={"source": "legacy"}),
        ]
        job_id = self.job(rows, status="partial", errors=errors, warnings=["原提示"])
        detail, _ = self.manager.submit_collection_batch(job_id)
        self.assertEqual(detail["job_status"], "partial")
        frozen = detail["snapshot"]
        self.assertEqual(frozen["errors"], errors)
        self.assertEqual(frozen["warnings"], ["原提示"])
        self.assertEqual([task["pre_dependency"] for task in frozen["tasks"]], ["pre_node", "weak", "strong"])
        self.assertEqual(frozen["tasks"][1]["pre_task_id"], frozen["tasks"][0]["task_id"])
        self.assertEqual(frozen["tasks"][2]["source_status"], "-2")
        self.assertEqual(frozen["tasks"][2]["dependency_error"], {"source": "legacy"})
        _sheet, values = self.workbook_rows(job_id)
        self.assertEqual(len(values), 3)
        self.assertTrue(all(all(value is None for value in row[7:]) for row in values))

    def test_formula_like_strings_are_literal_text_with_original_contents(self):
        row = self.row("source", task_uuid="=CASE(1)", task='=HYPERLINK("https://example.invalid","任务")',
                       app="+AppA", scene="-场景", capability="@能力", sub_capability="  原始子能力  ")
        job_id = self.job([row])
        self.manager.submit_collection_batch(job_id)
        sheet, values = self.workbook_rows(job_id)
        self.assertEqual(list(values[0][:7]), [row["task_uuid"], "单APP", row["app"], row["scene"], row["capability"], row["sub_capability"], row["task"]])
        self.assertTrue(all(sheet.cell(2, column).data_type == "s" for column in range(1, 8)))

    def test_invalid_excel_text_or_excess_length_rejects_without_publishing_or_truncating_source(self):
        cases = [
            ("augmentation", self.row(**{"用例编号": "CASE-\u0001"})),
            ("task_generation", self.row(task="任务\u0001文本")),
            ("task_generation", self.row(task="长" * 32768)),
        ]
        for kind, row in cases:
            with self.subTest(kind=kind, length=len(row["task"])):
                job_id = self.job([row], kind=kind)
                before = self.all_source_files()
                with self.assertRaises(ValueError):
                    self.manager.submit_collection_batch(job_id)
                self.assertFalse((self.manager.collection_batches_dir / job_id).exists())
                self.assertEqual(self.manager.collection_batches(job_id), [])
                self.assertEqual(self.all_source_files(), before)

    def test_duplicate_collection_case_ids_reject_without_publishing(self):
        cases = [
            [{"result_id": "a", "task": "变体甲", "app": "AppA", "用例编号": "SAME"}, {"result_id": "b", "task": "变体乙", "app": "AppA", "用例编号": "SAME"}],
            [{"result_id": "a", "task": "变体甲", "app": "AppA", "用例编号": "b"}, {"result_id": "b", "task": "变体乙", "app": "AppA"}],
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                job_id = self.job(rows, kind="augmentation")
                with self.assertRaises(ValueError):
                    self.manager.submit_collection_batch(job_id)
                self.assertEqual(self.manager.collection_batches(job_id), [])
                self.assertFalse((self.manager.collection_batches_dir / job_id).exists())

    def test_invalid_state_empty_and_invalid_dependency_inputs_do_not_publish(self):
        for state in ("queued", "running", "awaiting_confirmation", "failed", "interrupted"):
            with self.subTest(state=state), self.assertRaises(ValueError):
                self.manager.submit_collection_batch(self.job(status=state))
        for rows in ([], [self.row(deleted=True)], [self.row(pre_dependency="weak", pre_task_uuid="missing")], [self.row(task="")]):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.manager.submit_collection_batch(self.job(rows))
        with self.assertRaises(FileNotFoundError):
            self.manager.submit_collection_batch("missing-job")
        self.assertEqual(self.manager.collection_batches(), [])

    def test_workbook_write_failure_leaves_no_visible_batch_and_retry_succeeds(self):
        job_id = self.job()

        def fail_after_partial_write(snapshot, path):
            path.write_bytes(b"partial workbook")
            raise OSError("模拟写表格失败")

        with patch("backend.task_generation.collection_batches.write_collection_workbook", side_effect=fail_after_partial_write):
            with self.assertRaises(ValueError):
                self.manager.submit_collection_batch(job_id)
        self.assertEqual(self.manager.collection_batches(), [])
        self.assertFalse((self.manager.collection_batches_dir / job_id).exists())
        self.assertEqual(list(self.manager.collection_batches_dir.iterdir()), [])
        detail, created = self.manager.submit_collection_batch(job_id)
        self.assertTrue(created)
        self.assertEqual(self.manager.collection_batch(job_id), detail)

    def test_directory_metadata_and_publication_failures_never_expose_partial_batch(self):
        for stage in ("mkdir", "metadata", "rename"):
            with self.subTest(stage=stage):
                job_id = self.job()
                original_mkdir = Path.mkdir
                original_rename = Path.rename
                original_write = self.manager._write_json

                def mkdir(path, *args, **kwargs):
                    if path == self.manager.collection_batches_dir:
                        raise OSError("模拟创建批次目录失败")
                    return original_mkdir(path, *args, **kwargs)

                def write_json(path, value):
                    if path.name == "batch.json":
                        raise OSError("模拟写元数据失败")
                    return original_write(path, value)

                def rename(path, target):
                    if Path(target) == self.manager.collection_batches_dir / job_id:
                        raise OSError("模拟目录发布失败")
                    return original_rename(path, target)

                target = patch.object(Path, "mkdir", mkdir) if stage == "mkdir" else patch.object(self.manager, "_write_json", side_effect=write_json) if stage == "metadata" else patch.object(Path, "rename", rename)
                with target, self.assertRaises(ValueError):
                    self.manager.submit_collection_batch(job_id)
                self.assertEqual(self.manager.collection_batches(job_id), [])
                self.assertFalse((self.manager.collection_batches_dir / job_id).exists())

    def test_list_filters_and_ignores_hidden_staging_and_corrupt_directories(self):
        first = self.job()
        second = self.job([self.row("second", app="AppB")])
        detail1, _ = self.manager.submit_collection_batch(first)
        detail2, _ = self.manager.submit_collection_batch(second)
        staging = self.manager.collection_batches_dir / ".interrupted-tmp"
        shutil.copytree(self.manager.collection_batches_dir / first, staging)
        damaged = self.manager.collection_batches_dir / "broken"
        damaged.mkdir()
        (damaged / "batch.json").write_text("broken JSON", encoding="utf-8")
        listed = self.manager.collection_batches()
        self.assertEqual({batch["batch_id"] for batch in listed}, {first, second})
        self.assertTrue(all(set(batch) == SUMMARY_FIELDS for batch in listed))
        self.assertTrue(all("snapshot" not in batch for batch in listed))
        self.assertEqual(self.manager.collection_batches(first), [{key: value for key, value in detail1.items() if key != "snapshot"}])
        self.assertEqual(self.manager.collection_batches(second), [{key: value for key, value in detail2.items() if key != "snapshot"}])
        self.assertEqual(self.manager.collection_batches("not-created"), [])

    def test_missing_or_corrupt_batch_is_not_silently_recreated(self):
        with self.assertRaises(FileNotFoundError):
            self.manager.collection_batch("missing")
        job_id = self.job()
        self.manager.submit_collection_batch(job_id)
        workbook = self.manager.collection_batch_workbook(job_id)
        workbook.unlink()
        with self.assertRaises(ValueError):
            self.manager.collection_batch(job_id)
        with self.assertRaises(ValueError):
            self.manager.collection_batch_workbook(job_id)
        with self.assertRaises(ValueError):
            self.manager.submit_collection_batch(job_id)
        self.assertFalse(workbook.exists())

    def assert_corrupt_batch_rejected(self, job_id):
        for operation in (self.manager.collection_batch, self.manager.submit_collection_batch, self.manager.collection_batch_workbook):
            with self.subTest(operation=operation.__name__), self.assertRaises(ValueError):
                operation(job_id)
        self.assertEqual(self.manager.collection_batches(job_id), [])
        self.assertNotIn(job_id, [item["batch_id"] for item in self.manager.collection_batches()])
        app = FastAPI()
        app.include_router(router)
        prefix = "/api/task-generation"
        with patch("backend.task_generation.router.manager", self.manager), TestClient(app) as client:
            for method, url in (
                ("get", f"{prefix}/collection-batches/{job_id}"),
                ("post", f"{prefix}/jobs/{job_id}/collection-batch"),
                ("get", f"{prefix}/collection-batches/{job_id}/workbook"),
            ):
                response = getattr(client, method)(url)
                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.headers.get("Cache-Control"), "no-store")
            self.assertEqual(client.get(f"{prefix}/collection-batches", params={"job_id": job_id}).json(), {"batches": []})

    def test_truncated_or_replaced_workbook_fails_integrity_checks(self):
        for content in (b"", b"replaced and no longer an xlsx workbook"):
            with self.subTest(content=content):
                job_id = self.job()
                self.manager.submit_collection_batch(job_id)
                workbook = self.manager.collection_batch_workbook(job_id)
                workbook.write_bytes(content)
                self.assert_corrupt_batch_rejected(job_id)
                self.assertEqual(workbook.read_bytes(), content, "损坏的原文件不应被悄悄重建")

    def test_valid_json_with_changed_snapshot_fields_fails_integrity_checks(self):
        for field in ("job_id", "task"):
            with self.subTest(field=field):
                job_id = self.job()
                self.manager.submit_collection_batch(job_id)
                metadata = self.manager.collection_batches_dir / job_id / "batch.json"
                payload = json.loads(metadata.read_text(encoding="utf-8"))
                if field == "job_id":
                    payload["snapshot"]["job_id"] = "another-source-job"
                else:
                    payload["snapshot"]["tasks"][0]["task"] = "文件中的任务内容被意外替换"
                metadata.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
                damaged = metadata.read_bytes()
                self.assert_corrupt_batch_rejected(job_id)
                self.assertEqual(metadata.read_bytes(), damaged, "损坏的快照不应被源数据覆盖")

    def test_invalid_batch_ids_are_rejected_without_reading_outside_batch_root(self):
        for batch_id in ("../escape", ".hidden", "a/b", "a\\b", "", "x" * 129):
            with self.subTest(batch_id=batch_id), self.assertRaises(ValueError):
                self.manager.collection_batch(batch_id)

    def test_read_endpoints_and_download_reuse_files_without_writes(self):
        job_id = self.job()
        detail, _ = self.manager.submit_collection_batch(job_id)
        files = {str(path): (path.read_bytes(), path.stat().st_mtime_ns) for path in self.base.rglob("*") if path.is_file()}
        with patch.object(self.manager, "_write_json", side_effect=AssertionError("批次读取不得写文件")), \
             patch("backend.task_generation.collection_batches.write_collection_workbook", side_effect=AssertionError("下载不得重新生成 Excel")):
            self.assertEqual(self.manager.collection_batch(job_id), detail)
            self.assertEqual(len(self.manager.collection_batches()), 1)
            self.assertTrue(self.manager.collection_batch_workbook(job_id).is_file())
        after = {str(path): (path.read_bytes(), path.stat().st_mtime_ns) for path in self.base.rglob("*") if path.is_file()}
        self.assertEqual(after, files)

    def test_http_submit_list_detail_download_and_errors(self):
        app = FastAPI()
        app.include_router(router)
        job_id = self.job()
        empty = self.job([])
        running = self.job(status="running")
        prefix = "/api/task-generation"
        with patch("backend.task_generation.router.manager", self.manager), TestClient(app) as client:
            response = client.post(f"{prefix}/jobs/{job_id}/collection-batch")
            self.assertEqual(response.status_code, 201)
            detail = response.json()
            self.assertEqual(response.headers.get("Cache-Control"), "no-store")
            repeat = client.post(f"{prefix}/jobs/{job_id}/collection-batch")
            self.assertEqual(repeat.status_code, 200)
            self.assertEqual(repeat.json(), detail)
            listed = client.get(f"{prefix}/collection-batches", params={"job_id": job_id})
            self.assertEqual(listed.status_code, 200)
            self.assertEqual(listed.json()["batches"][0]["batch_id"], job_id)
            self.assertNotIn("snapshot", listed.json()["batches"][0])
            self.assertEqual(client.get(f"{prefix}/collection-batches/{job_id}").json(), detail)
            download = client.get(detail["download_url"])
            self.assertEqual(download.status_code, 200)
            self.assertEqual(download.content, self.manager.collection_batch_workbook(job_id).read_bytes())
            self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", download.headers["Content-Type"])
            self.assertIn(detail["filename"], download.headers["Content-Disposition"])
            book = load_workbook(io.BytesIO(download.content))
            self.assertEqual([cell.value for cell in book["Sheet1"][1]], EXPECTED_COLUMNS)
            book.close()
            for identifier, expected in (("missing", 404), (empty, 409), (running, 409)):
                error = client.post(f"{prefix}/jobs/{identifier}/collection-batch")
                self.assertEqual(error.status_code, expected)
                self.assertEqual(error.headers.get("Cache-Control"), "no-store")
                self.assertTrue(error.json()["detail"])
            self.assertEqual(client.get(f"{prefix}/collection-batches/missing").status_code, 404)
            self.assertEqual(client.get(f"{prefix}/collection-batches/missing/workbook").status_code, 404)

    def test_openapi_includes_batch_routes_and_structured_details(self):
        app = FastAPI()
        app.include_router(router)
        schema = app.openapi()
        prefix = "/api/task-generation"
        submit = schema["paths"][prefix + "/jobs/{job_id}/collection-batch"]["post"]
        self.assertTrue({"200", "201", "404", "409"}.issubset(submit["responses"]))
        self.assertIn("get", schema["paths"][prefix + "/collection-batches"])
        self.assertIn("get", schema["paths"][prefix + "/collection-batches/{batch_id}"])
        self.assertIn("get", schema["paths"][prefix + "/collection-batches/{batch_id}/workbook"])
        ref = submit["responses"]["201"]["content"]["application/json"]["schema"]["$ref"].split("/")[-1]
        self.assertEqual(set(schema["components"]["schemas"][ref]["properties"]), SUMMARY_FIELDS | {"snapshot"})


if __name__ == "__main__":
    unittest.main()
