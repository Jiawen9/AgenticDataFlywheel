"""Offline two-engine isolation through transfer, preprocessing, tree and release."""
from __future__ import annotations

import copy
from functools import partial
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image

from backend.batch_lifecycle import lifecycle
from backend.batch_results import current_quality_payload, current_tree_payload, merge_quality_results
from backend.collection_runs import CollectionRunStore
from backend.collection_transfer import CollectionTransferManager
from backend.data_publishing.service import DatasetReleaseRegistry
from backend.data_store import ArtifactStore
from backend.data_store.release_provenance import FrozenReleaseSources
from backend.manual_collection import ManualCollectionStore
from backend import pipeline_release as publishing
from backend.preprocessing_jobs import PreprocessingJobManager
from backend.preprocessing_service import SHEET
from backend.quality_input_builder import build_quality_workbook
from backend.stage_artifacts import fingerprint
from backend.task_generation.collection_batches import COLLECTION_COLUMNS
from backend.trajectory_context import resolve_batch_context, validate_batch_sources
from backend.tree_build_service import build_tree_run
from backend.tests.test_collection_result_batches import workbook
from backend.tests.test_collection_transfer import MockRemote
from backend.tests.test_preprocessing_jobs import CONFIG, ManualQueue
from backend.tests.test_trajectories_preprocessing import AcceptingReviewer, create_step, UI_XML
from backend.tests.test_tree_build_service import FakeAlignmentReviewer, FakeClassifier, FakeSummarizer


class EngineRemote(MockRemote):
    def get_run(self, run_id):
        return {**super().get_run(run_id), "agent_type": self.manifest["agent_type"]}


class ResultBatchDownstreamTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "platform"
        self.source = ManualCollectionStore(self.root).register("same-tasks.xlsx", workbook(), request_id="same-source")
        self.runs, self.store = CollectionRunStore(self.root), ArtifactStore(self.root)
        self.task_id = self.source["snapshot"]["tasks"][0]["task_id"]
        self.queue = ManualQueue()
        self.preprocessing = PreprocessingJobManager(self.root, source_store=self.runs, executor=self.queue,
            config_loader=lambda: copy.deepcopy(CONFIG), reviewer_factory=lambda **_: AcceptingReviewer())
        self.registry = DatasetReleaseRegistry(data_root=self.root, releases_file=self.root / "registry.json")

    def transfer(self, engine, token, color):
        metadata = {"agent_type": engine, "run_mode": "generate", "result_transport": "http",
            "filename": self.source["filename"], "vla": "http://mock.invalid/v1",
            "phone_id": "phone-a", "phone_ids": ["phone-a"], "app": "App A",
            "phone_apps": [{"phone_id": "phone-a", "app": "App A"}],
            "config": {"capture_logprobs": True, "sampling_enabled": False}}
        run, _ = self.runs.create_result_run(self.source["batch_id"], dispatch_key=engine,
            metadata=metadata, workbook_sha256=self.source["workbook_sha256"])
        remote_root = Path(self.temp.name) / "collector" / engine
        directory = remote_root / "case-A" / "same-source-trajectory"
        directory.mkdir(parents=True)
        action = {"action": "click", "coordinate": [50, 50]}
        create_step(directory, 1, action, engine + " action")
        for suffix in ("input.jpg", "input_stability.jpg"):
            Image.new("RGB", (200, 300), color).save(directory / ("step001_vla_" + suffix))
        xml = UI_XML.replace('class="android.widget.Button"', f'class="android.widget.Button" text="{engine}"')
        (directory / "step001_vla_input_stability_ui.xml").write_text(xml, encoding="utf-8")
        # Deliberately distinct old XML proves paired stability XML is preferred.
        (directory / "step001_vla_input_ui.xml").write_text(UI_XML.replace('class="android.widget.Button"',
            'class="android.widget.Button" text="old XML"'), encoding="utf-8")
        (directory / "step001_logprobs.json").write_text(json.dumps({"status": "ok",
            "response_token_ids": [token], "behavior_logprobs": [-token / 1000],
            "response_loss_mask": [1]}), encoding="utf-8")
        trajectory = {"agent_type": engine, "actions_flat": [{"global_step": 1, "action": action}]}
        (directory / "_trajectory_for_evaluate.json").write_text(json.dumps(trajectory), encoding="utf-8")
        if engine == "pipelineagent":
            (directory / "step001_vla_model_response.json").write_text(json.dumps({
                "agent_type": engine, "global_step": 1, "content": "原始Pipeline回复保持不变",
                "executed_action": action, "actual_action_coordinate_system": "pixel",
                "parsed_action": action, "parsed_summary": "pipeline action", "parse_status": "parsed",
                "repair_applied": False}, ensure_ascii=False), encoding="utf-8")
        entry = {"collection_case_id": "case-A", "source_trajectory_id": directory.name,
            "phone_id": "phone-a", "relative_dir": directory.relative_to(remote_root).as_posix(),
            "files": [{"path": path.relative_to(remote_root).as_posix(), "sha256": fingerprint(path), "size": path.stat().st_size}
                      for path in sorted(directory.iterdir())]}
        remote = EngineRemote(remote_root, run, [entry])
        remote.manifest["agent_type"] = engine
        manager = CollectionTransferManager(self.root, self.runs, remote)
        try:
            result = manager.sync(run["collection_run_id"])
            self.assertEqual(result["transfer_status"], "completed")
            self.assertEqual(manager.sync(run["collection_run_id"]), result)
            self.assertEqual(remote.download_calls, 1)
        finally:
            manager.close()
        raw = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in directory.iterdir()}
        return result, raw

    def preprocess(self, run):
        job = self.preprocessing.submit(run["batch_id"])
        self.queue.drain()
        job = self.preprocessing.get(job["job_id"])
        self.assertEqual(job["status"], "succeeded", job.get("error"))
        return resolve_batch_context(run["batch_id"], root=self.root)

    def check_assets(self, run, payload, expected_hashes, token):
        rows = payload["sheets"][SHEET]
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["agent_type"], run["agent_type"])
        self.assertEqual(row["collection_run_id"], run["collection_run_id"])
        self.assertEqual(row["task_id"], self.task_id)
        self.assertEqual(row["collection_case_id"], "case-A")
        raw_root = Path(payload["raw_root"])
        self.assertEqual(raw_root, self.runs.raw_root(run["batch_id"]))
        for field in ("image", "xml", "logprobs_path"):
            reference = Path(row[field])
            self.assertEqual(reference.parts[:2], ("runs", run["collection_run_id"]))
            path = raw_root / reference
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), expected_hashes[path.name])
        self.assertEqual(Path(row["xml"]).name, "step001_vla_input_stability_ui.xml")
        self.assertIn(run["agent_type"], (raw_root / row["xml"]).read_text(encoding="utf-8"))
        self.assertEqual(json.loads((raw_root / row["logprobs_path"]).read_text())["response_token_ids"], [token])
        return row

    def tree_and_quality(self, run, context, score):
        with patch("backend.tree_build_service.configure_reviewer_environment", return_value="fake-model"), \
             patch("backend.tree_build_service.QwenIntermediateStateClassifier", FakeClassifier), \
             patch("backend.tree_build_service.QwenStateAlignmentReviewer", FakeAlignmentReviewer):
            build_tree_run([self.task_id], job_id="tree-" + run["batch_id"], progress=lambda _: None,
                batch_id=run["batch_id"], annotation_version=context.annotation_version, data_root=self.root,
                env_path=self.root / "unused.env", quality_builder=partial(build_quality_workbook, summarizer=FakeSummarizer()))
        trees = current_tree_payload(run["batch_id"], self.root)
        row = context.payload["sheets"][SHEET][0]
        for step in trees["quality_input"]["sheets"]["Steps"]:
            identity = json.loads(step["action_input_json"])["source_identity"]
            self.assertEqual(identity["collection_run_id"], run["collection_run_id"])
            self.assertEqual(identity["trajectory_id"], row["trajectory_id"])
        # Model inference is simulated; persistence, source guards and selection are real.
        merge_quality_results(run["batch_id"], [{"task_id": self.task_id, "run_id": run["batch_id"],
            "status": "succeeded", "average_score": score,
            "evaluations": {row["trajectory_id"]: {"global_score": score}}}],
            {self.task_id: trees["tree_hashes"][self.task_id]}, {self.task_id: "test-quality-" + run["batch_id"]},
            job_id="quality-" + run["batch_id"], completed_at="2026-10-28T08:00:00Z", root=self.root)
        quality = current_quality_payload(run["batch_id"], self.root)
        self.assertEqual(set(quality["tasks"][0]["evaluations"]), {row["trajectory_id"]})
        return row

    def test_two_engines_keep_evidence_quality_candidates_and_publications_independent(self):
        collected = [self.transfer("newagent", 111, "red"), self.transfer("pipelineagent", 222, "blue")]
        runs = [item[0] for item in collected]
        contexts, rows, pipelines = [], [], []
        for (run, hashes), token, score in zip(collected, (111, 222), (4.1, 4.9)):
            context = self.preprocess(run)
            contexts.append(context)
            for stage in ("01_conversion", "02_annotation"):
                self.check_assets(run, self.store.read_payload(self.store.get(run["batch_id"], stage)), hashes, token)
            row = self.tree_and_quality(run, context, score)
            rows.append(row)
            tree = current_tree_payload(run["batch_id"], self.root)
            self.assertEqual(tree["source_annotation"]["sheets"][SHEET], context.payload["sheets"][SHEET])
            selection = publishing.build_selection(run["batch_id"], [self.task_id], "automatic", 4, root=self.root)
            self.assertEqual([item["trajectory_id"] for item in selection["candidates"]], [row["trajectory_id"]])
            self.assertEqual(selection["selected_trajectories"], {self.task_id: row["trajectory_id"]})
            pipelines.append({"pipeline_id": "publish-" + run["batch_id"], "name": run["agent_type"],
                              "batch_id": run["batch_id"], "mode": "automatic", "selection": selection})
        self.assertNotEqual(rows[0]["trajectory_id"], rows[1]["trajectory_id"])
        self.assertEqual(rows[0]["source_trajectory_id"], rows[1]["source_trajectory_id"])
        peer_before = {ref["stage"]: ref["content_hash"] for ref in self.store.list(runs[1]["batch_id"])}
        first_release = publishing.publish_pipeline(pipelines[0], root=self.root, registry=self.registry)
        self.assertEqual(lifecycle(runs[0]["batch_id"], self.root)["status"], "published")
        self.assertEqual(lifecycle(runs[1]["batch_id"], self.root)["status"], "active")
        self.assertEqual(lifecycle(self.source["batch_id"], self.root)["status"], "active")
        self.assertEqual(peer_before, {ref["stage"]: ref["content_hash"] for ref in self.store.list(runs[1]["batch_id"])})
        second_release = publishing.publish_pipeline(pipelines[1], root=self.root, registry=self.registry)
        self.assertNotEqual(first_release["release_id"], second_release["release_id"])
        self.assertEqual(len(self.registry.list_releases()), 2)
        for release, run, row in zip((first_release, second_release), runs, rows):
            self.assertEqual(release["batch_ids"], [run["batch_id"]])
            self.assertEqual(release["selection"]["selected_trajectories"], {self.task_id: row["trajectory_id"]})
            child = self.runs.batch(run["batch_id"])
            collection = self.store.get(run["batch_id"], "00_collection")
            frozen = FrozenReleaseSources(self.root, release)
            frozen_collection = json.loads(frozen.resolve_file(collection, "result.json").read_text(encoding="utf-8"))
            self.assertEqual(frozen_collection["task_source_batch_id"], self.source["batch_id"])
            self.assertEqual(frozen_collection["snapshot"]["tasks"][0]["source_values"],
                             self.source["snapshot"]["tasks"][0]["source_values"])
            self.assertEqual(set(frozen_collection["snapshot"]["tasks"][0]["source_values"]), set(COLLECTION_COLUMNS))
            source_file = self.store.resolve_file(collection, child["filename"])
            book = load_workbook(source_file, read_only=True)
            try:
                values = list(book.active.values)
                self.assertEqual(list(values[0]), COLLECTION_COLUMNS)
                self.assertEqual(values[1][12:17], ("出现结果", "复杂", "点击", "=literal SOP", "变量说明"))
                self.assertEqual(len(values), 2)
            finally:
                book.close()
            self.assertEqual(child["snapshot"]["tasks"][0], self.source["snapshot"]["tasks"][0])
            export_file, _ = self.registry.excel_file(release["release_id"], 0)
            exported = load_workbook(export_file, read_only=True)
            try:
                values = list(exported.active.values)
                columns = list(values[0])
                self.assertEqual(len(values), 2)
                self.assertEqual(values[1][columns.index("logprobs_path")], row["logprobs_path"])
                self.assertEqual(values[1][columns.index("image")], row["image"])
            finally:
                exported.close()

    def test_corrupt_raw_evidence_in_one_result_does_not_invalidate_the_other(self):
        first, _ = self.transfer("newagent", 111, "red")
        second, _ = self.transfer("pipelineagent", 222, "blue")
        first_context, second_context = self.preprocess(first), self.preprocess(second)
        row = first_context.payload["sheets"][SHEET][0]
        image = first_context.raw_root / row["image"]
        image.write_bytes(image.read_bytes() + b"changed")
        with self.assertRaisesRegex(ValueError, "校验值变化"):
            validate_batch_sources(first_context)
        validate_batch_sources(second_context)
        self.assertEqual(self.runs.ready_input(second["batch_id"])["runs"][0]["collection_run_id"], second["collection_run_id"])
        self.assertEqual(resolve_batch_context(second["batch_id"], root=self.root).annotation_version,
                         second_context.annotation_version)


if __name__ == "__main__":
    unittest.main()
