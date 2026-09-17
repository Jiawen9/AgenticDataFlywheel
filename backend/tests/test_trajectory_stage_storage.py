from __future__ import annotations

import json
import asyncio
import importlib
import sys
import tempfile
import unittest
from contextlib import ExitStack
from functools import partial
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace

from openpyxl import load_workbook

from backend.data_store import ArtifactStore, RecordStore
from backend.stage_artifacts import (observation_payload, publish_workbook, quality_tables,
    read_workbook_payload, store_root, write_payload_workbook, write_sidecar)
from backend.trajectories_tree.tree_builder import load_trajectories
from backend.trajectories_tree.intermediate_state_classifier import IntermediateStateResult
from backend.trajectories_preprocessing import run_pipeline
from backend.tree_build_service import build_tree_run
from backend.quality_input_builder import build_quality_workbook
from backend.quality_data import quality_task
from backend import trajectory_data
from backend.tests.test_tree_build_service import (prepare_source, FakeClassifier,
    FakeAlignmentReviewer, FakeSummarizer)
from backend.tests import test_trajectories_preprocessing as preprocessing_fixture


class TrajectoryStageStorageTests(unittest.TestCase):
    def test_preprocessing_keeps_two_versions_without_extra_model_calls(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_workbook, trajectory_root = preprocessing_fixture.WorkbookAnnotationTests()._create_workbook_fixture(root)
            # The fixture return is (workbook, trajectory_root).
            output = root / "conversion.xlsx"
            annotation = root / "annotation.xlsx"
            with patch("backend.trajectories_preprocessing.configure_reviewer_environment", return_value="mock"), \
                 patch("backend.trajectories_preprocessing.QwenBoxReviewer", return_value=preprocessing_fixture.AcceptingReviewer()):
                result = run_pipeline(source=trajectory_root, export_output=output, annotated_output=annotation,
                    env_file=root / ".env", max_review_rounds=2, batch_id="batch-one", data_root=root / "data")
            self.assertEqual(result["batch_id"], "batch-one")
            artifacts = ArtifactStore(root / "data").list(batch_id="batch-one")
            self.assertEqual({item["stage"] for item in artifacts}, {"01_conversion", "02_annotation"})
            for artifact in artifacts:
                store = ArtifactStore(root / "data")
                value = json.loads(store.resolve_file(artifact, "result.json").read_text(encoding="utf-8"))
                self.assertEqual(sum(len(rows) for rows in value["sheets"].values()), 4)
                self.assertTrue(any(item["kind"] == "excel" for item in artifact["files"]))
            conversion_ref, annotation_ref = result["artifacts"]
            self.assertEqual(annotation_ref["source_refs"][0]["version"], conversion_ref["version"])

    def test_json_is_authoritative_when_excel_is_modified_or_deleted(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _, workbook = prepare_source(root)
            first = publish_workbook(workbook, batch_id="batch", stage="02_annotation", data_root=root / "data")
            with patch("backend.stage_artifacts.load_workbook", side_effect=AssertionError("Excel was parsed")):
                trajectories = load_trajectories(workbook, None)
            self.assertEqual(len(trajectories), 2)
            changed = load_workbook(workbook)
            changed.active["E2"] = "changed summary"
            changed.save(workbook)
            changed.close()
            self.assertEqual(load_trajectories(workbook, None)[0][1][0].summary, "TASK-A summary")
            workbook.unlink()
            self.assertEqual(load_trajectories(workbook, None)[0][1][0].summary, "TASK-A summary")
            frozen = json.loads(ArtifactStore(root / "data").resolve_file(first, "result.json").read_text(encoding="utf-8"))
            self.assertEqual(next(iter(frozen["sheets"].values()))[0]["summary"], "TASK-A summary")

    def test_observation_keeps_ignored_advert_and_full_reason(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            _, workbook = prepare_source(root)
            trajectory, steps = load_trajectories(workbook, None)[0]
            step = steps[0]
            step.apply_classification(IntermediateStateResult(True, "advertisement", .98, "广告浮层", "{}", False, "显示广告弹窗"), .8)
            step.counted_in_tree = False
            step.effective_intermediate = True
            payload, rows = observation_payload({"TASK-A": [(trajectory, steps)]}, .8)
            self.assertEqual(len(payload["trajectories"][0]["steps"]), 1)
            self.assertEqual(rows[0]["Observation"], "显示广告弹窗")
            self.assertEqual(rows[0]["中间态类别"], "advertisement")
            self.assertEqual(rows[0]["置信度"], .98)
            self.assertEqual(rows[0]["判断原因"], "广告浮层")
            self.assertFalse(rows[0]["计入树"])

    def test_tree_pins_annotation_and_reuses_batch_with_all_stage_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trajectory_root, workbook = prepare_source(root)
            data_root = root / "data"
            original = publish_workbook(workbook, batch_id="pipeline", stage="02_annotation", data_root=data_root)
            with patch("backend.tree_build_service.configure_reviewer_environment", return_value="fake-model"), \
                 patch("backend.tree_build_service.QwenIntermediateStateClassifier", FakeClassifier), \
                 patch("backend.tree_build_service.QwenStateAlignmentReviewer", FakeAlignmentReviewer):
                run_id, manifest = build_tree_run(["TASK-A", "TASK-B"], job_id="j", progress=lambda _: None,
                    xlsx_path=workbook, trajectory_root=trajectory_root, runs_dir=root / "runs", data_root=data_root,
                    env_path=root / ".env", quality_builder=partial(build_quality_workbook, summarizer=FakeSummarizer()))
            self.assertEqual(manifest["batch_id"], "pipeline")
            self.assertEqual(manifest["source_refs"][0]["version"], original["version"])
            self.assertEqual({a["stage"] for a in manifest["artifacts"]}, {"03_observation", "04_tree"})
            self.assertTrue((root / "runs" / run_id / manifest["source_annotated_file"]).is_file())
            self.assertTrue((root / "runs" / run_id / manifest["quality_input_json"]).is_file())
            observation = next(a for a in manifest["artifacts"] if a["stage"] == "03_observation")
            book = load_workbook(ArtifactStore(data_root).resolve_file(observation, "result.xlsx"), read_only=True)
            self.assertEqual(book.active.max_row, 3)
            book.close()
            workbook.unlink()
            self.assertEqual(len(load_trajectories(root / "runs" / run_id / manifest["source_annotated_json"], None)), 2)

    def test_quality_artifact_has_three_sheets_and_sqlite_current_view(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            task = {"task_id": "A", "evaluations": {"A-1": {"global_score": 4.5, "passed_threshold": True,
                    "step_evaluations": [{"step_id": 1, "dimension_scores": [{"dimension_name": "accuracy", "score": 4.5, "reasoning": "正确"}]}]}},
                    "rubric": {"dimensions": [{"name": "accuracy", "description": "动作准确"}]}}
            artifact = ArtifactStore(root).publish("b", "05_quality", {"tasks": [task]}, tables=quality_tables([task]))
            book = load_workbook(ArtifactStore(root).resolve_file(artifact, "result.xlsx"), read_only=True)
            self.assertEqual(book.sheetnames, ["轨迹汇总", "步骤评分", "评分标准"])
            self.assertEqual([sheet.max_row for sheet in book], [2, 2, 2])
            book.close()
            results_root = root / "results"
            RecordStore(store_root(results_root)).put("quality_results", "run:A", task)
            self.assertEqual(quality_task("run", "A", results_root)["evaluations"]["A-1"]["global_score"], 4.5)

    def test_custom_root_cannot_escape_test_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.assertTrue(store_root(root).is_relative_to(root))
            self.assertTrue(store_root(root / "out.xlsx").is_relative_to(root))

    def test_old_sentinels_are_not_scanned_or_merged(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            legacy = root / "backend_workspace"
            new_runs = root / "data" / "system" / "trajectory_tree_runs"
            for run_id, target in [("old", legacy / "trajectory_tree_runs"), ("new", new_runs)]:
                (target / run_id).mkdir(parents=True)
                (target / run_id / "manifest.json").write_text(json.dumps({"run_id": run_id, "completed_at": run_id}), encoding="utf-8")
            sentinel = legacy / "trajectory_tree_runs" / "old" / "manifest.json"
            before = sentinel.read_bytes()
            with patch.object(trajectory_data, "TREE_RUNS_DIR", new_runs), patch.object(trajectory_data, "LEGACY_ROOT", legacy, create=True):
                self.assertEqual({row["run_id"] for row in trajectory_data.list_tree_runs(new_runs)}, {"new"})
                self.assertEqual(trajectory_data.resolve_tree_run_dir("old", new_runs), (new_runs / "old").resolve())
            new_workbook = root / "current" / "annotated.xlsx"
            payload = {"columns": {"Steps": ["文件夹名", "image", "xml", "action", "summary", "actions_box"]},
                "sheets": {"Steps": [{"文件夹名": "OLD-1", "image": "OLD/OLD-1/step001_vla_input.jpg",
                                       "xml": "x", "action": '{"action":"wait"}', "summary": "old", "actions_box": ""}]}}
            write_payload_workbook(legacy / "annotated_trajectories.xlsx", payload)
            payload["sheets"]["Steps"][0].update({"文件夹名": "NEW-1", "image": "NEW/NEW-1/step001_vla_input.jpg"})
            write_payload_workbook(new_workbook, payload)
            write_sidecar(new_workbook, payload)
            with patch.object(trajectory_data, "ANNOTATED_XLSX", new_workbook), patch.object(trajectory_data, "LEGACY_ROOT", legacy, create=True):
                self.assertEqual(set(trajectory_data.load_trajectory_index(new_workbook)), {"NEW"})
            self.assertEqual(sentinel.read_bytes(), before)

    def test_missing_and_corrupt_json_fail_even_when_excel_exists(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            trajectory_root, workbook = prepare_source(root)
            workbook.with_suffix(".json").unlink()
            with self.assertRaisesRegex(FileNotFoundError, "JSON snapshot"):
                read_workbook_payload(workbook)
            with self.assertRaisesRegex(FileNotFoundError, "JSON"):
                build_tree_run(["TASK-A"], job_id="missing", progress=lambda _: None,
                    xlsx_path=workbook, trajectory_root=trajectory_root, runs_dir=root / "runs")
            self.assertEqual(len(load_trajectories(workbook, None, allow_excel_import=True)), 2)
            workbook.with_suffix(".json").write_text("broken JSON", encoding="utf-8")
            with self.assertRaises(json.JSONDecodeError):
                read_workbook_payload(workbook)

    def test_jobs_do_not_import_or_write_json_mirrors(self):
        from backend.tree_build_jobs import TreeBuildJobManager
        from backend.quality_jobs import QualityJobManager
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, manager_type, namespace in [
                    ("tree", TreeBuildJobManager, "tree_jobs"),
                    ("quality", QualityJobManager, "quality_jobs")]:
                job_root = root / name
                job_root.mkdir()
                sentinel = job_root / "old.json"
                sentinel.write_text('{"job_id":"old","status":"running"}', encoding="utf-8")
                before = sentinel.read_bytes()
                manager = manager_type(job_root, runner=lambda *args, **kwargs: {})
                try:
                    self.assertIsNone(manager.get("old"))
                    manager._write({"job_id": "new", "status": "succeeded"})
                    self.assertEqual(manager.get("new")["status"], "succeeded")
                    self.assertFalse((job_root / "new.json").exists())
                    self.assertEqual(sentinel.read_bytes(), before)
                finally:
                    manager.shutdown()

    def test_quality_worker_uses_json_without_excel_and_publishes_frozen_results(self):
        rubric_module = Path(__file__).resolve().parents[1] / "DevelopRubrics"
        with patch.object(sys, "path", [str(rubric_module), *sys.path]):
            worker = importlib.import_module("backend.DevelopRubrics.quality_job_runner")
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            root = Path(directory)
            trajectory_root, workbook = prepare_source(root)
            data_root = root / "data"
            runs = root / "runs"
            with patch("backend.tree_build_service.configure_reviewer_environment", return_value="fake-model"), \
                 patch("backend.tree_build_service.QwenIntermediateStateClassifier", FakeClassifier), \
                 patch("backend.tree_build_service.QwenStateAlignmentReviewer", FakeAlignmentReviewer):
                run_id, manifest = build_tree_run(["TASK-A"], job_id="tree", progress=lambda _: None,
                    xlsx_path=workbook, trajectory_root=trajectory_root, runs_dir=runs, data_root=data_root,
                    env_path=root / ".env", quality_builder=partial(build_quality_workbook, summarizer=FakeSummarizer()))
            (runs / run_id / manifest["quality_input_file"]).unlink()
            for name, value in {"DATA_ROOT": data_root, "TREE_RUNS": runs,
                                "RESULTS_ROOT": root / "results", "CHECKPOINT_ROOT": root / "checkpoints"}.items():
                stack.enter_context(patch.object(worker, name, value))
            stack.enter_context(patch.object(worker, "configure_model_environment"))
            stack.enter_context(patch.object(worker, "quality_manifest", return_value={"tasks": []}))
            stack.enter_context(patch.object(worker, "_matching_rubric", return_value=root / "rubric.json"))
            stack.enter_context(patch.object(worker.GEN, "load_config", return_value={}))
            rubric = SimpleNamespace(model_dump_json=lambda: json.dumps({"dimensions": [{"name": "正确性"}]}))
            stack.enter_context(patch.object(worker.EVAL, "_load_rubric", return_value=rubric))
            stack.enter_context(patch.object(worker.EVAL, "_build_pipeline", return_value=object()))
            stack.enter_context(patch.object(worker.EVAL, "_evaluation_settings", return_value={}))
            stack.enter_context(patch.object(worker.EVAL, "_settings_signature", return_value="test"))
            stack.enter_context(patch.object(worker.EVAL, "initialize_evaluations_jsonl"))
            stack.enter_context(patch.object(worker.EVAL, "load_existing_evaluations_jsonl", return_value={}))

            async def evaluate(**kwargs):
                trajectory = kwargs["trajectories"][0]
                self.assertEqual(trajectory.steps[0].observation, "页面保持稳定")
                value = {"trajectory_id": trajectory.trajectory_id, "global_score": 4.5, "passed_threshold": True,
                         "step_evaluations": [{"step_id": 1, "dimension_scores": [{"score": 4.5, "reasoning": "正确"}]}]}
                evaluation = SimpleNamespace(trajectory_id=trajectory.trajectory_id, global_score=4.5,
                    passed_threshold=True, model_dump_json=lambda **_: json.dumps(value))
                kwargs["on_trajectory_complete"](evaluation)
                return SimpleNamespace(all_evaluations=[evaluation])

            stack.enter_context(patch.object(worker.EVAL, "evaluate_run_incrementally", side_effect=evaluate))
            result = asyncio.run(worker.run(run_id, ["TASK-A"], "quality"))
            self.assertEqual(result["artifact"]["stage"], "05_quality")
            saved = RecordStore(data_root).get("quality_results", f"{run_id}:TASK-A")
            self.assertEqual(saved["average_score"], 4.5)
            frozen = ArtifactStore(data_root).resolve_file(result["artifact"], "result.xlsx")
            self.assertTrue(frozen.is_file())
            previous_manifest = RecordStore(data_root).get("quality_manifests", run_id)
            original_save = RecordStore._save

            def fail_manifest(connection, namespace, key, value, revision):
                if namespace == "quality_manifests":
                    raise RuntimeError("synthetic transaction failure")
                return original_save(connection, namespace, key, value, revision)

            with patch.object(worker, "quality_manifest", return_value=previous_manifest), \
                 patch.object(RecordStore, "_save", side_effect=fail_manifest):
                with self.assertRaisesRegex(RuntimeError, "synthetic transaction failure"):
                    asyncio.run(worker.run(run_id, ["TASK-A"], "rollback"))
            self.assertEqual(RecordStore(data_root).get("quality_results", f"{run_id}:TASK-A"), saved)
            self.assertEqual(RecordStore(data_root).get("quality_manifests", run_id), previous_manifest)

            self.assertFalse((root / "results" / run_id).exists())




if __name__ == "__main__":
    unittest.main()
