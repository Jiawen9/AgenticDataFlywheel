from __future__ import annotations

import copy
import hashlib
import io
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from backend.collection_runs import CollectionRunError, CollectionRunStore
from backend.collection_result_batches import CollectionResultBatchStore
from backend.data_store import RecordStore
from backend.data_store.artifacts import ArtifactStore
from backend.manual_collection import ManualCollectionStore
from backend.task_generation.collection_batches import COLLECTION_COLUMNS, payload_digest
from backend.tests.test_collection_runs import raw_trajectory, seed_batch


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.append(COLLECTION_COLUMNS)
    for case, app in (("case-A", "App A"), ("case-B", "App B")):
        row = [case, "单APP", app, "一级", "二级", "三级", "打开 " + app,
               "困难", "页面标题", 5, "手工", "测试者", "出现结果", "复杂", "点击",
               "=literal SOP", "变量说明"]
        sheet.append(row)
        sheet.cell(sheet.max_row, 16).data_type = "s"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class CollectionResultBatchTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / "data"
        self.source = ManualCollectionStore(self.root).register("tasks.xlsx", workbook(), request_id="upload")
        self.runs = CollectionRunStore(self.root)
        self.metadata = {"agent_type": "newagent", "run_mode": "generate", "result_transport": "http",
            "filename": self.source["filename"], "vla": "http://mock.invalid/v1",
            "phone_id": "", "app": "", "phone_ids": ["phone-a"],
            "phone_apps": [{"phone_id": "phone-a", "app": "App A", "status": "空闲"}],
            "config": {"sampling_enabled": False, "capture_logprobs": False}}

    def create(self, request="request-1", metadata=None):
        return self.runs.create_result_run(self.source["batch_id"], dispatch_key=request,
            metadata=metadata or self.metadata, workbook_sha256=self.source["workbook_sha256"])

    def publish_marker(self, batch):
        RecordStore(self.root).put("batch_lifecycle", batch,
            {"batch_id": batch, "status": "published", "release_id": "release-test"})

    def test_engines_and_repeated_collection_own_distinct_batches_with_shared_source(self):
        first, _ = self.create()
        second, _ = self.create("pipeline", {**self.metadata, "agent_type": "pipelineagent"})
        third, _ = self.create("again")
        self.assertEqual(len({run["batch_id"] for run in (first, second, third)}), 3)
        self.assertEqual(len({run["collection_run_id"] for run in (first, second, third)}), 3)
        self.assertEqual(len({run["output_dir"] for run in (first, second, third)}), 3)
        self.assertEqual({run["task_source_batch_id"] for run in (first, second, third)}, {self.source["batch_id"]})
        self.assertEqual({run["batch_tasks"]["case-A"]["task_id"] for run in (first, second, third)},
                         {self.source["snapshot"]["tasks"][0]["task_id"]})
        self.assertEqual(self.runs.list_runs(self.source["batch_id"]), [])

    def test_selected_app_preserves_all_seventeen_columns_and_source_identity(self):
        run, _ = self.create()
        child = self.runs.batch(run["batch_id"], require_workbook=True)
        self.assertEqual(child["batch_role"], "collection_result")
        self.assertEqual(child["kind"], "manual_collection")
        self.assertEqual(child["task_count"], 1)
        self.assertEqual(child["snapshot"]["tasks"], self.source["snapshot"]["tasks"][:1])
        self.assertEqual(child["task_source_snapshot_sha256"], payload_digest(self.source["snapshot"]))
        book = load_workbook(child["workbook_path"], data_only=False)
        try:
            values = list(book.active.values)
            self.assertEqual(list(values[0]), COLLECTION_COLUMNS)
            self.assertEqual(len(values), 2)
            self.assertEqual(values[1][7:], ("困难", "页面标题", 5, "手工", "测试者", "出现结果", "复杂", "点击", "=literal SOP", "变量说明"))
            self.assertEqual(book.active.cell(2, 16).data_type, "s")
        finally:
            book.close()

    def test_same_app_multiple_phones_share_one_selected_task_without_changing_associations(self):
        metadata = {**self.metadata, "phone_ids": ["phone-a", "phone-b"], "phone_apps": [
            {"phone_id": phone, "app": "App A"} for phone in ("phone-a", "phone-b")]}
        run, _ = self.create(metadata=metadata)
        child = self.runs.batch(run["batch_id"])
        self.assertEqual(child["task_count"], 1)
        self.assertEqual(child["execution_metadata"]["phone_apps"], metadata["phone_apps"])

    def test_generated_tasks_keep_generated_identity_and_kind(self):
        source, path = seed_batch(self.root, "generated")
        metadata = {**self.metadata, "filename": source["filename"],
                    "phone_apps": [{"phone_id": "phone-a", "app": "Demo"}]}
        run, _ = self.runs.create_result_run("generated", dispatch_key="one", metadata=metadata,
                                            workbook_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
        child = self.runs.batch(run["batch_id"])
        self.assertEqual(child["kind"], "task_generation")
        self.assertEqual(child["source_job_id"], "generated")
        self.assertEqual(child["snapshot"]["job_id"], "generated")
        self.assertEqual(child["snapshot"]["tasks"], source["snapshot"]["tasks"])
        self.assertEqual(run["batch_tasks"]["task-one"]["source_result_id"], "result-one")

    def test_concurrent_request_is_one_batch_and_one_run_and_config_changes_conflict(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda _: self.create(), range(4)))
        self.assertEqual(sum(created for _, created in results), 1)
        self.assertEqual(len(self.runs.result_batches.list()), 1)
        self.assertEqual(len(self.runs.list_runs()), 1)
        original = results[0][0]
        self.assertEqual(self.create(), (original, False))
        for change in ({"agent_type": "pipelineagent"}, {"vla": "http://other.invalid"}, {"app": "App B"}):
            with self.subTest(change=change), self.assertRaisesRegex(CollectionRunError, "request_id"):
                self.create(metadata={**self.metadata, **change})
        self.assertEqual(self.runs.get(original["collection_run_id"]), original)

    def test_zero_matching_tasks_or_wrong_upload_digest_create_no_batch_or_run(self):
        with self.assertRaisesRegex(CollectionRunError, "没有匹配"):
            self.create(metadata={**self.metadata, "phone_apps": [{"phone_id": "phone-a", "app": "absent"}]})
        with self.assertRaisesRegex(CollectionRunError, "上传文件"):
            self.runs.create_result_run(self.source["batch_id"], dispatch_key="bad", metadata=self.metadata,
                                       workbook_sha256="0" * 64)
        self.assertEqual(self.runs.result_batches.list(), [])
        self.assertEqual(self.runs.list_runs(), [])

    def test_frozen_child_and_request_replay_do_not_depend_on_parent_files(self):
        run, _ = self.create()
        parent_path = Path(self.source["workbook_path"])
        parent_path.unlink()
        child = self.runs.batch(run["batch_id"], require_workbook=True)
        self.assertTrue(Path(child["workbook_path"]).is_file())
        self.assertEqual(self.create(), (run, False))
        with self.assertRaises(CollectionRunError):
            self.create("new-request")

    def test_source_publication_and_peer_publication_do_not_prevent_new_results(self):
        self.publish_marker(self.source["batch_id"])
        first, _ = self.create()
        self.publish_marker(first["batch_id"])
        self.assertEqual(self.create(), (first, False))
        second, created = self.create("new-request")
        self.assertTrue(created)
        self.assertNotEqual(first["batch_id"], second["batch_id"])
        with self.assertRaisesRegex(CollectionRunError, "单次运行"):
            self.runs.create(first["batch_id"], dispatch_key="append", metadata=self.metadata)

    def test_result_batch_only_accepts_its_reserved_run_and_cannot_be_a_new_source(self):
        child, _ = self.runs.prepare_result_batch(self.source["batch_id"], dispatch_key="reserved", metadata=self.metadata)
        with self.assertRaisesRegex(CollectionRunError, "单次运行"):
            self.runs.create(child["batch_id"], dispatch_key="different", metadata=self.metadata)
        with self.assertRaisesRegex(CollectionRunError, "原任务来源"):
            self.runs.create_result_run(child["batch_id"], dispatch_key="different", metadata=self.metadata)
        run, created = self.runs.create(child["batch_id"], dispatch_key=child["execution_request_id"],
                                       metadata=child["execution_metadata"], workbook_sha256=child["workbook_sha256"])
        self.assertTrue(created)
        self.assertEqual(run["collection_run_id"], child["collection_run_id"])

    def test_legacy_request_replay_stays_on_original_batch_without_creating_child(self):
        previous, _ = self.runs.create(self.source["batch_id"], dispatch_key="request-1", metadata=self.metadata,
                                      workbook_sha256=self.source["workbook_sha256"])
        self.publish_marker(self.source["batch_id"])
        self.assertEqual(self.create(), (previous, False))
        self.assertEqual(self.runs.result_batches.list(), [])
        with self.assertRaisesRegex(CollectionRunError, "request_id"):
            self.create(metadata={**self.metadata, "agent_type": "pipelineagent"})

    def test_artifact_commit_failure_is_retryable_and_does_not_leave_partial_batch(self):
        original = RecordStore.put_many
        def fail_result_commit(store, entries, *args, **kwargs):
            if any(entry["namespace"] == "collection_result_batches" for entry in entries):
                raise OSError("simulated commit failure")
            return original(store, entries, *args, **kwargs)
        with patch.object(RecordStore, "put_many", fail_result_commit), self.assertRaises(OSError):
            self.create()
        self.assertEqual(self.runs.result_batches.list(), [])
        self.assertEqual(self.runs.list_runs(), [])
        run, created = self.create()
        self.assertTrue(created)
        self.assertTrue(Path(self.runs.batch(run["batch_id"], require_workbook=True)["workbook_path"]).is_file())

    def test_completed_inputs_remain_separate_and_reject_foreign_run_registration(self):
        first, _ = self.create()
        second, _ = self.create("request-2")
        for run in (first, second):
            self.runs.complete(run["collection_run_id"], {"batch_id": run["batch_id"],
                "trajectories": [raw_trajectory(run, "case-A")]})
        ready = [self.runs.ready_input(run["batch_id"]) for run in (first, second)]
        self.assertEqual([len(value["trajectories"]) for value in ready], [1, 1])
        self.assertNotEqual(ready[0]["raw_root"], ready[1]["raw_root"])
        self.assertEqual(ready[0]["task_source_batch_id"], ready[1]["task_source_batch_id"])
        forged = copy.deepcopy(first)
        forged.update(collection_run_id="cr_foreign", run_id="cr_foreign",
                      output_dir=str(self.runs.run_root(first["batch_id"], "cr_foreign")))
        self.runs.records.put("collection_runs", "cr_foreign", forged)
        with self.assertRaisesRegex(CollectionRunError, "未绑定"):
            self.runs.ready_input(first["batch_id"])


if __name__ == "__main__":
    unittest.main()
