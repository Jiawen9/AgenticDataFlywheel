"""Regression coverage for independent manual Thought / Summary edits."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from backend.trajectory_correction import draft_store, service
from backend.trajectory_correction.exporter import (
    export_full_dataset_workbook,
    export_session_workbook,
)
from backend.trajectory_correction.workbook import load_snapshot


def action_hash(action):
    return hashlib.sha256(
        json.dumps(action, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


class CorrectionTextEditTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "annotated.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["文件夹名", "image", "action", "summary", "thought", "actions_box"])
        for step in (1, 2):
            sheet.append([
                "TASK-1", f"TASK/TASK-1/step{step:03d}_vla_input.jpg",
                '{"action":"wait"}', f"original summary {step}", "", "wait()",
            ])
        workbook.save(self.source)
        workbook.close()
        self.source_bytes = self.source.read_bytes()
        self.snapshot = load_snapshot(self.source)
        self.group_id = self.snapshot["groups"][0]["group_id"]
        self.session_id = "abcdef1234567890"
        self.action = {"action": "type", "text": "corrected"}
        self.generated = {
            "thought": "generated thought", "summary": "generated summary",
            "content_tag": "thought_summary", "action_hash": action_hash(self.action),
            "bbox_hash": "original-bbox-hash", "model": "test-model",
            "generated_at": "2026-09-08T00:00:00+00:00",
        }
        sessions_dir = self.root / "sessions"
        sessions_dir.mkdir()
        for patcher in (
            patch.object(draft_store, "CORRECTION_SESSIONS_DIR", sessions_dir),
            patch.object(draft_store, "ensure_correction_dirs"),
            patch.object(service, "_snapshot", side_effect=lambda _: deepcopy(self.snapshot)),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        self.reset_session()

    def reset_session(self):
        draft_store.save_session({
            "session_id": self.session_id,
            "row_edits": {"2": {
                "actions": json.dumps(self.action), "original_actions": '{"action":"wait"}',
            }},
            "cot": {
                "2": deepcopy(self.generated),
                "3": {**self.generated, "action_hash": action_hash({"action": "wait"})},
            },
            "group_exports": {self.group_id: True},
        })

    def save_text(self, **values):
        return service.patch_row(self.session_id, 2, values)["row"]

    def assert_values(self, thought, summary, manual_fields):
        # Exercise the public reads after loading the persisted JSON again.
        session = draft_store.load_session(self.session_id)
        self.assertEqual(session["cot"]["2"], self.generated)
        self.assertEqual(session["cot"]["3"], {
            **self.generated, "action_hash": action_hash({"action": "wait"}),
        })
        group = service.get_group(self.session_id, self.group_id)
        row = group["rows"][0]
        cot_row = service.get_cot(self.session_id)["groups"][0]["rows"][0]
        for current in (row, cot_row):
            self.assertEqual(current["thought"], thought)
            self.assertEqual(current["summary"], summary)
            for field in ("thought", "summary"):
                self.assertEqual(current[field + "_source"], "manual" if field in manual_fields else "generated")
        self.assertEqual(row["cot_status"], "manual")
        self.assertEqual(cot_row["status"], "manual")
        self.assertEqual(row["cot_generated_at"], self.generated["generated_at"])
        self.assertEqual(row["cot_action_hash"], self.generated["action_hash"])
        self.assertEqual(group["rows"][1]["thought"], "generated thought")

        for exporter in (export_session_workbook, export_full_dataset_workbook):
            with self.subTest(exporter=exporter.__name__):
                result = exporter(
                    workbook_path=self.source, snapshot=self.snapshot, session=session,
                    output_dir=self.root / "exports", export_id="text-edits",
                )
                exported = load_workbook(self.root / "exports" / result["filename"], read_only=True)
                try:
                    sheet = exported.active
                    headers = [cell.value for cell in sheet[1]]
                    # Excel stores an explicitly empty string as an empty cell.
                    self.assertEqual(sheet.cell(2, headers.index("thought") + 1).value or "", thought)
                    self.assertEqual(sheet.cell(2, headers.index("summary") + 1).value or "", summary)
                    expected_rows = 2 if exporter is export_session_workbook else 3
                    self.assertEqual(sheet.max_row, expected_rows)
                    if exporter is export_full_dataset_workbook:
                        self.assertEqual(sheet.cell(3, headers.index("thought") + 1).value, "generated thought")
                finally:
                    exported.close()
        self.assertEqual(self.source.read_bytes(), self.source_bytes)

    def test_edit_summary_preserves_generated_thought_when_original_is_empty(self):
        row = self.save_text(summary="manual summary")
        self.assertEqual(row["thought"], "generated thought")
        self.assert_values("generated thought", "manual summary", {"summary"})

    def test_edit_thought_preserves_generated_summary(self):
        row = self.save_text(thought="manual thought")
        self.assertEqual(row["summary"], "generated summary")
        self.assert_values("manual thought", "generated summary", {"thought"})

    def test_sequential_edits_keep_both_latest_manual_values(self):
        self.save_text(summary="first summary")
        self.save_text(thought="manual thought")
        self.save_text(summary="latest summary")
        self.assert_values("manual thought", "latest summary", {"summary", "thought"})

    def test_clearing_each_field_is_an_explicit_override(self):
        for field in ("summary", "thought"):
            with self.subTest(field=field):
                self.reset_session()
                self.save_text(**{field: ""})
                self.assertIn(field, draft_store.load_session(self.session_id)["row_edits"]["2"])
                self.assert_values(
                    "" if field == "thought" else "generated thought",
                    "" if field == "summary" else "generated summary", {field},
                )

    def test_saving_original_text_does_not_reactivate_generated_value(self):
        self.save_text(summary="original summary 1")
        self.assert_values("generated thought", "original summary 1", {"summary"})

    def test_both_fields_can_be_saved_in_one_patch(self):
        self.save_text(thought="manual thought", summary="manual summary")
        self.assert_values("manual thought", "manual summary", {"summary", "thought"})

    def test_stale_generated_text_is_not_used_by_reads_or_exports(self):
        session = draft_store.load_session(self.session_id)
        session["cot"]["2"]["action_hash"] = "stale"
        draft_store.save_session(session)
        row = self.save_text(summary="manual summary")
        self.assertEqual(row["thought"], "")
        self.assertEqual(row["thought_source"], "original")
        session = draft_store.load_session(self.session_id)
        self.assertEqual(session["cot"]["2"]["action_hash"], "stale")
        for exporter in (export_session_workbook, export_full_dataset_workbook):
            result = exporter(
                workbook_path=self.source, snapshot=self.snapshot, session=session,
                output_dir=self.root / "exports", export_id="stale",
            )
            workbook = load_workbook(self.root / "exports" / result["filename"], read_only=True)
            try:
                headers = [cell.value for cell in workbook.active[1]]
                self.assertFalse(workbook.active.cell(2, headers.index("thought") + 1).value)
                self.assertEqual(workbook.active.cell(2, headers.index("summary") + 1).value, "manual summary")
            finally:
                workbook.close()

    def test_action_edit_still_invalidates_generated_cot(self):
        self.save_text(summary="manual summary")
        row = service.patch_row(self.session_id, 2, {"actions": '{"action":"open","text":"app"}'})["row"]
        session = draft_store.load_session(self.session_id)
        self.assertNotIn("2", session["cot"])
        self.assertIn("3", session["cot"])
        self.assertEqual(row["summary"], "manual summary")
        self.assertEqual(row["thought"], "")

    def test_rl_restored_action_does_not_reuse_corrected_action_thought(self):
        self.save_text(summary="manual summary")
        session = draft_store.load_session(self.session_id)
        session["row_edits"]["3"] = {
            "actions": '{"action":"open","text":"app"}',
            "original_actions": '{"action":"wait"}',
        }
        result = export_session_workbook(
            workbook_path=self.source, snapshot=self.snapshot, session=session,
            output_dir=self.root / "exports", export_id="rl",
        )
        self.assertEqual(result["sheets"], {"SFT_人工精修": 1, "RL_负向反思": 2})
        workbook = load_workbook(self.root / "exports" / result["filename"], read_only=True)
        try:
            sheet = workbook["RL_负向反思"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(sheet.cell(2, headers.index("action") + 1).value, '{"action":"wait"}')
            self.assertFalse(sheet.cell(2, headers.index("thought") + 1).value)
            self.assertEqual(sheet.cell(2, headers.index("summary") + 1).value, "manual summary")
        finally:
            workbook.close()
