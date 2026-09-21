"""Frozen, real collection batches supplied by a manually prepared task workbook."""
from __future__ import annotations

import hashlib
import io
import os
import uuid
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from .batch_lifecycle import is_batch_active
from .batch_operations import active_batch_lock
from .collection_runs import CollectionRunError, _component, _identifier
from .data_store import ArtifactStore, DATA_ROOT, RecordStore
from .data_store.paths import contained_path
from .data_store.registry import utc_now
from .task_generation.collection_batches import COLLECTION_COLUMNS, payload_digest


class ManualCollectionError(CollectionRunError):
    pass


def _cell(value):
    return value.isoformat() if isinstance(value, (date, datetime, time)) else value


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class ManualCollectionStore:
    def __init__(self, root: Path = DATA_ROOT):
        self.root = contained_path(Path(root))
        self.records = RecordStore(self.root)
        self.artifacts = ArtifactStore(self.root)

    def _parse(self, content: bytes, batch_id: str):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 512 * 1024 * 1024:
                    raise ManualCollectionError("Excel 解压后超过 512 MiB", 413)
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except ManualCollectionError:
            raise
        except Exception as exc:
            raise ManualCollectionError("无法读取采集任务 Excel", 422) from exc
        try:
            sheet = workbook.active
            if sheet not in workbook.worksheets:
                raise ManualCollectionError("活动工作表必须为采集任务明细表", 422)
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            while headers and headers[-1] is None:
                headers.pop()
            if headers != COLLECTION_COLUMNS:
                raise ManualCollectionError("采集任务表必须按模板保留全部 17 列表头及顺序", 422)
            tasks, warnings, seen = [], [], set()
            for row_number, row in enumerate(rows, 2):
                if not any(value is not None and value != "" for value in row):
                    continue
                if any(value is not None and value != "" for value in row[17:]):
                    raise ManualCollectionError(f"{sheet.title} 第 {row_number} 行有模板外的列", 422)
                values = dict(zip(COLLECTION_COLUMNS, [_cell(value) for value in row[:17]]))
                values = {column: values.get(column) for column in COLLECTION_COLUMNS}
                case, app, task = (_text(values[key]) for key in ("用例编号", "涉及APP", "任务"))
                if case.startswith("="):
                    raise ManualCollectionError(f"{sheet.title} 第 {row_number} 行用例编号不能为公式", 422)
                try:
                    _component(case, "用例编号")
                except CollectionRunError as exc:
                    raise ManualCollectionError(f"{sheet.title} 第 {row_number} 行：{exc}", 422) from exc
                if case.casefold() in seen:
                    raise ManualCollectionError(f"{sheet.title} 第 {row_number} 行用例编号重复：{case}", 422)
                seen.add(case.casefold())
                for field, value in (("涉及APP", app), ("任务", task)):
                    if not value or value.startswith("="):
                        raise ManualCollectionError(f"{sheet.title} 第 {row_number} 行 {field} 必须为非空文本", 422)
                missing = [field for field in ("一级场景", "二级场景") if not _text(values[field])]
                if missing:
                    warnings.append({"sheet": sheet.title, "row": row_number, "field": "、".join(missing),
                                     "message": "场景缺失，后续汇总显示为未分类"})
                tasks.append({"collection_case_id": case, "case_id": case,
                    "task_id": "manual_task_" + payload_digest([batch_id, case]),
                    "source_result_id": None, "source_row_id": f"{sheet.title}!{row_number}",
                    "source_row": row_number, "source_kind": "manual_collection", "task": task,
                    "app": app, "scene": _text(values["一级场景"]) or None,
                    "capability": _text(values["二级场景"]) or None,
                    "sub_capability": _text(values["三级场景"]) or None,
                    "pre_dependency": "unknown", "pre_task_id": None, "source_status": None,
                    "source_seed_id": None, "source_task": None, "dependency_error": None,
                    "source_values": values})
            if not tasks:
                raise ManualCollectionError("采集任务表没有有效任务", 422)
            return sheet.title, tasks, warnings
        finally:
            workbook.close()

    def register(self, filename: str, content: bytes, description: str = "",
                 request_id: str | None = None) -> dict:
        _component(filename, "文件名")
        if Path(filename).suffix.lower() != ".xlsx":
            raise ManualCollectionError("采集任务表只支持 .xlsx", 422)
        if not content or len(content) > 50 * 1024 * 1024:
            raise ManualCollectionError("采集任务表不能为空且不得超过 50 MiB", 413)
        if request_id is not None:
            _identifier(request_id, "request_id")
        batch_id = "manual_" + (payload_digest(request_id) if request_id else uuid.uuid4().hex)
        checksum = hashlib.sha256(content).hexdigest()
        request_hash = payload_digest({"filename": filename, "sha256": checksum, "description": description})
        with active_batch_lock(batch_id, self.root):
            existing = self.records.get("manual_collection_batches", batch_id)
            if existing:
                if existing["request_hash"] != request_hash:
                    raise ManualCollectionError("相同 request_id 对应不同的上传内容", 409)
                return self.get(batch_id, require_workbook=True)
            sheet, tasks, warnings = self._parse(content, batch_id)
            internal_name = f"collection-batch-{batch_id}.xlsx"
            created = utc_now()
            snapshot = {"schema_version": 1, "job_id": None, "source_job_id": None,
                "batch_id": batch_id, "kind": "manual_collection", "job_status": None,
                "knowledge_base_version": None, "task_count": len(tasks), "tasks": tasks,
                "errors": [], "warnings": warnings}
            payload = {"schema_version": 1, "batch_id": batch_id, "source_job_id": None,
                "kind": "manual_collection", "job_status": None, "knowledge_base_version": None,
                "created_at": created, "task_count": len(tasks), "apps": list(dict.fromkeys(task["app"] for task in tasks)),
                "filename": internal_name, "original_filename": filename, "description": description,
                "download_url": f"/api/task-generation/collection-batches/{batch_id}/workbook",
                "sheet_name": sheet, "snapshot": snapshot, "warnings": warnings,
                "workbook_sha256": checksum}
            temporary = contained_path(self.root, "tmp", "manual_collection", uuid.uuid4().hex + ".xlsx")
            temporary.parent.mkdir(parents=True, exist_ok=True)
            try:
                with temporary.open("xb") as stream:
                    stream.write(content)
                    stream.flush()
                    os.fsync(stream.fileno())
                self.artifacts.publish_many(batch_id, [{"stage": "00_collection", "payload": payload,
                    "workbooks": {internal_name: temporary}, "metadata": {"kind": "manual_collection", "task_count": len(tasks)},
                    "source_refs": [{"kind": "manual_collection", "filename": filename, "sha256": checksum}]}],
                    record_entries=[{"namespace": "manual_collection_batches", "key": batch_id,
                        "payload": {"batch_id": batch_id, "request_hash": request_hash, "created_at": created,
                                    "payload_sha256": payload_digest(payload)}, "expected_revision": 0}])
            finally:
                temporary.unlink(missing_ok=True)
            return self.get(batch_id, require_workbook=True)

    def get(self, batch_id: str, *, require_workbook: bool = False) -> dict | None:
        _identifier(batch_id, "采集批次编号")
        record = self.records.get("manual_collection_batches", batch_id)
        if record is None:
            return None
        artifact = self.artifacts.get(batch_id, "00_collection")
        if artifact is None:
            raise ManualCollectionError("手工采集批次快照缺失")
        payload = self.artifacts.read_payload(artifact)
        if (payload.get("batch_id") != batch_id or payload.get("kind") != "manual_collection"
                or payload_digest(payload) != record["payload_sha256"]):
            raise ManualCollectionError("手工采集批次快照校验失败")
        workbook_path = contained_path(self.root, "batches", batch_id, "00_collection", payload["filename"])
        if require_workbook:
            workbook_path = self.artifacts.resolve_file(artifact, payload["filename"])
        return {**payload, "workbook_path": str(workbook_path)}

    def list(self, *, include_published: bool = False) -> list[dict]:
        return [self.get(item["batch_id"]) for item in reversed(self.records.list("manual_collection_batches"))
                if include_published or is_batch_active(item["batch_id"], self.root)]
