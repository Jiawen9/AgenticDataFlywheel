"""A durable single-device main -> evaluation worker, launched in its own process group."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from .core import CollectorError, atomic_json, boundary, component, now


def run(config):
    phone, workdir, output = config["phone_id"], Path(config["workdir"]), Path(config["output"])
    result = {"trajectories": [], "reports": [], "errors": []}
    if (output / "cancelled.json").exists() or (output.parent / "cancelled.json").exists():
        result["errors"].append({"phone_id": phone, "error": "运行已取消，未执行设备命令"})
        atomic_json(output / "completion.json", result)
        return
    runs = workdir / ".runs"
    # Preserve failed/unarchived output; a new run never consumes it as its own result.
    if runs.exists():
        previous = output / "previous-device-output"
        if previous.exists():
            raise CollectorError("当前设备仍有未完成的工作目录，请先检查运行记录")
        shutil.move(str(boundary(workdir, runs)), str(previous))
    runs.mkdir()
    old_reports = list(workdir.glob("evaluator_results_*.xlsx"))
    if old_reports:
        previous_reports = output / "previous-reports"
        previous_reports.mkdir()
        for path in old_reports:
            shutil.move(str(boundary(workdir, path)), str(previous_reports / path.name))
    workbook = load_workbook(config["task_workbook"])
    try:
        sheet = workbook.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
        index = headers.index("用例编号") + 1
        for row in range(sheet.max_row, 1, -1):
            if str(sheet.cell(row, index).value or "").strip() not in config["cases"]:
                sheet.delete_rows(row)
        # Preserve all other column values, styles and sheets used by the execution template.
        workbook.save(workdir / "test.xlsx")
    finally:
        workbook.close()
    env = dict(os.environ)
    params = config["params"]
    env.update(ADB_SERIAL=phone, ADB_PATH=config["adb_path"],
               GUI_AGENT_CORE_VLA_SAMPLING_ENABLED=str(params["sampling_enabled"]).lower(),
               GUI_AGENT_CORE_VLA_TEMPERATURE=str(params["temperature"]),
               GUI_AGENT_CORE_VLA_TOP_P=str(params["top_p"]),
               GUI_AGENT_CORE_VLA_ENABLE_SKILLS=str(params["use_experience_lib"]).lower(),
               GUI_AGENT_VLA_API_BASE_URL=params["vla"])
    def call(args, timeout):
        return subprocess.run(args, cwd=workdir, env=env, check=True, timeout=timeout)
    try:
        check = subprocess.run([config["adb_path"], "-s", phone, "get-state"], capture_output=True,
                               text=True, timeout=config.get("adb_timeout", 10), env=env)
        if check.returncode or check.stdout.strip() != "device":
            raise CollectorError("ADB 预检失败：设备未连接")
        yadb = subprocess.run([config["adb_path"], "-s", phone, "shell", "test", "-x", "/data/local/tmp/yadb"],
                              capture_output=True, timeout=config.get("adb_timeout", 10), env=env)
        if yadb.returncode:
            if not (workdir / "yadb").is_file():
                raise CollectorError("设备缺少 yadb 且执行模板没有 yadb 文件")
            call([config["adb_path"], "-s", phone, "push", str(workdir / "yadb"), "/data/local/tmp/yadb"], 30)
            call([config["adb_path"], "-s", phone, "shell", "chmod", "755", "/data/local/tmp/yadb"], 10)
        shutil.copyfile(workdir / "config_trajectory.py", workdir / "config.py")
        try:
            call([sys.executable, "-c", "import main; main.main()"], config["execution_timeout"])
        except (subprocess.SubprocessError, OSError) as exc:
            result["errors"].append({"phone_id": phone, "error": f"轨迹执行失败：{exc}"})
        shutil.copyfile(workdir / "config_report.py", workdir / "config.py")
        try:
            call([sys.executable, str(workdir / "run_evaluator_batch.py")], config["evaluator_timeout"])
        except (subprocess.SubprocessError, OSError) as exc:
            result["errors"].append({"phone_id": phone, "error": f"评估失败：{exc}"})
        # Freeze a copy before returning any successful entries; no "latest directory" guessing.
        snapshot = output / "raw-output"
        shutil.copytree(runs, snapshot)
        known = set(config["cases"])
        for evaluation in sorted(snapshot.rglob("_trajectory_for_evaluate.json")):
            trajectory = evaluation.parent
            rel = trajectory.relative_to(snapshot)
            # Deployed template emits .runs/<timestamp>/<case>/<trajectory> (or <case>/<trajectory>).
            if len(rel.parts) not in {2, 3} or rel.parts[-2] not in known:
                result["errors"].append({"phone_id": phone, "error": f"无法关联任务目录：{rel.as_posix()}"})
                continue
            case, source_id = rel.parts[-2:]
            component(case); component(source_id)
            try:
                payload = json.loads(evaluation.read_text(encoding="utf-8-sig"))
                if not isinstance(payload, dict):
                    raise ValueError("trajectory JSON must be an object")
            except (ValueError, OSError) as exc:
                result["errors"].append({"phone_id": phone, "collection_case_id": case, "error": f"轨迹JSON损坏：{exc}"})
                continue
            result["trajectories"].append({"collection_case_id": case, "phone_id": phone,
                "source_trajectory_id": source_id, "source_dir": str(trajectory), "collected_at": now()})
        report_root = output / "report-output"
        report_root.mkdir(exist_ok=True)
        for path in sorted(workdir.glob("evaluator_results_*.xlsx")):
            dest = report_root / path.name
            shutil.copyfile(path, dest)
            result["reports"].append({"phone_id": phone, "path": str(dest)})
    except Exception as exc:
        result["errors"].append({"phone_id": phone, "error": f"执行预检/归档失败：{exc}"})
    finally:
        shutil.copyfile(workdir / "config_trajectory.py", workdir / "config.py")
    atomic_json(output / "completion.json", result)


if __name__ == "__main__":
    path = Path(sys.argv[1])
    config = json.loads(path.read_text(encoding="utf-8"))
    atomic_json(Path(config["output"]) / "process.json", {"pid": os.getpid(), "config_path": str(path)})
    try:
        run(config)
    except Exception as exc:
        atomic_json(Path(config["output"]) / "completion.json",
                    {"trajectories": [], "reports": [], "errors": [{"phone_id": config["phone_id"], "error": str(exc)}]})
        raise
