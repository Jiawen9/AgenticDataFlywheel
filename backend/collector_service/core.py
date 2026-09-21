"""Durable phone-factory runs and immutable, verifiable result archives."""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import math
import os
import re
import shutil
import sqlite3
import threading
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from pathlib import Path
from openpyxl import load_workbook

TERMINAL = {"succeeded", "partial", "failed", "interrupted"}
ACTIVE = {"queued", "running"}


class CollectorError(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def now():
    return datetime.now(timezone.utc).isoformat()


def component(value, label="编号"):
    if (not isinstance(value, str) or not value or len(value) > 180 or value in {".", ".."}
            or value.endswith((".", " ")) or re.search(r'[<>:"/\\|?*\x00-\x1f]', value)
            or re.fullmatch(r"(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(\..*)?", value, re.I)):
        raise CollectorError(f"{label} 无效")
    return value


def phone_identifier(value, label="手机编号"):
    if not isinstance(value, str) or not value or len(value) > 180 or value.startswith("-") or re.search(r"[\s/\\\\\x00-\x1f]", value):
        raise CollectorError(f"{label} 无效")
    return value


def cell_text(value):
    """Match manual-batch identity normalization without rewriting uploaded bytes."""
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat().strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def identifier(value, label="运行编号"):
    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,127}", value):
        raise CollectorError(f"{label} 无效")
    return value


def token(value):
    return hashlib.sha256(value.encode()).hexdigest()[:20]


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":"), allow_nan=False).encode()).hexdigest()


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def boundary(base, path):
    base, path = Path(base).resolve(), Path(path)
    # Check links before resolve; resolving first would hide a junction.
    for item in (path, *path.parents):
        if item == base:
            break
        if item.is_symlink() or (hasattr(item, "is_junction") and item.is_junction()):
            raise CollectorError("路径不允许符号链接或目录联接")
    resolved = path.resolve()
    if resolved == base or base not in resolved.parents:
        raise CollectorError("路径超出允许目录")
    return resolved


def atomic_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(value, f, ensure_ascii=False, allow_nan=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)


@dataclass(frozen=True)
class CollectorConfig:
    root: Path
    template_dir: Path | None = None
    adb_path: str = "adb"
    execution_timeout: float = 86400
    evaluator_timeout: float = 3600
    adb_timeout: float = 10
    max_upload_bytes: int = 50 * 1024 * 1024

    @classmethod
    def from_env(cls):
        root = Path(os.environ.get("PHONE_FACTORY_ROOT", "./phone_factory_workspace")).absolute()
        template = os.environ.get("PHONE_FACTORY_TEMPLATE_DIR")
        return cls(root=root, template_dir=Path(template).absolute() if template else None,
                   adb_path=os.environ.get("PHONE_FACTORY_ADB", "adb"),
                   execution_timeout=float(os.environ.get("PHONE_FACTORY_EXECUTION_TIMEOUT", "86400")),
                   evaluator_timeout=float(os.environ.get("PHONE_FACTORY_EVALUATOR_TIMEOUT", "3600")),
                   adb_timeout=float(os.environ.get("PHONE_FACTORY_ADB_TIMEOUT", "10")))


class Collector:
    def __init__(self, config, *, device_adapter, execution_adapter, executor=None):
        self.config = config
        self.root = Path(config.root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.devices, self.execution = device_adapter, execution_adapter
        self.executor = executor or ThreadPoolExecutor(max_workers=16, thread_name_prefix="collector")
        self._owns_executor = executor is None
        self._threads = {}
        self._cancel = {}
        self._mutex = threading.RLock()
        self._resource_locks = {}
        with self.db() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS runs (id TEXT PRIMARY KEY, payload TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS phones (id TEXT PRIMARY KEY, payload TEXT NOT NULL);
            """)
        # Persisted active runs reattach through the executor's durable completion marker.
        for run in self.list_runs():
            if run["status"] in ACTIVE:
                self._schedule(run["run_id"])

    @contextlib.contextmanager
    def db(self):
        conn = sqlite3.connect(self.root / "collector.sqlite", timeout=30)
        try:
            conn.execute("PRAGMA busy_timeout=30000")
            conn.execute("PRAGMA journal_mode=WAL")
            yield conn
            conn.commit()
        except BaseException:
            conn.rollback()
            raise
        finally:
            conn.close()

    @contextlib.contextmanager
    def lock(self, name):
        # Thread + OS locks serialize the same root even with multiple service instances.
        with self._mutex:
            resource_lock = self._resource_locks.setdefault(name, threading.RLock())
        with resource_lock:
            folder = self.root / ".locks"
            folder.mkdir(exist_ok=True)
            with (folder / (token(name) + ".lock")).open("a+b") as f:
                f.seek(0)
                if f.read(1) == b"":
                    f.write(b"0")
                    f.flush()
                f.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(f.fileno(), msvcrt.LK_LOCK, 1)
                else:
                    import fcntl
                    fcntl.flock(f.fileno(), fcntl.LOCK_EX)
                try:
                    yield
                finally:
                    f.seek(0)
                    if os.name == "nt":
                        msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
                    else:
                        fcntl.flock(f.fileno(), fcntl.LOCK_UN)

    def _get(self, conn, run_id):
        row = conn.execute("SELECT payload FROM runs WHERE id=?", (run_id,)).fetchone()
        return json.loads(row[0]) if row else None

    def _save(self, conn, run):
        conn.execute("INSERT OR REPLACE INTO runs VALUES (?,?)",
                     (run["run_id"], json.dumps(run, ensure_ascii=False, allow_nan=False)))

    def get(self, run_id):
        identifier(run_id)
        with self.db() as conn:
            run = self._get(conn, run_id)
        if run is None:
            raise CollectorError("采集运行不存在", 404)
        return run

    def list_runs(self):
        with self.db() as conn:
            return [json.loads(row[0]) for row in conn.execute("SELECT payload FROM runs ORDER BY rowid")]

    def run_dir(self, run_id):
        return boundary(self.root, self.root / "runs" / identifier(run_id))

    def phone_dir(self, phone_id):
        return boundary(self.root, self.root / "devices" / token(phone_identifier(phone_id)))

    def add_phone(self, phone_id):
        phone_identifier(phone_id)
        with self.lock("mutations"):
            path = self.phone_dir(phone_id)
            if not path.exists():
                self.execution.prepare_device(path)
            with self.db() as conn:
                conn.execute("INSERT OR REPLACE INTO phones VALUES (?,?)",
                             (phone_id, json.dumps({"phone_id": phone_id, "created_at": now()})))
        return {"ok": True, "message": "手机已添加", "phone_id": phone_id}

    def delete_phone(self, phone_id):
        phone_identifier(phone_id)
        with self.lock("mutations"):
            owned = set(self.execution.owned_runs(phone_id)) if hasattr(self.execution, "owned_runs") else set()
            known = {r["run_id"]: r for r in self.list_runs()}
            affected = [r for r in known.values() if (r["status"] in ACTIVE and phone_id in r["phone_tasks"]) or r["run_id"] in owned]
            affected.extend({"run_id": run_id} for run_id in owned if run_id not in known)
            for run in affected:
                # Stop the whole assigned run (and its owned process groups) before touching directories.
                self.execution.cancel(run, self.run_dir(run["run_id"]))
                self._cancel.setdefault(run["run_id"], threading.Event()).set()
            for run in affected:
                with self.db() as conn:
                    latest = self._get(conn, run["run_id"])
                    if latest and (latest["status"] in ACTIVE or (latest["status"] == "failed" and run["run_id"] in owned)):
                        latest.update(status="interrupted", completed_at=now(),
                                      errors=[*latest.get("errors", []), {"phone_id": phone_id, "error": "手机已删除，运行已中断"}])
                        self._save(conn, latest)
            path = self.phone_dir(phone_id)
            if path.exists():
                shutil.rmtree(boundary(self.root / "devices", path))
            with self.db() as conn:
                conn.execute("DELETE FROM phones WHERE id=?", (phone_id,))
        return {"ok": True, "message": "手机已删除", "phone_id": phone_id,
                "processes_killed": len(affected), "folder_removed": True}

    def statuses(self, phones):
        active = {p for r in self.list_runs() if r["status"] in ACTIVE for p in r["phone_tasks"]}
        result = []
        for phone in phones:
            phone_identifier(phone)
            running = phone in active
            if not running and hasattr(self.execution, "owned_runs"):
                running = bool(self.execution.owned_runs(phone))
            result.append({"phone_id": phone, "status": "运行中" if running else "空闲"})
        return result

    def submit(self, *, task_bytes, apps_bytes, form):
        if max(len(task_bytes), len(apps_bytes)) > self.config.max_upload_bytes:
            raise CollectorError("上传文件超过 50 MiB", 413)
        try:
            apps = json.loads(apps_bytes.decode("utf-8-sig"))
            manifest = json.loads(form.get("task_manifest") or "{}")
        except (ValueError, UnicodeError) as exc:
            raise CollectorError("任务清单或手机关联文件不是有效 JSON") from exc
        if not isinstance(apps, list) or not isinstance(manifest, dict):
            raise CollectorError("手机关联文件必须为数组，task_manifest 必须为对象")
        mode = form.get("runmode") or "generate"
        if mode not in {"generate", "modeliter"}:
            raise CollectorError("runmode 必须为 generate 或 modeliter")
        legacy = not form.get("collection_run_id")
        run_id = identifier(form.get("collection_run_id") or "legacy_" + uuid.uuid4().hex)
        batch_id = form.get("batch_id") or None
        if batch_id is not None:
            identifier(batch_id, "批次编号")
        if not legacy and mode == "generate" and not batch_id:
            raise CollectorError("生产采集缺少 batch_id")
        vla = str(form.get("vla") or "").strip()
        if not vla:
            raise CollectorError("缺少 VLA 接口参数 vla")
        def boolean(key):
            value = str(form.get(key) or "false").lower()
            if value not in {"true", "false", "1", "0", "yes", "no"}:
                raise CollectorError(f"{key} 必须为布尔值")
            return value in {"true", "1", "yes"}
        try:
            temperature, top_p = float(form.get("temperature") or .7), float(form.get("top_p") or .85)
        except ValueError as exc:
            raise CollectorError("采样参数无效") from exc
        if not math.isfinite(temperature) or not 0 <= temperature <= 2 or not math.isfinite(top_p) or not 0 <= top_p <= 1:
            raise CollectorError("temperature 应为 0–2，top_p 应为 0–1")
        params = {"vla": vla, "runmode": mode, "phone_id": str(form.get("phone_id") or ""),
                  "app": str(form.get("app") or ""), "sampling_enabled": boolean("sampling_enabled"),
                  "temperature": temperature, "top_p": top_p, "use_experience_lib": boolean("use_experience_lib")}
        try:
            with zipfile.ZipFile(io.BytesIO(task_bytes)) as zipped:
                if sum(entry.file_size for entry in zipped.infolist()) > 512 * 1024 * 1024:
                    raise CollectorError("Excel 解压后超过允许大小")
            wb = load_workbook(io.BytesIO(task_bytes), read_only=True, data_only=False)
            try:
                ws = wb.active
                ws.reset_dimensions()
                rows = list(ws.iter_rows(values_only=True))
            finally:
                wb.close()
        except Exception as exc:
            raise CollectorError("任务文件不是有效 Excel 工作表") from exc
        if not rows:
            raise CollectorError("任务文件为空")
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]
        if len(set(headers)) != len(headers) or not {"用例编号", "涉及APP", "任务"}.issubset(headers):
            raise CollectorError("任务文件缺少用例编号/涉及APP/任务列或包含重复表头")
        table = {}
        for number, cells in enumerate(rows[1:], start=2):
            if not any(c is not None for c in cells):
                continue
            row = dict(zip(headers, cells))
            case = component(cell_text(row.get("用例编号")), f"第 {number} 行用例编号")
            if case.casefold() in {c.casefold() for c in table}:
                raise CollectorError(f"第 {number} 行用例编号重复")
            if any(not cell_text(row.get(k)) or cell_text(row.get(k)).startswith("=") for k in ("任务", "涉及APP")):
                raise CollectorError(f"第 {number} 行任务或涉及APP 为空")
            table[case] = row
        tasks = manifest.get("tasks")
        if legacy and not tasks:
            tasks = [{"collection_case_id": c, "task_id": c, "task": r["任务"], "app": r["涉及APP"]} for c, r in table.items()]
        if not isinstance(tasks, list) or not tasks:
            raise CollectorError("task_manifest.tasks 为空")
        by_case = {}
        for task in tasks:
            if not isinstance(task, dict):
                raise CollectorError("任务清单项必须为对象")
            case = component(task.get("collection_case_id"), "任务清单用例编号")
            if case in by_case or case not in table or not isinstance(task.get("task_id"), str) or not task["task_id"].strip():
                raise CollectorError("任务清单包含未知/重复用例或缺少 task_id")
            if cell_text(task.get("task")) != cell_text(table[case]["任务"]) or cell_text(task.get("app")) != cell_text(table[case]["涉及APP"]):
                raise CollectorError(f"用例 {case} 的任务文本/App 与表格不匹配")
            by_case[case] = task
        if set(by_case) != set(table):
            raise CollectorError("任务清单与表格用例集合不匹配")
        phone_apps = {}
        for item in apps:
            if not isinstance(item, dict):
                raise CollectorError("手机关联项必须为对象")
            pid, app = phone_identifier(item.get("phone_id")), str(item.get("app") or "")
            if params["phone_id"] and params["phone_id"] != pid:
                continue
            if params["app"] and params["app"] != app:
                continue
            if app:
                phone_apps.setdefault(pid, set()).add(app)
        phone_tasks = {pid: [c for c, task in by_case.items() if cell_text(task["app"]) in allowed]
                       for pid, allowed in phone_apps.items()}
        phone_tasks = {p: cases for p, cases in phone_tasks.items() if cases}
        if not phone_tasks:
            raise CollectorError("没有匹配的手机/APP 任务")
        request_hash = digest({"task_sha256": hashlib.sha256(task_bytes).hexdigest(),
                               "apps": apps, "tasks": tasks, "params": params, "batch_id": batch_id})
        with self.lock("mutations"):
            with self.db() as conn:
                existing = self._get(conn, run_id)
            if existing:
                if existing["request_digest"] != request_hash:
                    raise CollectorError("相同运行编号对应不同的实际运行参数", 409)
                return self.public(existing)
            self.execution.preflight(params)
            busy = {p for r in self.list_runs() if r["status"] in ACTIVE for p in r["phone_tasks"]}
            if busy.intersection(phone_tasks):
                raise CollectorError("手机已有运行中的任务：" + "、".join(sorted(busy.intersection(phone_tasks))), 409)
            for phone in phone_tasks:
                if hasattr(self.execution, "owned_runs") and self.execution.owned_runs(phone):
                    raise CollectorError(f"手机 {phone} 仍有所属执行进程，请先结束该运行", 409)
                path = self.phone_dir(phone)
                if not path.exists():
                    self.execution.prepare_device(path)
            root = self.run_dir(run_id)
            root.mkdir(parents=True, exist_ok=True)
            for name, content in (("tasks.xlsx", task_bytes), ("apps.json", apps_bytes)):
                with (root / name).open("wb") as f:
                    f.write(content); f.flush(); os.fsync(f.fileno())
            run = {"run_id": run_id, "collection_run_id": run_id, "batch_id": batch_id, "run_mode": mode,
                   "status": "queued", "created_at": now(), "completed_at": None, "params": params,
                   "request_digest": request_hash, "tasks": by_case, "phone_tasks": phone_tasks,
                   "errors": [], "manifest": None, "legacy": legacy}
            atomic_json(root / "request.json", run)
            with self.db() as conn:
                self._save(conn, run)
        self._schedule(run_id)
        return self.public(self.get(run_id))

    def _schedule(self, run_id):
        if run_id in self._threads and not self._threads[run_id].done():
            return
        self._cancel.setdefault(run_id, threading.Event())
        self._threads[run_id] = self.executor.submit(self._execute, run_id)

    def _execute(self, run_id):
        # Lock a run, but allow deletion/cancellation and other devices to continue.
        with self.lock("run:" + run_id):
            run = self.get(run_id)
            if run["status"] not in ACTIVE:
                return
            try:
                root = self.run_dir(run_id)
                frozen = root / "frozen" / "manifest.json"
                if frozen.is_file():
                    manifest = json.loads(frozen.read_text(encoding="utf-8"))
                    self._verify_archive(root / "frozen", manifest)
                else:
                    with self.db() as conn:
                        current = self._get(conn, run_id)
                        if current["status"] not in ACTIVE:
                            return
                        current["status"] = "running"; self._save(conn, current)
                    output = root / "execution"
                    output.mkdir(exist_ok=True)
                    result = self.execution.execute(run, {p: self.phone_dir(p) for p in run["phone_tasks"]},
                                                    output, self._cancel[run_id])
                    if self._cancel[run_id].is_set() or self.get(run_id)["status"] not in ACTIVE:
                        return
                    manifest = self._freeze(run, result)
                with self.db() as conn:
                    latest = self._get(conn, run_id)
                    if latest["status"] not in ACTIVE:
                        return
                    has_results = bool(manifest["trajectories"] or (run["run_mode"] == "modeliter" and manifest["reports"]))
                    latest.update(manifest=manifest, errors=manifest["errors"], completed_at=now(),
                                  status=("partial" if manifest["errors"] else "succeeded") if has_results else "failed")
                    self._save(conn, latest)
            except Exception as exc:
                with self.db() as conn:
                    latest = self._get(conn, run_id)
                    if latest and latest["status"] in ACTIVE:
                        latest.update(status="failed", completed_at=now(),
                                      errors=[*latest.get("errors", []), {"error": str(exc)}])
                        self._save(conn, latest)

    def _freeze(self, run, result):
        root = self.run_dir(run["run_id"])
        execution = root / "execution"
        stage, final = root / ".freezing", root / "frozen"
        if stage.exists():
            shutil.rmtree(boundary(root, stage))
        stage.mkdir()
        raw = stage / "raw"
        raw.mkdir()
        trajectories, reports, errors, seen = [], [], list(result.get("errors", [])), set()
        for entry in result.get("trajectories", []):
            try:
                case = component(entry.get("collection_case_id"), "结果用例编号")
                phone = phone_identifier(entry.get("phone_id"))
                if case not in run["tasks"] or case not in run["phone_tasks"].get(phone, []):
                    raise CollectorError("结果用例/设备不属于本次运行")
                source_id = component(entry.get("source_trajectory_id"), "原轨迹编号")
                rel = f"{case}/{token(phone)}__{source_id}"
                if rel.casefold() in seen:
                    raise CollectorError("同一设备存在重复轨迹")
                source = boundary(execution, Path(entry["source_dir"]))
                if not source.is_dir():
                    raise CollectorError("轨迹目录不存在")
                content = json.loads((source / "_trajectory_for_evaluate.json").read_text(encoding="utf-8-sig"))
                if not isinstance(content, dict) or not isinstance(content.get("actions_flat"), list) or not content["actions_flat"]:
                    raise CollectorError("轨迹 evaluation JSON 缺少有效动作")
                from ..export_vla_trajectories import collect_rows
                rows, warnings = collect_rows(source)
                if not rows or any(not row[3] for row in rows) or any("skipped because" in warning for warning in warnings):
                    raise CollectorError("轨迹缺少完整可转换的步骤文件")
                files = []
                for src in sorted(source.rglob("*")):
                    boundary(source, src)
                    if not src.is_file():
                        continue
                    local = src.relative_to(source).as_posix()
                    for part in local.split("/"):
                        component(part, "文件名")
                    dst = raw / rel / local
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(src, dst)
                    files.append({"path": f"{rel}/{local}", "size": dst.stat().st_size, "sha256": sha(dst)})
                if not files or not any(x["path"].endswith("/_trajectory_for_evaluate.json") for x in files):
                    raise CollectorError("轨迹缺少 _trajectory_for_evaluate.json 或没有文件")
                seen.add(rel.casefold())
                trajectories.append({"collection_case_id": case, "relative_dir": rel, "phone_id": phone,
                                     "source_trajectory_id": source_id, "collected_at": entry.get("collected_at") or now(),
                                     "files": files})
            except (ValueError, OSError, KeyError) as exc:
                errors.append({"collection_case_id": entry.get("collection_case_id") if entry.get("collection_case_id") in run["tasks"] else None,
                               "phone_id": entry.get("phone_id"), "error": f"{entry.get('collection_case_id')}: {exc}"})
        for entry in result.get("reports", []):
            try:
                src = boundary(execution, Path(entry["path"]))
                component(src.name, "报告文件名")
                if not src.is_file():
                    raise CollectorError("报告文件不存在")
                rid = token(src.relative_to(execution).as_posix())
                rel = f"reports/{rid}/{src.name}"
                dst = raw / rel
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(src, dst)
                reports.append({"id": rid, "name": src.name, "path": rel,
                                "size": dst.stat().st_size, "sha256": sha(dst)})
            except (ValueError, OSError, KeyError) as exc:
                errors.append({"error": f"报告归档失败：{exc}"})
        for phone, cases in run["phone_tasks"].items():
            for case in cases:
                if not any(t["phone_id"] == phone and t["collection_case_id"] == case for t in trajectories):
                    errors.append({"collection_case_id": case, "phone_id": phone, "error": "没有可验证的完成轨迹"})
        # Only declared files are archived. Failed copies cannot slip into an otherwise partial run.
        files = [f["path"] for t in trajectories for f in t["files"]] + [r["path"] for r in reports]
        archive = stage / "archive.zip"
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as z:
            for name in files:
                z.write(raw / name, name)
        with archive.open("r+b") as f:
            os.fsync(f.fileno())
        manifest = {"schema_version": 1, "batch_id": run["batch_id"], "collection_run_id": run["run_id"],
                    "run_mode": run["run_mode"], "trajectories": trajectories, "reports": reports, "errors": errors,
                    "archive": {"size": archive.stat().st_size, "sha256": sha(archive)}}
        atomic_json(stage / "manifest.json", manifest)
        os.replace(stage, final)
        return manifest

    def _verify_archive(self, root, manifest):
        archive = root / "archive.zip"
        if not archive.is_file() or archive.stat().st_size != manifest["archive"]["size"] or sha(archive) != manifest["archive"]["sha256"]:
            raise CollectorError("冻结归档校验失败", 409)
        return archive

    def archive(self, run_id):
        run = self.get(run_id)
        if run["status"] not in TERMINAL or not run.get("manifest"):
            raise CollectorError("结果尚未完成归档", 409)
        return self._verify_archive(self.run_dir(run_id) / "frozen", run["manifest"])

    def reports(self, mode):
        if mode not in {"generate", "modeliter"}:
            raise CollectorError("runmode 必须为 generate 或 modeliter")
        folders = []
        for run in self.list_runs():
            if run["run_mode"] != mode or not run.get("manifest"):
                continue
            modified = int(datetime.fromisoformat(run["completed_at"]).timestamp())
            files = [{**report, "file_id": report["id"], "run_id": run["run_id"], "modified": modified}
                     for report in run["manifest"]["reports"]]
            folders.append({"index": len(folders) + 1, "dir_name": run["run_id"],
                            "run_id": run["run_id"], "modified": modified, "files": files})
        return {"ok": True, "folders": folders}

    def report(self, *, mode, run_id=None, file_id=None, folder=None, name=None):
        folders = self.reports(mode)["folders"]
        target = run_id or folder or (folders[-1]["run_id"] if folders else "")
        run = self.get(target)
        if run["run_mode"] != mode or not run.get("manifest"):
            raise CollectorError("报告不存在", 404)
        matches = [r for r in run["manifest"]["reports"] if (r["id"] == file_id if file_id else r["name"] == name)]
        if len(matches) != 1:
            raise CollectorError("报告不存在或名称不唯一，请使用 file_id", 404)
        report = matches[0]
        path = boundary(self.run_dir(target) / "frozen" / "raw",
                        self.run_dir(target) / "frozen" / "raw" / report["path"])
        if not path.is_file() or sha(path) != report["sha256"]:
            raise CollectorError("报告文件校验失败", 409)
        return path

    @staticmethod
    def public(run):
        return {"ok": True, **{k: run[k] for k in ("run_id", "collection_run_id", "batch_id", "run_mode",
                 "status", "created_at", "completed_at", "errors", "manifest")},
                "phones": [{"phone_id": p} for p in run["phone_tasks"]]}

    def close(self):
        if self._owns_executor:
            self.executor.shutdown(wait=False)
