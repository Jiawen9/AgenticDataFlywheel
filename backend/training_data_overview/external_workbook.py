"""Validate an external step workbook without consulting raw files or sessions."""
from __future__ import annotations

from collections import Counter
from datetime import date, datetime, time
import hashlib
import json
import math
from pathlib import Path, PurePosixPath
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .converter import _classification_known, _field, _image, _json_action, _text

PARSER_VERSION = 1
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
METADATA_FIELDS = ("data_source", "data_date", "app", "level1", "level2", "sheet_name")
IDENTITY_FIELDS = ("trajectory_id", "文件夹名", "meta_task", "轨迹")
CLASSIFICATION_FIELDS = (("APP", "app", "涉及APP"), ("一级场景", "scene"), ("二级场景", "capability"))


def _cell(value: Any) -> Any:
    """Keep every cell in the content fingerprint, including SOP and extra columns."""
    if isinstance(value, (datetime, date, time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, float) and not math.isfinite(value):
        return {"type": "float", "value": str(value)}
    if value is None or isinstance(value, (bool, int, float, str)):
        return value.replace("\r\n", "\n").replace("\r", "\n") if isinstance(value, str) else value
    return {"type": type(value).__name__, "value": str(value)}


def _issue(sheet: str, row: int | None, field: str, message: str) -> dict:
    return {"sheet": sheet, "row": row, "field": field, "message": message}


def _summary(rows: list[dict]) -> dict:
    apps: dict[str, dict] = {}
    scenes: dict[tuple[str, str], dict] = {}
    for row in rows:
        app = apps.setdefault(row["APP"], {"name": row["APP"], "trajectory_count": 0, "step_count": 0})
        scene = scenes.setdefault((row["一级场景"], row["二级场景"]), {
            "level1": row["一级场景"], "level2": row["二级场景"], "trajectory_count": 0, "step_count": 0})
        for value in (app, scene):
            value["trajectory_count"] += 1
            value["step_count"] += row["step数量"]
    return {"trajectory_count": len(rows), "step_count": sum(row["step数量"] for row in rows),
            "apps": list(apps.values()), "scenes": list(scenes.values())}


def parse_external_workbook(path: Path | str, *, data_source: str, data_date: str,
                            app: str = "", level1: str = "", level2: str = "",
                            sheet_name: str = "") -> dict:
    """Return preview diagnostics and canonical aggregate rows for one selected sheet.

    Blank classifications inherit a nonblank value from the same trajectory before
    applying upload defaults. Images are identity text only and are never opened.
    """
    path = Path(path)
    metadata = {"data_source": _text(data_source) or "人工采集", "data_date": _text(data_date),
                "app": _text(app), "level1": _text(level1), "level2": _text(level2),
                "sheet_name": _text(sheet_name)}
    result = {"valid": False, "sheets": [], "sheet_name": metadata["sheet_name"],
              "metadata": metadata, "summary": _summary([]), "errors": [], "warnings": [],
              "rows": [], "content_hash": ""}
    errors, warnings = result["errors"], result["warnings"]
    try:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", metadata["data_date"]):
            raise ValueError("date format")
        date.fromisoformat(metadata["data_date"])
    except ValueError:
        errors.append(_issue(metadata["sheet_name"], None, "data_date", "数据日期必须是有效的 YYYY-MM-DD 日期"))
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        errors.append(_issue("", None, "file", "仅支持 .xlsx 或 .xlsm 步骤表"))
        return result
    workbook = None
    try:
        # XLSX/XLSM are ZIP archives: the upload byte limit alone does not bound
        # expansion. Inspect all entries before openpyxl reads XML or shared strings.
        with ZipFile(path) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                errors.append(_issue(metadata["sheet_name"], None, "file",
                                     "工作簿解压后超过 512 MiB 上限，请拆分后上传"))
                return result
            if any(entry.flag_bits & 1 for entry in archive.infolist()):
                errors.append(_issue(metadata["sheet_name"], None, "file", "不支持加密工作簿，请先另存为未加密的步骤表"))
                return result
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        result["sheets"] = list(workbook.sheetnames)
        selected = metadata["sheet_name"] or (workbook.active.title if workbook.active is not None else "")
        if selected not in workbook.sheetnames:
            errors.append(_issue(selected, None, "sheet_name", "指定的工作表不存在"))
            return result
        result["sheet_name"] = metadata["sheet_name"] = selected
        sheet = workbook[selected]
        if not callable(getattr(sheet, "iter_rows", None)):
            errors.append(_issue(selected, None, "sheet_name", "所选标签不是步骤工作表，请选择包含轨迹和动作列的工作表"))
            return result
        # Ignore declared dimensions, which can be stale or claim billions of
        # empty cells even when the actual XML contains only a few rows.
        sheet.reset_dimensions()
        values = sheet.iter_rows(values_only=True)
        raw_headers = list(next(values, ()))
        # Formatting-only trailing columns do not change the business table.
        while raw_headers and raw_headers[-1] is None:
            raw_headers.pop()
        headers = [_text(value) for value in raw_headers]
        counts = Counter(header for header in headers if header)
        for header, count in counts.items():
            if count > 1:
                errors.append(_issue(selected, 1, header, "表头重复，请为每列保留唯一名称"))
        if not any(name in headers for name in IDENTITY_FIELDS):
            errors.append(_issue(selected, 1, "trajectory_id", "缺少轨迹编号列（trajectory_id、文件夹名、meta_task 或轨迹）"))
        if not any(name in headers for name in ("action", "actions")):
            errors.append(_issue(selected, 1, "action", "缺少最终动作列（action 或 actions）"))
        if any(issue["row"] == 1 for issue in errors):
            return result
        groups: dict[str, dict] = {}
        # Stream the same canonical JSON used by parser v1; do not retain every
        # step's SOP, image path and auxiliary cells in a second in-memory table.
        fingerprint = {"parser_version": PARSER_VERSION,
                       "metadata": {key: value for key, value in metadata.items() if key != "sheet_name"},
                       "headers": headers}
        digest = hashlib.sha256()
        digest.update(json.dumps(fingerprint, ensure_ascii=False, sort_keys=True,
                                 separators=(",", ":"), allow_nan=False).encode("utf-8")[:-1])
        digest.update(b',"steps":[')
        has_step = False
        for row_number, cells in enumerate(values, start=2):
            if not any(cell is not None and cell != "" for cell in cells):
                continue
            cells = list(cells)
            for index, value in enumerate(cells):
                if value is not None and value != "" and (index >= len(headers) or not headers[index]):
                    errors.append(_issue(selected, row_number, f"column_{index + 1}", "有内容的列缺少表头"))
            # Preserve source step order and all cell contents, including unparsed columns.
            normalized = [_cell(value) for value in cells[:len(headers)]]
            normalized.extend([None] * (len(headers) - len(normalized)))
            if has_step:
                digest.update(b",")
            digest.update(json.dumps(normalized, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":"), allow_nan=False).encode("utf-8"))
            has_step = True
            row = {name: value for name, value in zip(headers, cells) if name}
            identity = _field(row, ("trajectory_id",), "")
            if not identity:
                original = _field(row, ("文件夹名", "meta_task", "轨迹"), "")
                if not original:
                    errors.append(_issue(selected, row_number, "trajectory_id", "该步骤缺少轨迹编号"))
                    continue
                image = _image(row.get("image"))
                identity = f"{PurePosixPath(image).parent.as_posix()}/{original}" if image else original
            group = groups.setdefault(identity, {"steps": 0, "actions": Counter(),
                "fields": [{}, {}, {}], "first_row": row_number})
            group["steps"] += 1
            for target, aliases in zip(group["fields"], CLASSIFICATION_FIELDS):
                value = _field(row, aliases, "")
                if value:
                    if target and value not in target:
                        first_value, first_row = next(iter(target.items()))
                        errors.append(_issue(selected, row_number, aliases[0],
                            f"同一轨迹的分类冲突：第 {first_row} 行为“{first_value}”，本行为“{value}”"))
                    target.setdefault(value, row_number)
            action = _json_action(row.get("actions", row.get("action")))
            name = _text(action.get("action")) if action else ""
            if not name:
                name = "unknown"
                warnings.append(_issue(selected, row_number, "actions" if "actions" in row else "action",
                                       "最终动作无法解析，已计入 unknown"))
            group["actions"][name] += 1
        rows = []
        for identity, group in groups.items():
            labels = []
            for values, key, aliases in zip(group["fields"], ("app", "level1", "level2"), CLASSIFICATION_FIELDS):
                fallback = "未记录 App" if key == "app" else "未分类"
                label = next(iter(values), "") or metadata[key] or fallback
                labels.append(label)
                if not values and not metadata[key]:
                    warnings.append(_issue(selected, group["first_row"], aliases[0], f"该轨迹未提供分类，已归入{fallback}"))
            app_value, first, second = labels
            rows.append({"trajectory_identity": identity, "step数量": group["steps"],
                "APP": app_value, "一级场景": first, "二级场景": second,
                "时间": metadata["data_date"], "生产方式": metadata["data_source"],
                "人工精修步骤数量": None, "manual_known": False,
                "action_box": dict(sorted(group["actions"].items())),
                "app_known": _classification_known(app_value), "level1_known": _classification_known(first),
                "level2_known": _classification_known(second)})
        if not rows:
            errors.append(_issue(selected, None, "file", "工作表没有可统计的轨迹步骤"))
        result["summary"] = _summary(rows)
        if not errors:
            warnings.append(_issue(selected, None, "人工精修步骤数量", "外部步骤表未提供可核验的精修记录，人工精修数量记为未知"))
            result["rows"] = rows
            digest.update(b"]}")
            result["content_hash"] = digest.hexdigest()
            result["valid"] = True
    except (OSError, ValueError, KeyError, TypeError, SyntaxError, BadZipFile, InvalidFileException) as exc:
        errors.append(_issue(metadata["sheet_name"], None, "file", f"无法读取步骤表：{exc}"))
    finally:
        if workbook is not None:
            workbook.close()
    return result
