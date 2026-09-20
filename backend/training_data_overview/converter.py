"""Read frozen published workbooks and derive the DataVue all_data contract.

No current correction/session values, raw directories, or latest stage versions
are consulted. Conversion never writes application data.
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook

from ..data_store import ArtifactStore
from ..data_store.release_provenance import FrozenReleaseSources

ALL_DATA_COLUMNS = [
    "轨迹", "子任务轨迹", "step数量", "APP", "一级场景", "二级场景",
    "时间", "生产方式", "人工精修步骤数量", "action_box",
]


class ConversionError(ValueError):
    """A published input or an explicitly registered source is invalid."""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _image(value: Any) -> str:
    raw = _text(value).replace("\\", "/")
    return PurePosixPath(raw).as_posix() if raw else ""


def _sha(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def _warn(warnings: list[str], value: str) -> None:
    if value not in warnings:
        warnings.append(value)


def _json_action(value: Any) -> dict | None:
    if isinstance(value, (dict, list)):
        parsed = value
    else:
        try:
            parsed = json.loads(_text(value))
        except (TypeError, ValueError):
            return None
    if isinstance(parsed, list) and len(parsed) == 1 and isinstance(parsed[0], dict):
        parsed = parsed[0]
    if not isinstance(parsed, dict):
        return None
    return {**parsed, "action": _text(parsed.get("action")).lower()}


def _same_action(left: Any, right: Any) -> bool:
    first, second = _json_action(left), _json_action(right)
    if first is not None and second is not None:
        return first == second
    return _text(left) == _text(right)


class _Sources:
    def __init__(self, store: ArtifactStore, source: dict, workbook_hash: str,
                 warnings: list[str], label: str):
        self.store, self.warnings, self.label = store, warnings, label
        self.rows: dict[str, dict] = {}
        self.collection_tasks: list[dict] = []
        self.manual_rows: dict[int, dict] = {}
        self.edits: dict = {}
        self.manual_known = False
        self.batch_id = ""
        self._seen: set[tuple[str, str, str]] = set()
        self._stack: set[tuple[str, str, str]] = set()
        self._collection_versions: set[str] = set()
        self._annotation_loaded = False
        reference = source.get("artifact")
        if not reference:
            _warn(warnings, f"{label}：缺少冻结修正元数据，人工精修数量及缺失分类未记录")
            return
        manifest, payload = self._read(reference)
        if manifest["stage"] != "07_cot":
            raise ConversionError(f"{label}：发布来源不是完整导出的 07_cot 版本")
        self.batch_id = manifest["batch_id"]
        excel = next((item for item in manifest["files"] if item["name"] == "full_dataset.xlsx"), None)
        if excel is None or excel.get("sha256") != workbook_hash:
            raise ConversionError(f"{label}：发布表与冻结完整导出版本校验值不一致")
        # Verify the registered full export as well as the independently frozen release copy.
        self.store.resolve_file(manifest, "full_dataset.xlsx")
        if source.get("id") and payload.get("session_id") != source["id"]:
            raise ConversionError(f"{label}：发布来源修正会话不一致")
        if payload.get("batch_id") != self.batch_id:
            raise ConversionError(f"{label}：修正快照批次不一致")
        groups, edits = payload.get("groups"), payload.get("row_edits")
        if isinstance(groups, list) and isinstance(edits, dict):
            self.manual_known = True
            self.edits = edits
            for group in groups:
                for row in group.get("rows", []):
                    index = row.get("excel_row")
                    if not isinstance(index, int) or index < 2 or index in self.manual_rows:
                        raise ConversionError(f"{label}：修正步骤源行号无效或重复")
                    self.manual_rows[index] = row
        else:
            _warn(warnings, f"{label}：冻结修正快照未记录完整人工编辑元数据")
        run_id = _text(payload.get("tree_run_id") or source.get("tree_run_id"))
        if source.get("tree_run_id") and run_id != source["tree_run_id"]:
            raise ConversionError(f"{label}：发布来源建树运行编号不一致")
        if not run_id:
            _warn(warnings, f"{label}：未记录建树来源，无法补充缺失分类")
            return
        matches = [item for item in store.list(self.batch_id, "04_tree")
                   if item.get("metadata", {}).get("run_id") == run_id]
        if len(matches) != 1:
            raise ConversionError(f"{label}：建树来源 {run_id} 必须对应唯一冻结版本")
        tree, tree_payload = self._read(matches[0])
        if tree_payload.get("run_id") != run_id:
            raise ConversionError(f"{label}：建树快照运行编号不一致")
        self._visit(tree)
        if not self.collection_tasks:
            _warn(warnings, f"{label}：未记录采集分类，缺失 App 和场景显示为未记录")

    def _read(self, reference: dict) -> tuple[dict, dict]:
        if not isinstance(reference, dict) or not all(reference.get(k) for k in ("batch_id", "stage", "version")):
            raise ConversionError(f"{self.label}：阶段来源缺少精确版本")
        manifest = self.store.get(reference["batch_id"], reference["stage"], reference["version"])
        if manifest is None:
            raise ConversionError(f"{self.label}：已引用的阶段版本不存在")
        if self.batch_id and manifest["batch_id"] != self.batch_id:
            raise ConversionError(f"{self.label}：阶段来源跨越批次")
        registered = {item["name"]: item for item in manifest["files"]}
        # Source refs carry frozen manifests; don't accept changed registrations.
        for item in reference.get("files", []):
            actual = registered.get(item.get("name"))
            if actual is None or any(item.get(k) != actual.get(k) for k in ("path", "sha256", "size")):
                raise ConversionError(f"{self.label}：冻结阶段引用与登记信息不一致")
        path = self.store.resolve_file(manifest, "result.json")
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ConversionError(f"{self.label}：阶段 JSON 格式无效")
        return manifest, payload

    def _visit(self, reference: dict) -> None:
        manifest, payload = self._read(reference)
        key = (manifest["batch_id"], manifest["stage"], manifest["version"])
        if key in self._stack:
            raise ConversionError(f"{self.label}：阶段来源存在循环")
        if key in self._seen:
            return
        self._stack.add(key)
        stage = manifest["stage"]
        if stage == "00_collection":
            self._collection_versions.add(manifest["version"])
            if len(self._collection_versions) > 1:
                raise ConversionError(f"{self.label}：同一发布来源引用多个采集分类版本")
            snapshot = payload.get("snapshot", {})
            tasks = snapshot.get("tasks")
            if not isinstance(tasks, list) or any(not isinstance(task, dict) for task in tasks):
                raise ConversionError(f"{self.label}：冻结采集分类格式无效")
            if payload.get("batch_id") != self.batch_id:
                raise ConversionError(f"{self.label}：采集分类批次不一致")
            self.collection_tasks = tasks
        if stage in {"02_annotation", "01_conversion"} and not self._annotation_loaded:
            sheets = payload.get("sheets")
            if not isinstance(sheets, dict):
                raise ConversionError(f"{self.label}：轨迹来源缺少步骤 JSON")
            for rows in sheets.values():
                for row in rows:
                    image = _image(row.get("image"))
                    if image:
                        if image in self.rows:
                            raise ConversionError(f"{self.label}：轨迹来源存在重复图片身份")
                        self.rows[image] = row
            self._annotation_loaded = True
        allowed = {
            "04_tree": {"03_observation"}, "03_observation": {"02_annotation"},
            "02_annotation": {"02_annotation", "01_conversion"}, "01_conversion": {"00_collection"},
            "00_collection": set(),
        }
        if stage not in allowed:
            raise ConversionError(f"{self.label}：不支持的轨迹来源阶段 {stage}")
        for ref in manifest.get("source_refs", []):
            if all(ref.get(k) for k in ("batch_id", "stage", "version")):
                if ref["stage"] not in allowed[stage]:
                    raise ConversionError(f"{self.label}：阶段来源顺序无效")
                self._visit(ref)
        self._stack.remove(key)
        self._seen.add(key)

    def task_for(self, row: dict) -> dict:
        case_id, task_id = _text(row.get("collection_case_id")), _text(row.get("task_id"))
        if not self.collection_tasks or not (case_id or task_id):
            return {}
        field, identifier = ("collection_case_id", case_id) if case_id else ("task_id", task_id)
        matches = [task for task in self.collection_tasks if _text(task.get(field)) == identifier]
        if len(matches) != 1:
            raise ConversionError(f"{self.label}：采集用例 {identifier} 无法唯一关联冻结任务")
        task = matches[0]
        for key in ("task_id", "source_result_id"):
            if row.get(key) and task.get(key) and _text(row[key]) != _text(task[key]):
                raise ConversionError(f"{self.label}：采集用例来源身份冲突")
        return task

    def manual(self, excel_row: int, values: dict) -> bool | None:
        if not self.manual_known:
            return None
        edits = self.edits.get(str(excel_row), {})
        if not isinstance(edits, dict):
            raise ConversionError(f"{self.label}：人工编辑记录格式无效")
        relevant = "actions" in edits or "sop" in edits
        if not relevant:
            return False
        original = self.manual_rows.get(excel_row)
        if original is None:
            raise ConversionError(f"{self.label}：人工编辑记录缺少冻结源行")
        if _image(original.get("image")) != _image(values.get("image")):
            raise ConversionError(f"{self.label}：人工编辑源行与最终发布表不一致")
        final_action = values.get("actions", values.get("action"))
        baseline_action = edits.get("original_actions", original.get("original_action",
                                       original.get("values", {}).get("actions",
                                       original.get("values", {}).get("action"))))
        action_changed = ("actions" in edits and baseline_action is not None
                          and _same_action(final_action, edits["actions"])
                          and not _same_action(final_action, baseline_action))
        baseline_sop = original.get("values", {}).get("sop", "")
        sop_changed = ("sop" in edits and "sop" in values
                       and _text(values["sop"]) == _text(edits["sop"])
                       and _text(values["sop"]) != _text(baseline_sop))
        if "actions" in edits and baseline_action is None and not sop_changed:
            _warn(self.warnings, f"{self.label}：部分动作缺少人工修改基线，精修数量未记录")
            return None
        return action_changed or sop_changed


def _classification_known(value: str) -> bool:
    # Preserve source text for filtering, but placeholders aren't real coverage.
    normalized = "".join(value.casefold().split())
    return bool(normalized) and normalized not in {
        "unclassified", "unknown", "未分类", "未知", "未记录", "未记录场景", "未记录app",
    }


def _field(row: dict, names: tuple[str, ...], fallback: Any) -> str:
    return next((_text(row.get(name)) for name in names if _text(row.get(name))), _text(fallback))



def _external_frozen_file(registry: Any, release_id: str, reference: dict, label: str) -> tuple[Path, bytes]:
    if not isinstance(reference, dict) or not reference.get("path") or not reference.get("sha256"):
        raise ConversionError(f"{label}：冻结文件登记信息不完整")
    path = registry.resolve_project_path(reference["path"])
    expected_root = (Path(registry.data_root) / "releases" / release_id).resolve()
    if not path.resolve().is_relative_to(expected_root):
        raise ConversionError(f"{label}：只允许读取该发布的冻结文件")
    if not path.is_file():
        raise ConversionError(f"{label}：冻结文件不存在")
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != reference["sha256"]:
        raise ConversionError(f"{label}：SHA256 校验失败")
    return path, data


def _convert_external_release(registry: Any, release: dict, release_id: str) -> dict:
    """Use the frozen parser result; never infer external provenance from sessions."""
    from .external_workbook import METADATA_FIELDS, PARSER_VERSION

    try:
        imported = release.get("external_import")
        if not isinstance(imported, dict) or imported.get("parser_version") != PARSER_VERSION:
            raise ConversionError("外部发布缺少受支持的导入记录")
        metadata = {key: imported.get(key) for key in METADATA_FIELDS}
        if any(not isinstance(value, str) for value in metadata.values()):
            raise ConversionError("外部导入元数据格式无效")
        try:
            if datetime.strptime(metadata["data_date"], "%Y-%m-%d").date().isoformat() != metadata["data_date"]:
                raise ValueError("date format")
        except ValueError as exc:
            raise ConversionError("外部导入数据日期无效") from exc
        if not metadata["data_source"] or not metadata["sheet_name"]:
            raise ConversionError("外部导入来源或工作表为空")
        files = release.get("excel_paths")
        if not isinstance(files, list) or len(files) != 1:
            raise ConversionError("每个外部发布必须对应一份原始步骤表")
        original, _ = _external_frozen_file(registry, release_id, files[0], "外部步骤表")
        if original.suffix.lower() not in {".xlsx", ".xlsm"} or files[0]["sha256"] != imported.get("source_sha256"):
            raise ConversionError("外部步骤表格式或导入来源校验失败")
        canonical_ref, manifest_ref = imported.get("canonical"), imported.get("manifest")
        _, canonical_data = _external_frozen_file(registry, release_id, canonical_ref, "规范化数据")
        _, manifest_data = _external_frozen_file(registry, release_id, manifest_ref, "导入清单")
        canonical = json.loads(canonical_data)
        manifest = json.loads(manifest_data)
        if not isinstance(canonical, dict) or not isinstance(manifest, dict):
            raise ConversionError("外部规范化数据或导入清单格式无效")
        for value in (canonical, manifest):
            if value.get("schema_version") != 1 or value.get("parser_version") != PARSER_VERSION:
                raise ConversionError("外部规范化数据或导入清单版本无效")
            if value.get("metadata") != metadata:
                raise ConversionError("外部导入元数据与冻结文件不一致")
            if not imported.get("content_hash") or value.get("content_hash") != imported["content_hash"]:
                raise ConversionError("外部导入内容指纹与冻结文件不一致")
        batches = release.get("batch_ids")
        if (not isinstance(batches, list) or len(batches) != 1 or not batches[0]
                or manifest.get("release_id") != release_id or manifest.get("batch_id") != batches[0]
                or not imported.get("import_id") or manifest.get("import_id") != imported["import_id"]
                or manifest.get("sheet_name") != metadata["sheet_name"]):
            raise ConversionError("外部导入清单的发布、批次或导入身份不一致")
        for field, expected in (("source", files[0]), ("canonical", canonical_ref)):
            actual = manifest.get(field)
            if not isinstance(actual, dict) or any(actual.get(key) != expected.get(key) for key in ("path", "sha256")):
                raise ConversionError("外部导入清单与冻结文件登记不一致")
        rows = canonical.get("rows")
        if not isinstance(rows, list) or not rows:
            raise ConversionError("外部规范化数据没有可统计的轨迹")
        result, identities = [], set()
        for row in rows:
            if not isinstance(row, dict):
                raise ConversionError("外部规范化轨迹格式无效")
            identity, steps, actions = row.get("trajectory_identity"), row.get("step数量"), row.get("action_box")
            if not isinstance(identity, str) or not identity or identity in identities:
                raise ConversionError("外部规范化轨迹身份无效或重复")
            identities.add(identity)
            if (type(steps) is not int or steps < 1 or not isinstance(actions, dict)
                    or any(not isinstance(key, str) or not key or type(value) is not int or value < 1
                           for key, value in actions.items()) or sum(actions.values()) != steps):
                raise ConversionError("外部规范化轨迹的步骤或动作数量无效")
            if (row.get("时间") != metadata["data_date"] or row.get("生产方式") != metadata["data_source"]
                    or row.get("人工精修步骤数量") is not None or row.get("manual_known") is not False):
                raise ConversionError("外部规范化轨迹的日期、来源或精修信息不一致")
            for field, known in (("APP", "app_known"), ("一级场景", "level1_known"), ("二级场景", "level2_known")):
                value = row.get(field)
                if not isinstance(value, str) or not value or row.get(known) is not _classification_known(value):
                    raise ConversionError("外部规范化轨迹分类无效")
            unique_id = f"{release_id}/001/{quote(identity, safe='')}"
            result.append({**row, "轨迹": unique_id, "子任务轨迹": f"{unique_id}/subtask-1",
                           "release_id": release_id, "source_file_index": 1})
        warnings = canonical.get("warnings")
        if not isinstance(warnings, list):
            raise ConversionError("外部规范化数据缺少校验信息")
        messages: list[str] = []
        for warning in warnings:
            if not isinstance(warning, dict) or not isinstance(warning.get("message"), str):
                raise ConversionError("外部规范化数据校验信息格式无效")
            location = _text(warning.get("sheet"))
            if warning.get("row") is not None:
                location += f" 第 {warning['row']} 行"
            _warn(messages, f"{location}：{warning['message']}" if location else warning["message"])
        return {"release_id": release_id, "rows": result, "warnings": messages}
    except ConversionError:
        raise
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ConversionError(f"外部发布数据转换失败：{exc}") from exc


def convert_release(registry: Any, release: dict) -> dict:
    """Convert all frozen full export files, using only each active business sheet."""
    release_id = _text(release.get("release_id"))
    if not release_id or any(value in release_id for value in ("/", "\\", "..")):
        raise ConversionError("发布记录缺少有效发布编号")
    if release.get("source_kind") == "external_manual":
        return _convert_external_release(registry, release, release_id)
    try:
        created = datetime.fromisoformat(_text(release.get("created_at")).replace("Z", "+00:00"))
        if created.tzinfo is None:
            raise ValueError("missing timezone")
        date = created.astimezone(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError) as exc:
        raise ConversionError("发布记录时间无效或缺少时区") from exc
    files = release.get("excel_paths")
    if not isinstance(files, list) or not files:
        raise ConversionError("发布记录没有完整数据表")
    refs = release.get("source_refs") or []
    if refs and len(refs) != len(files):
        raise ConversionError("发布文件与冻结来源数量不一致")
    warnings: list[str] = []
    result: list[dict] = []
    try:
        # Published statistics read only the self-contained release lineage.
        store = FrozenReleaseSources(registry.data_root, release) if any(
            ref.get("artifact") for ref in refs) else None
        for file_index, item in enumerate(files, start=1):
            label = f"文件 {file_index}"
            path = registry.resolve_project_path(item.get("path", ""))
            expected_root = (Path(registry.data_root) / "releases" / release_id).resolve()
            if not path.resolve().is_relative_to(expected_root):
                raise ConversionError(f"{label}：统计只允许读取该发布的冻结文件")
            if not path.is_file():
                raise ConversionError(f"{label}：发布表不存在")
            if path.suffix.lower() not in {".xlsx", ".xlsm"} or not item.get("sha256") or _sha(path) != item["sha256"]:
                raise ConversionError(f"{label}：发布表格式或 SHA256 校验失败")
            source = refs[file_index - 1] if refs else {}
            metadata = _Sources(store, source, item["sha256"], warnings, label)
            workbook = load_workbook(path, read_only=True, data_only=False)
            groups: dict[str, dict] = {}
            try:
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                headers = [_text(cell) for cell in next(values, ())]
                if len(set(headers)) != len(headers) or not any(
                        key in headers for key in ("trajectory_id", "文件夹名", "meta_task", "轨迹")):
                    raise ConversionError(f"{label}：发布表轨迹编号列缺失或表头重复")
                if "action" not in headers and "actions" not in headers:
                    raise ConversionError(f"{label}：发布表缺少最终动作列")
                for row_number, cells in enumerate(values, start=2):
                    if not any(cell is not None and cell != "" for cell in cells):
                        continue
                    row = dict(zip(headers, cells))
                    image = _image(row.get("image"))
                    frozen = metadata.rows.get(image, {})
                    if metadata.rows and image and not frozen:
                        raise ConversionError(f"{label}：第 {row_number} 行不属于冻结轨迹来源")
                    identity = _field(row, ("trajectory_id",), frozen.get("trajectory_id"))
                    if not identity:
                        original = _field(row, ("文件夹名", "meta_task", "轨迹"), "")
                        if not original:
                            raise ConversionError(f"{label}：第 {row_number} 行缺少轨迹编号")
                        identity = f"{PurePosixPath(image).parent.as_posix()}/{original}" if image else original
                    task = metadata.task_for(frozen or row)
                    classifications = (
                        _field(row, ("APP", "app", "涉及APP"), task.get("app")),
                        _field(row, ("一级场景", "scene"), task.get("scene")),
                        _field(row, ("二级场景", "capability"), task.get("capability")),
                    )
                    group = groups.setdefault(identity, {"steps": 0, "actions": Counter(),
                        "manual": 0, "manual_known": True, "fields": [set(), set(), set()]})
                    group["steps"] += 1
                    for target, value in zip(group["fields"], classifications):
                        if value:
                            target.add(value)
                    action = _json_action(row.get("actions", row.get("action")))
                    action_name = _text(action.get("action")) if action else ""
                    if not action_name:
                        action_name = "unknown"
                        _warn(warnings, f"{label}：存在无法解析的最终动作，已计入 unknown")
                    group["actions"][action_name] += 1
                    manual = metadata.manual(row_number, row)
                    if manual is None:
                        group["manual_known"] = False
                    elif manual:
                        group["manual"] += 1
                for identity, group in groups.items():
                    if any(len(values) > 1 for values in group["fields"]):
                        raise ConversionError(f"{label}：同一轨迹存在冲突的 App 或场景分类")
                    app, level1, level2 = [next(iter(values), "") for values in group["fields"]]
                    unique_id = f"{release_id}/{file_index:03d}/{quote(identity, safe='')}"
                    result.append({
                        "轨迹": unique_id, "子任务轨迹": f"{unique_id}/subtask-1",
                        "step数量": group["steps"], "APP": app or "未记录 App",
                        "一级场景": level1 or "未分类", "二级场景": level2 or "未分类",
                        "时间": date, "生产方式": "数据飞轮",
                        "人工精修步骤数量": group["manual"] if group["manual_known"] else None,
                        "action_box": dict(sorted(group["actions"].items())),
                        "release_id": release_id, "source_file_index": file_index,
                        "trajectory_identity": identity, "manual_known": group["manual_known"],
                        "app_known": _classification_known(app), "level1_known": _classification_known(level1),
                        "level2_known": _classification_known(level2),
                    })
            finally:
                workbook.close()
            if _sha(path) != item["sha256"]:
                raise ConversionError(f"{label}：转换期间发布表发生变化")
    except ConversionError:
        raise
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise ConversionError(f"发布数据转换失败：{exc}") from exc
    if not result:
        raise ConversionError("发布表没有可统计的轨迹步骤")
    return {"release_id": release_id, "rows": result, "warnings": warnings}

