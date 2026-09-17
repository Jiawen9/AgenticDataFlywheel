from pathlib import Path
import json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT = Path(__file__).resolve().parents[2]
OUT = PROJECT / "outputs"
QA = OUT / "docx_design_qa"
QA.mkdir(parents=True, exist_ok=True)
ROOT = "backend_workspace"
FONT = "Microsoft YaHei"
FONT_FILE = r"C:\Windows\Fonts\msyh.ttc"
BOLD_FILE = r"C:\Windows\Fonts\msyhbd.ttc"
TEAL = "#0f766e"
LINE = "#cbd5e1"
INK = "#111827"
MUTED = "#475569"

def font(size, bold=False):
    return ImageFont.truetype(BOLD_FILE if bold else FONT_FILE, size)
def text_center(draw, xy, text, size=32, bold=False, color=INK):
    box = draw.multiline_textbbox((0,0), text, font=font(size,bold), spacing=12, align="center")
    draw.multiline_text((xy[0]-(box[2]-box[0])/2-box[0],xy[1]-(box[3]-box[1])/2-box[1]),text,font=font(size,bold),fill=color,spacing=12,align="center")
def box(draw, rect, title, sub="", fill="#f8fafc", outline=LINE, size=33):
    draw.rounded_rectangle(rect, radius=15, fill=fill, outline=outline, width=3)
    cx=(rect[0]+rect[2])/2
    if sub:
        text_center(draw,(cx,rect[1]+35),title,size,True)
        sub_box=draw.multiline_textbbox((0,0),sub,font=font(27),spacing=12)
        text_center(draw,(cx,rect[3]-18-(sub_box[3]-sub_box[1])/2),sub,27,color=MUTED)
    else: text_center(draw,(cx,(rect[1]+rect[3])/2),title,size,True)
def arrow(draw, start, end, color=TEAL, width=5):
    draw.line([start,end], fill=color,width=width)
    angle=math.atan2(end[1]-start[1],end[0]-start[0])
    points=[end,(end[0]-18*math.cos(angle-.5),end[1]-18*math.sin(angle-.5)),(end[0]-18*math.cos(angle+.5),end[1]-18*math.sin(angle+.5))]
    draw.polygon(points,fill=color)

def business_figure():
    im=Image.new("RGB",(1440,1630),"white");d=ImageDraw.Draw(im)
    box(d,(60,20,680,132),"普通任务生成","场景树与种子任务 → 生成和人工审核")
    box(d,(760,20,1380,132),"失败用例扩增","冻结知识库 → 场景匹配 → 变体审核")
    arrow(d,(370,132),(500,180));arrow(d,(1070,132),(950,180))
    box(d,(340,180,1100,290),"提交轨迹采集 00_collection","每个生成或扩增作业首次提交冻结 JSON 与 17 列 Excel",fill="#f0fdfa",outline=TEAL,size=32)
    arrow(d,(720,290),(720,335))
    box(d,(340,335,1100,455),"采集运行与完成登记","下发批次和运行编号 → 原文件落地 → 清单及 SHA256 核验",size=30)
    arrow(d,(720,455),(720,505))
    box(d,(340,505,1100,617),"01 轨迹转换","读取冻结原始清单 → 步骤 JSON 与轨迹 Excel",fill="#f0fdfa",outline=TEAL)
    box(d,(35,485,290,635),"CLI 支路","显式指定\n自备原轨迹",size=30)
    arrow(d,(290,560),(340,560))
    arrow(d,(720,617),(720,665))
    box(d,(340,665,1100,777),"02 动作标框","生成和复核动作框 → 标框 JSON 与 Excel",fill="#f0fdfa",outline=TEAL)
    box(d,(55,825,685,945),"03 Observation","保留全部步骤与广告 中间态判断",size=31)
    box(d,(755,825,1385,945),"04 轨迹建树","树 JSON 与同版本质检输入",size=31)
    arrow(d,(685,885),(755,885))
    d.line([(720,777),(720,803),(370,803)],fill=TEAL,width=5)
    arrow(d,(370,803),(370,825))
    arrow(d,(1070,945),(1070,983));d.line([(1070,983),(720,983)],fill=TEAL,width=5)
    arrow(d,(720,983),(720,1010))
    box(d,(340,1010,1100,1122),"05 轨迹质检","评分标准与逐步评分 → 质检 JSON 与 Excel",fill="#f0fdfa",outline=TEAL)
    box(d,(55,1170,685,1290),"06 人工修正","当前编辑进 SQLite → 保存完整过程版本",size=31)
    box(d,(755,1170,1385,1290),"07 COT","修正步生成 → 完整过程表与数据集导出",size=31)
    d.line([(720,1122),(720,1146),(370,1146)],fill=TEAL,width=5)
    arrow(d,(370,1146),(370,1170))
    arrow(d,(685,1230),(755,1230))
    arrow(d,(1070,1290),(1070,1340));d.line([(1070,1340),(720,1340)],fill=TEAL,width=5)
    arrow(d,(720,1340),(720,1365))
    box(d,(80,1365,850,1485),"数据发布","release_id → 冻结所选完整数据集 Excel",fill="#f0fdfa",outline=TEAL)
    arrow(d,(850,1425),(915,1425))
    box(d,(915,1365,1385,1485),"云道S3上传","接入方实现登录 上传与远端去重",fill="#f8fafc",size=31)
    text_center(d,(720,1560),"箭头表示数据关系  每一步仍由各自按钮或作业触发  本图不代表全部自动执行",25,color=MUTED)
    target=QA/"business_flow.png";im.save(target);return target

def storage_figure():
    im=Image.new("RGB",(1440,1300),"white");d=ImageDraw.Draw(im)
    box(d,(365,10,1075,110),"统一数据根目录 backend_workspace","ADF_DATA_ROOT 可在服务启动前覆盖",fill="#f0fdfa",outline=TEAL,size=31)
    arrow(d,(720,110),(720,158))
    box(d,(460,158,980,270),"system/app.sqlite","当前作业 人工编辑 结果索引 发布与回执",fill="#e0f2fe",outline="#94a3b8",size=32)
    box(d,(30,340,440,485),"raw","原轨迹 截图 XML\n保持实际源文件",size=32)
    box(d,(515,340,925,485),"resources 与 inputs","场景树等配置资源\n明确上传的输入表",size=30)
    box(d,(1000,340,1410,485),"system 工作文件","冻结输入 运行输出\n导出与诊断文件",size=30)
    arrow(d,(720,270),(720,325));d.line([(235,325),(1205,325)],fill=TEAL,width=4)
    arrow(d,(235,325),(235,340));arrow(d,(720,325),(720,340));arrow(d,(1205,325),(1205,340))
    box(d,(105,585,1335,860),"batches/<batch_id>/<stage>/<version>","result.json 为下游完整数据     Excel 为核对投影\nmanifest.json 记录来源关系 文件清单 SHA256 与版本",fill="#f0fdfa",outline=TEAL,size=34)
    arrow(d,(235,485),(235,585));arrow(d,(720,485),(720,585));arrow(d,(1205,485),(1205,585))
    text_center(d,(720,695),"完整写入并登记后才对外可见   旧版本不覆盖",30,True,color=TEAL)
    arrow(d,(720,860),(720,925))
    box(d,(100,925,865,1070),"releases/<release_id>","manifest.json + 按序号保存的冻结 Excel\n上传与发布下载读取同一份字节",fill="#e0f2fe",size=31)
    box(d,(930,925,1370,1070),"cache logs tmp","加速缓存 运行日志\n临时未完成文件",size=32)
    text_center(d,(720,1160),"SQLite 管当前状态  JSON 传完整版本  Excel 供人查看和交付",32,True)
    text_center(d,(720,1230),"raw 和 resources 仍必须保留  单独备份数据库或 Excel 都不完整",27,color=MUTED)
    target=QA/"storage_relationship.png";im.save(target);return target

doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.7);section.bottom_margin=Inches(.65)
section.left_margin=Inches(.8);section.right_margin=Inches(.8)
section.footer_distance=Inches(.25)
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Caption"]:
    st=doc.styles[name]
    st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.18
doc.styles["Normal"].font.size=Pt(10.8)
doc.styles["Title"].font.size=Pt(25)
doc.styles["Title"].paragraph_format.space_after=Pt(12)
doc.styles["Subtitle"].font.size=Pt(12)
doc.styles["Heading 1"].font.size=Pt(17)
doc.styles["Heading 1"].paragraph_format.space_after=Pt(10)
doc.styles["Heading 2"].font.size=Pt(12.5)
doc.styles["Heading 2"].paragraph_format.space_before=Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after=Pt(5)
doc.styles["Caption"].font.size=Pt(9.5)
doc.styles["Caption"].paragraph_format.space_after=Pt(8)
for name in ["Code"]:
    if name not in doc.styles:doc.styles.add_style(name,1)
    doc.styles[name].font.name="Consolas";doc.styles[name].font.size=Pt(9.5)
    doc.styles[name]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
    doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.styles[name].paragraph_format.line_spacing=1.08
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");r._r.addnext(fld)
footer.style=doc.styles["Caption"]

def p(text="",boldlead=None,style=None):
    para=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        para.add_run(boldlead).bold=True;para.add_run(text[len(boldlead):])
    else:para.add_run(text)
    return para
def h(text,level=1):doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title)
def code(text):p(text,style="Code")
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[6.9/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,label,w in zip(t.rows[0].cells,headers,widths):c.text=label;c.width=Inches(w)
    header=t.rows[0]._tr.get_or_add_trPr();flag=OxmlElement("w:tblHeader");header.append(flag)
    for values in rows:
        for c,value,w in zip(t.add_row().cells,values,widths):c.text=str(value);c.width=Inches(w)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");trpr.append(cant)
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right","insideH","insideV"]:
                element=OxmlElement("w:"+edge);element.set(qn("w:val"),"single");element.set(qn("w:sz"),"5");element.set(qn("w:color"),"D9D9D9");borders.append(element)
            props.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge,value in [("top",85),("bottom",85),("left",115),("right",115)]:
                e=OxmlElement("w:"+edge);e.set(qn("w:w"),str(value));e.set(qn("w:type"),"dxa");margins.append(e)
            props.append(margins)
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"DDEBF7" if ri==0 else ("FFFFFF" if ri%2 else "F8FAFC"));props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.12
                if ci==0 and len(headers)>2:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    p("")
    return t
def picture(path,width,alt):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=para.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set("descr",alt)
    return shape

p("数据流转与存储设计",style="Title")
p("AgenticDataFlywheel 业务流程和维护说明",style="Subtitle")
p("适用对象：业务使用者、后端维护工程师和采集及云道S3接入同事\n版本日期：2026年9月16日")
p("系统统一使用 backend_workspace 作为运行数据根目录。模块之间读取结构化 JSON，当前作业和人工编辑保存在 SQLite；各阶段保留可下载的 Excel 核对表。原轨迹、图片和 XML 继续保存为文件，并通过批次、轨迹和版本标识关联。这样既能按批次追溯全过程，也能维持现有人工审核和表格交付方式。")
h("三种数据形式的职责",2)
table(["形式","保存内容","使用方式"],[
["SQLite","当前作业状态、结果、人工编辑、索引和上传回执","通过后端事务读写；不是供用户手改的表格"],
["JSON","冻结输入、完整步骤、阶段结果、上游关系和版本信息","作为模块间数据协议；必需文件缺失或校验失败时停止"],
["Excel","阶段核对表、采集输入表及训练数据交付表","供查看、下载和交付；修改已导出表不会回写当前数据"],
],[1.0,2.85,3.05])
h("最重要的使用规则",2)
p("先通过页面保存人工修改，再导出或提交下游。一次阶段完成会新增版本，已经发布的版本保持不变。文件名、批次号和页面任务状态各有用途，不能用某个 Excel 的存在代替成功状态，也不能根据目录名称猜测任务身份。")
p("业务 API 地址继续保留。停用的是旧目录扫描和旧文件回退；后端不再从历史工作目录、旧 JSON 或旧 Excel 中自动恢复当前结果。业务下载响应指向本次保存的文件，浏览器下载副本仍进入用户自己的下载目录。")
h("本次真实验证范围",2)
p("实例批次 validation-20260915-155706 已从自备原始轨迹处理到 05_quality，包含 10 条轨迹和 142 个步骤。本次没有重新执行真实手机采集，没有继续修正、COT、数据发布或云道S3上传。作业技术完成仍需要人工核查模型判断。")

page("01 业务全链路")
picture(business_figure(),6.7,"普通生成和失败扩增冻结为采集批次，采集完成登记进入转换和标框，建树产生Observation与树，后续质检修正COT发布再上传云道S3；自备原轨迹通过CLI进入转换")
p("图1 业务数据关系与触发顺序",style="Caption")
p("页面预处理只执行转换和标框。Observation 与中间态判断在建树任务中执行，随后保存 03 和 04 两个阶段版本。质检、修正、COT、发布和上传仍由各自入口触发，不因前一阶段完成而自动全部执行。")

page("02 存储关系与目录")
picture(storage_figure(),6.6,"SQLite保存当前状态与索引，raw和resources保留源文件与资源，system提供工作文件，batches冻结JSON Excel和manifest，releases冻结发布Excel，cache日志和临时文件辅助运行")
p("图2 当前状态 工作文件 阶段版本与发布版本的关系",style="Caption")
p("默认绝对根目录为 D:\\work\\Code\\AgenticDataFlywheel\\backend_workspace。文中后续路径均相对这个根目录；部署时可在启动前设置 ADF_DATA_ROOT。backend_workspace 下直接放 system、batches 等目录，不再套一层 data。")
p("system 是业务工作区，其中可能有输入快照、运行 JSON、导出 Excel 和诊断文件；batches 是按阶段登记的可追溯版本。可变作业状态与人工编辑只进入 system/app.sqlite，不双写成另一套可变 JSON 状态库。")
code("batches/<batch_id>/<stage>/<version>/\n  result.json    result.xlsx 或原导出文件名    manifest.json")
p("清单中的 files 给出实际文件名、相对路径、大小和 SHA256；source_refs 记录上游版本及资源。建树阶段主要保存 JSON，不要求为树强行增加一份无业务意义的 Excel。")

page("03 任务生成 扩增与采集批次")
h("生成和扩增如何形成输入",2)
p("普通生成读取场景树和种子任务，扩增先上传失败用例并冻结知识库快照，再分类并保存关联路径。点击开始扩增后直接消费已保存的分类结果，不再次分类。页面审核、人工修改和删除更新 SQLite 当前值；筛选、分页和折叠只影响展示。")
p("任务生成结束保存 00_task_generation，扩增分类结束保存 00_scene_matching，扩增完成保存 00_augmentation。人工保存不会为每次按键生成 Excel；需要额外中间表时可请求当前结果 snapshot，不调用模型。")
h("第一次提交采集冻结什么",2)
p("普通生成和扩增共用 collection-batch 接口。只有 succeeded 或 partial 且存在有效未删除任务时可首次提交。系统校验编号、任务、App 和依赖引用，冻结当前全部未删除结果；后续修改源任务不会改动已提交批次，重复提交返回同一个批次。")
code("system/task_generation/collection_batches/<batch_id>/\n  batch.json    collection-batch-<batch_id>.xlsx\nbatches/<batch_id>/00_collection/<version>/")
table(["采集列","来源","规则"],[
["用例编号","collection_case_id","普通生成沿用 task_id；扩增优先原 case_id，缺失回退 task_id"],
["单APP跨APP","固定值","单APP"],
["涉及APP","app","对应场景树 L4"],
["一级场景","scene","对应 L1"],
["二级场景","capability","对应 L2"],
["三级场景","sub_capability","对应 L3"],
["任务","task","采用提交时已保存的文本"],
],[1.2,1.4,4.3])
p("后十列保留表头并留空：难易程度、页面关键元素、预期步数、用例来源、构建人、任务结束标志、任务复杂度、关键动作、SOP、动态变量说明。")
p("JSON 继续保留主任务、前置任务、强依赖和原行状态。17 列 Excel 不承担依赖调度，也不能把扩增任务默认解释成无依赖。")

page("04 采集完成与轨迹预处理")
h("原始文件先落地 再登记完成",2)
p("采集侧选择批次本身不启动手机。点击运行后先幂等登记采集 Excel，再下发 batch_id、collection_run_id 和 output_dir。远端采集服务必须把最终文件写到本后端能够读取的位置，并通过完成接口回传全部文件的相对路径和 SHA256。路径字段不会自动传输文件。")
code("raw/collection_batches/<batch_id>/runs/<collection_run_id>/\n  <collection_case_id>/<原轨迹目录>/ 原始 JSON 图片 XML")
p("完成登记校验批次、用例身份、完整目录和文件字节。相同清单重复回传不产生新记录；补采使用新的 collection_run_id。只有 completed 清单中的可用轨迹进入预处理，下发回执或页面运行中状态都不等于轨迹已就绪。")
h("转换与动作标框",2)
p("开始预处理时冻结该批次当前所有可用采集结果，保存 input.json 及输入摘要。01_conversion 将 response、图片和 XML 等转换为统一步骤 JSON 与轨迹 Excel，不调用模型。02_annotation 读取转换 JSON，生成并复核动作框，保存完整标框 JSON 与 Excel。标框失败时已完成的转换版本仍可下载。")
code("system/preprocessing/<batch_id>/<job_id>/input.json\nbatches/<batch_id>/01_conversion/<version>/\nbatches/<batch_id>/02_annotation/<version>/")
p("点击保存 bbox 时携带 batch_id 和 annotation_version，后端校验当前版本后发布新的 02。页面继续使用明确版本读取任务、轨迹和图片；建树提交也冻结这个版本，避免跨批次或同名轨迹串用。")
h("自备原轨迹的 CLI 入口",2)
p("自备数据先显式复制到 raw/rollout_trajectories，保留任务和轨迹层级。该支路不补造生成任务或真实采集登记；脚本尊重显式输入路径。")
code('python backend/trajectories_preprocessing.py --source "backend_workspace/raw/rollout_trajectories" --batch-id my-batch\n\npython backend/export_vla_trajectories.py "backend_workspace/raw/rollout_trajectories" --batch-id my-batch')
p("第一条执行转换和标框，第二条只转换。默认工作文件按批次与本次执行隔离；显式 --export-output 或 --annotated-output 才覆盖默认输出位置。")

page("05 Observation 建树与质检")
table(["阶段","输入与触发","保存内容"],[
["03_observation","页面提交建树后读取冻结的 02 JSON 和对应图片","所有步骤的 Observation、中间态或广告类别、判断原因和是否计入树；JSON 与 Excel"],
["04_tree","与 03 属于同一建树作业，基于已保存视觉判断构建","树 JSON、来源轨迹身份以及同版本完整质检输入"],
["05_quality","在轨迹质检页选任务集和任务后启动","评分标准、逐步与全轨迹评分 JSON；轨迹汇总和步骤等 Excel"],
],[1.45,2.4,3.05])
p("阶段路径分别为 batches/<batch_id>/03_observation/<version>、04_tree/<version> 和 05_quality/<version>。文件清单以 manifest 为准；不能只根据页面节点数量判断原步骤是否丢失。")
h("完整过程记录和树投影",2)
p("03 保留全部原步骤，即使广告、加载或其他中间态没有进入树也不会从过程表消失。04 保存用于展示与后续质检的树结构，并保留完整质检输入。树节点可以归并共享状态，节点数与原步骤数不必相等；比较时应使用明确的轨迹和步骤身份。")
p("同一次建树保存 batch_id、annotation_version、原始资源根和来源清单。新采集数据在提交、执行和发布边界接受来源校验；运行后新补采的轨迹不会悄悄加入已经冻结的任务集。")
h("质检结果如何进入修正",2)
p("质检读取选定 run_id 的完整输入，生成评分标准并对轨迹评分。质检作业完成后保存 05，再根据现有推荐和用户选择创建修正会话；修正会话冻结输入和原始资源关系。模型评分及推荐不等于人工批准，进入修正前仍应查看关键失败原因和被忽略步骤。")
code("system/trajectory_tree_runs/<run_id>/\nsystem/trajectory_quality_results/<run_id>/\nsystem/trajectory_correction/inputs/")
p("这些 system 文件属于运行和输入工作区；查看某一阶段已经交付的结果时，应读取对应 batches 版本。必需 JSON 缺失或校验不通过时明确失败，不改读 Excel，也不寻找其他任务集替代。")

page("06 修正 COT与数据发布")
h("人工修改独立保存",2)
p("修正会话的动作、Summary、Thought、bbox、删除状态和导出选择保存到 SQLite。修改 Summary 只修改 Summary，修改 Thought 只修改 Thought，不清除另一字段已经保存的生成结果。显式保存的空字符串和手动改回原始文本都保留为人工选择。")
p("Summary 和 Thought 分别按人工值、与当前动作匹配的生成值、原始值取值。动作修改后的旧生成结果匹配校验、主动重新生成的现有行为继续保留；不是所有原始步骤都会自动补生成 COT。")
h("过程快照与训练导出的区别",2)
p("在修正导出、完整导出或提交 COT 前保存 06_correction；COT 完成或完整数据集导出保存 07_cot。06 和 07 的过程 JSON 与 Excel 保留完整步骤、人工选择和已删除记录，便于追溯；真正的训练导出仍执行已有删除、勾选和分流规则。")
table(["导出类型","当前规则"],[
["精修与筛查","无动作修改但有 SOP 修改时进入人工精修；否则按原质量进入原生分支。有动作修改时，首个修正动作之前的前缀进入 SFT；存在第二个修正动作时按既有规则形成 RL 负向反思前缀。"],
["完整数据集","在原表结构上覆盖会话最终值，沿用现有组选择和删除规则；与完整过程快照不是同一种投影。"],
["阶段过程表","用于核对过程及追溯，不直接视为可训练或已审核数据。"],
],[1.4,5.5])
h("发布冻结与云道S3上传",2)
p("创建发布记录时，后端读取已导出的完整数据集 Excel，校验后按发布编号冻结副本。多个会话保留多份表，不合并，不上传原始轨迹，也不临时转换为 JSON。")
code("releases/<release_id>/manifest.json\nreleases/<release_id>/001/<原文件名>.xlsx")
p("页面云道S3上传只提交 release_id 和 target: internal。后端把经过校验的绝对文件路径、文件名、序号、SHA256 和稳定幂等键交给 internal_uploader.py。默认实现尚未配置，按钮不可用；目标名云道S3并不意味着本次实现了通用 S3 协议。")

page("07 标识符与版本关系")
table(["标识","含义","关联规则"],[
["job_id","一次生成 扩增 预处理 建树 质检或上传作业","只标识执行记录，不等于每种业务的数据版本"],
["batch_id","持续追踪同一批数据","生成采集批次等于来源生成或扩增 job_id；CLI 可显式指定"],
["task_id","生成侧稳定任务身份","优先原 task_uuid，缺失采用 result_id；重复读取不重建"],
["source_result_id","来源审核结果身份","用于从采集及后续轨迹追溯生成侧记录"],
["collection_case_id","采集 Excel 的用例编号","通过冻结批次映射回 task_id，不按目录名称猜测"],
["collection_run_id","同一批次的一次正式采集","一次补采新运行；网络重试使用原请求键和运行号"],
["trajectory_id","批次内供系统使用的稳定轨迹身份","新采集使用 tr_ 加来源组合摘要；原目录名另存 source_trajectory_id"],
["run_id","一次建树得到的任务集","后续质检和修正绑定它及所用标框版本"],
["version","阶段冻结版本","与 batch_id 和 stage 共同定位一个 manifest"],
["session_id 与 release_id","修正会话和发布记录","会话绑定冻结输入；发布绑定已导出 Excel 的冻结副本"],
],[1.6,2.25,3.05])
p("例子：同一个用例分两次采集，每次都产生名为 run-1 的原轨迹。两条数据的 collection_run_id 不同，因此内部 trajectory_id 不同；页面可以显示相同原名，但不能合并它们。步骤继续保留原 step 编号及来源行信息。")
p("新采集稳定轨迹摘要基于 batch_id、collection_run_id、task_id 和 relative_dir 的规范 JSON。所有下游保留显式来源字段；历史自备轨迹缺少采集字段时沿用其原身份，不伪造采集运行。")
p("版本使用时间与唯一后缀生成，但业务代码应把它当作标识符使用。文件必须通过 manifest 或业务接口解析；维护者和上传接入方不需要从编号中推算目录。")

page("08 关键接口与下载方式")
p("所有路径均从后端根地址开始。接口只负责各自动作；GET 查询不补生成模型结果，也不通过旧目录寻找缺失数据。请求结构错误通常为 422，不存在为 404，状态或数据完整性冲突为 409。")
table(["用途","接口"],[
["当前结果与采集冻结","GET /api/task-generation/jobs/{job_id}/collection-input\nPOST /api/task-generation/jobs/{job_id}/collection-batch"],
["采集批次查询与表格","GET /api/task-generation/collection-batches\nGET /api/task-generation/collection-batches/{batch_id}/workbook"],
["采集运行与完成登记","POST /api/phone-factory/remote/start-run\nPOST /api/phone-factory/collection-runs/{collection_run_id}/complete"],
["预处理查询与启动","GET /api/trajectory-preprocessing/batches\nPOST /api/trajectory-preprocessing/jobs"],
["预处理恢复与重试","GET /api/trajectory-preprocessing/jobs/{job_id}\nPOST /api/trajectory-preprocessing/jobs/{job_id}/retry"],
["建树与质检","POST /api/tree-builds\nPOST /api/quality-jobs"],
["修正和完整数据集导出","POST /api/correction/sessions/{session_id}/export\nPOST /api/correction/sessions/{session_id}/dataset-export"],
["发布与上传","POST /api/dataset-releases\nPOST /api/dataset-releases/{release_id}/upload"],
],[1.45,5.45])
h("统一阶段产物接口",2)
code("GET /api/data-storage\nGET /api/data-batches\nGET /api/data-batches/{batch_id}/artifacts\nGET /api/data-batches/{batch_id}/artifacts/{stage}/{version}\nGET /api/data-batches/{batch_id}/artifacts/{stage}/{version}/files/{filename}")
p("业务导出 URL 保持原样，后端将工作文件保存到 system 并登记对应阶段版本。查完整过程用统一产物接口，取训练导出用该次业务响应的 download_url。发布下载和上传使用 releases 中同一份冻结 Excel，并验证 SHA256。")

page("09 冻结 重试与并发规则")
h("三个不同的冻结时点",2)
p("提交采集冻结生成结果；开始预处理冻结已经完成登记的原轨迹范围；提交建树冻结所选标框版本。三次冻结解决的是不同边界。后续源任务编辑、新补采或者新的人工标框版本，都不会自动改写已经启动的下游任务。")
h("失败后怎样继续",2)
table(["情况","继续方式"],[
["同批次重复点击启动","锁内返回现有运行中作业；相同输入和配置已有成功结果时复用它"],
["标框失败","重试原预处理作业，保留已完成的转换版本并复用有效缓存"],
["服务重启","运行中作业标记中断；查询恢复记录，由用户选择重试"],
["代码或模型配置变化","原作业重试被拒绝，使用开始新预处理冻结新配置"],
["预处理期间人工改框","发布时发现基准版本冲突，保留人工版本，不用旧结果覆盖"],
["云道S3多文件部分失败","停止后续上传；重试跳过确认成功且哈希一致的文件"],
],[2.0,4.9])
p("启动后的状态和每份上传回执会持久化。前端切换批次清缓存与轮询，拒收旧请求的晚回包；未保存 bbox 的保存、放弃、取消保护先完成后再切换。复用旧成功作业时仍读取批次最新 02，不回退已保存的人工版本。")
h("本地和远端的幂等边界",2)
p("采集下发使用 request_id，同一次网络重试必须沿用原键；正式新采集使用新键。云道S3按发布编号、文件序号和 SHA256 组成上传幂等键。手机服务和云道S3接入方仍须落实远端去重；本地记录无法单独消除远端已接收但回执未保存的时间窗口。")
p("阶段文件先写入临时目录，JSON、Excel 和 manifest 全部完整后才发布并登记。半成品不进入可选版本列表。已经登记的文件缺失或字节变化会报错，不静默覆盖，不从另一版或 Excel 自动修补。")

page("10 部署 配置与备份")
h("启动时统一数据根目录",2)
p("默认根目录为项目下 backend_workspace。服务和 CLI 必须指向同一根目录；ADF_DATA_ROOT 在 Python 模块导入时读取，应在启动前设置，不依赖后续加载模型 .env。当前后台队列和管理锁按单后端进程运行，部署使用一个 worker。")
code('$env:ADF_DATA_ROOT = "D:/work/Code/AgenticDataFlywheel/backend_workspace"\npython -m uvicorn backend.api:app --host 0.0.0.0 --port 8765 --workers 1')
h("模型和资源配置",2)
p("后端模型配置保存在 backend/.env。通用连接使用 MODEL_URL、MODEL_NAME 和凭据变量；COT_MODEL_NAME 单独选择 COT 视觉模型。任务生成可用 TASK_GENERATION_* 覆盖，树分类、摘要和质检分别有并发参数。ADARUBRIC_PYTHON 指向可运行质检依赖的 Python 环境。凭据只留在后端，设计文档和导出表不保存密钥。")
p("知识库与先验资源放在 resources/task_generation/KnowledgeBase；通过明确上传或操作者复制资源版本来初始化。手机设备配置可通过页面设置，或明确调用 initialize_settings 导入 phones、apps、phoneApps、vla 和 config；不导入旧 tasks 或运行状态。")
h("备份和恢复",2)
p("停止后端和所有写入脚本后，备份整个 backend_workspace。数据库、raw 图片及 XML、resources、输入和阶段清单必须一起保留；仅备份 Excel 不包含当前人工编辑，仅备份 SQLite 不包含文件字节。需要在线备份时应使用 SQLite 的一致性备份机制并配套文件快照，不直接复制正在写入的数据库文件。")
p("cache 用于加速，不是唯一事实来源；logs 用于排查，tmp 保存尚未发布的文件。清理前先确认没有运行中作业及实际文件引用。根目录迁移应核对保存的来源路径、索引和可读取资源，不能只改环境变量就假定所有绝对路径已迁移。")
h("本次目录更名的维护命令",2)
code("python -m backend.data_store.migrate_root check --project-root .\npython -m backend.data_store.migrate_root apply --project-root .\npython -m backend.data_store.migrate_root verify --project-root .")
p("停止服务后检查并执行切换。verify 通过且页面验收完成后运行 finalize，删除暂存旧 workspace。finalize 前且数据库无新增写入时可用 rollback 恢复；确认阶段开始后不能回滚。以上子命令共用同一前缀和 --project-root 参数，重复执行会依据迁移日志检查状态。")
p("项目根目录 .data-root-migration/ 保存 state.json 迁移日志和 original.sqlite 原库备份。新根路径登记在 SQLite 中，由统一解析器映射旧绝对路径并校验边界；冻结 JSON、Excel 和原始文件不进行字符串替换。")
h("外部尚需接通的部分",2)
p("远端手机采集服务不在本仓库，需实现指定目录写入和完成登记。云道S3适配器需实现登录、上传、超时和去重，目前默认未配置。历史模拟上传回执不算真实云道S3成功。训练配比、模型发布等演示模块不属于本次生产轨迹存储闭环。")

page("11 实例核验与维护入口")
p("验证批次 validation-20260915-155706 使用已有原始任务 AT-YYSP-AQY-001。原始输入有 10 条轨迹、142 个 response，转换得到 142 条有效步骤，没有跳过步骤。该实例用于检查新存储链路，不能作为真实手机采集对接已完成的证据。")
table(["阶段","已验证结果"],[
["01 转换","142 条步骤；JSON 与 Excel 对应；10 条轨迹身份与数量保留"],
["02 标框","142 条步骤；84 条有动作框，58 条属于非标框目标动作"],
["03 Observation","全部 142 条步骤保留；包含中间态判断与原因"],
["04 建树","134 条步骤保留在树，8 条未计入树；完整过程仍可追溯"],
["05 质检","10 条轨迹完成评分；平均分 2.22801，2 条达到当前评分阈值"],
],[1.5,5.4])
code("建树作业  7c92d69d63814a0b8cc3b79df8a526d7\n任务集    20260915_173937\n质检作业  ac0af64c0a3a4040b8041141c4840a73")
p("本次成功阶段使用 qwen3.8-max。此前 qwen3.6 系列请求有上游容量失败，早期输出格式也经过适配；失败尝试不当作成功版本。此前重跑已核验原始文件副本；本次目录迁移核对 3139 项清单，除 SQLite 登记根目录映射外，原始文件、冻结产物与缓存字节保持不变。")
p("模型语义仍需人工核查。8 条未入树步骤中，4 条理由明确引用 After，另 2 条描述动作后画面，而提示词要求判断 Before。本次保留原输出和疑点，未擅自改判，也不把技术通过等同于判断绝对正确。")
p("当前停在 05_quality，等待用户检查。实例没有执行 06 修正、07 COT、发布或上传；这些后续功能由隔离测试验证，不能写成该实例已经真实跑通。")
h("从哪里开始查一份结果",2)
p("先在页面取得 batch_id 和阶段版本，再通过统一产物接口读取 manifest，核对 files 和 source_refs。若要查当前编辑，查对应作业或会话 API；若要查对外交付文件，查 release_id 的发布清单。上传同事直接使用 UploadContext 提供的绝对路径，不需推导任何编号。")
p("维护说明文件：DATA_STORAGE.md、backend/PREPROCESSING.md、backend/COLLECTION_RUNS.md、backend/task_generation/COLLECTION_BATCHES.md 和 backend/data_publishing/INTERNAL_UPLOAD.md。实现入口分别是 data_store、preprocessing_jobs、collection_runs、trajectory_correction 与 data_publishing。")

doc.core_properties.title="数据流转与存储设计"
doc.core_properties.subject="AgenticDataFlywheel 统一数据根目录 阶段快照与业务对接"
doc.core_properties.author="AgenticDataFlywheel"
doc.core_properties.keywords="数据流转 SQLite JSON Excel 批次 预处理"
output=OUT/"数据流转与存储设计.docx"
for root in (doc.styles.element, doc.element):
    for border in list(root.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.save(output)
print(output)
print("paragraphs",len(doc.paragraphs),"tables",len(doc.tables))
